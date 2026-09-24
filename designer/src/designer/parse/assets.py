"""Картинки и встроенные шрифты из pptx. Владелец: задача T-01."""
from __future__ import annotations

import io
import struct
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from PIL import Image

from designer.contracts import Asset

ICON_MAX_WIDTH_RATIO = 0.08
ICON_SQUARE_TOLERANCE = 0.25
BACKGROUND_MIN_AREA_RATIO = 0.9
LARGE_PHOTO_MIN_AREA_RATIO = 0.2
LOGO_MIN_SLIDE_SHARE = 0.3
LOGO_POSITION_ROUND = 0.02

TTF_MAGICS = (b"\x00\x01\x00\x00", b"true", b"ttcf")
OTF_MAGIC = b"OTTO"


def _iter_pictures(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_pictures(shape.shapes)


def extract_assets(pptx_path: Path, package_dir: Path) -> list[Asset]:
    """Кладёт файлы в package_dir/assets и package_dir/fonts, возвращает опись картинок."""
    prs = Presentation(str(pptx_path))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    slide_area = slide_w * slide_h

    assets_dir = package_dir / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)

    occurrences: dict[str, list[dict]] = defaultdict(list)
    blobs: dict[str, tuple[bytes, str]] = {}
    n_slides = sum(1 for _ in prs.slides)

    def record(shape, slide_no: int) -> None:
        image = shape.image
        sha1 = image.sha1
        if sha1 not in blobs:
            blobs[sha1] = (image.blob, image.ext)
        area_ratio = (shape.width * shape.height) / slide_area if slide_area else 0.0
        width_ratio = shape.width / slide_w if slide_w else 0.0
        aspect = shape.width / shape.height if shape.height else 0.0
        occurrences[sha1].append({
            "slide": slide_no,
            "x_frac": shape.left / slide_w if slide_w else 0.0,
            "y_frac": shape.top / slide_h if slide_h else 0.0,
            "area_ratio": area_ratio,
            "width_ratio": width_ratio,
            "aspect": aspect,
        })

    for master in prs.slide_masters:
        for shape in _iter_pictures(master.shapes):
            record(shape, 0)
        for layout in master.slide_layouts:
            for shape in _iter_pictures(layout.shapes):
                record(shape, 0)

    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape in _iter_pictures(slide.shapes):
            record(shape, slide_no)

    assets: list[Asset] = []
    for sha1, occ_list in occurrences.items():
        blob, ext = blobs[sha1]
        asset_id = sha1[:12]
        rel_path = f"assets/{asset_id}.{ext}"
        out_path = package_dir / rel_path
        if not out_path.exists():
            out_path.write_bytes(blob)

        try:
            with Image.open(io.BytesIO(blob)) as img:
                width_px, height_px = img.size
        except Exception:
            width_px, height_px = 0, 0

        kind = _classify_kind(occ_list, ext, n_slides)
        used_on = sorted({occ["slide"] for occ in occ_list})

        assets.append(Asset(
            id=asset_id,
            path=rel_path,
            kind=kind,
            width_px=width_px,
            height_px=height_px,
            sha1=sha1,
            used_on=used_on,
        ))

    fonts_dir = package_dir / "fonts"
    fonts_dir.mkdir(parents=True, exist_ok=True)
    _extract_embedded_fonts(prs, fonts_dir)

    return assets


def _classify_kind(occurrences: list[dict], ext: str, n_slides: int) -> str:
    if any(occ["area_ratio"] >= BACKGROUND_MIN_AREA_RATIO for occ in occurrences):
        return "background"

    in_layout_or_master = any(occ["slide"] == 0 for occ in occurrences)
    slide_positions: dict[tuple[float, float], set[int]] = defaultdict(set)
    for occ in occurrences:
        if occ["slide"] == 0:
            continue
        key = (
            round(occ["x_frac"] / LOGO_POSITION_ROUND) * LOGO_POSITION_ROUND,
            round(occ["y_frac"] / LOGO_POSITION_ROUND) * LOGO_POSITION_ROUND,
        )
        slide_positions[key].add(occ["slide"])
    same_position_share = max((len(v) for v in slide_positions.values()), default=0) / n_slides if n_slides else 0
    small_in_layout = in_layout_or_master and any(
        occ["slide"] == 0 and occ["area_ratio"] < BACKGROUND_MIN_AREA_RATIO for occ in occurrences
    )
    if same_position_share >= LOGO_MIN_SLIDE_SHARE or small_in_layout:
        return "logo"

    width_ratio = max((occ["width_ratio"] for occ in occurrences), default=0.0)
    aspects = [occ["aspect"] for occ in occurrences if occ["aspect"] > 0]
    near_square = aspects and all(abs(a - 1) <= ICON_SQUARE_TOLERANCE for a in aspects)
    if width_ratio <= ICON_MAX_WIDTH_RATIO and near_square:
        return "icon"

    max_area = max((occ["area_ratio"] for occ in occurrences), default=0.0)
    if ext.lower() in ("jpg", "jpeg") or max_area > LARGE_PHOTO_MIN_AREA_RATIO:
        return "photo"

    return "decor"


# ---------- Встроенные шрифты ----------

def _parse_eot_font_data(data: bytes) -> bytes | None:
    """Возвращает встроенные байты шрифта из контейнера EOT, либо None при неверном формате."""
    if len(data) < 82:
        return None
    try:
        eot_size, font_data_size = struct.unpack_from("<II", data, 0)
        magic = struct.unpack_from("<H", data, 34)[0]
    except struct.error:
        return None
    if magic != 0x504C or eot_size != len(data):
        return None
    off = eot_size - font_data_size
    if off < 82 or off > eot_size:
        return None
    return data[off:eot_size]


def _font_extension(payload: bytes) -> str | None:
    if payload[:4] == OTF_MAGIC:
        return "otf"
    if payload[:4] in TTF_MAGICS:
        return "ttf"
    return None


def _slug(text: str) -> str:
    safe = "".join(ch if ch.isalnum() else "-" for ch in text)
    while "--" in safe:
        safe = safe.replace("--", "-")
    return safe.strip("-") or "font"


def _extract_embedded_fonts(prs, fonts_dir: Path) -> dict[str, str]:
    """Сохраняет встроенные шрифты в fonts_dir, возвращает family -> путь внутри пакета.

    Для шрифта, чьи данные после снятия обёртки EOT не опознаются как TTF/OTF по первым
    байтам (частый случай — PowerPoint сжимает поток алгоритмом MicroType Express),
    файл сохраняется как есть, а ссылка в словаре не выставляется.
    """
    part = prs.part
    lst = part._element.find(qn("p:embeddedFontLst"))
    mapping: dict[str, str] = {}
    if lst is None:
        return mapping

    for embedded in lst.findall(qn("p:embeddedFont")):
        font_el = embedded.find(qn("p:font"))
        if font_el is None:
            continue
        family = font_el.get("typeface")
        if not family:
            continue
        for style_tag in ("regular", "bold", "italic", "boldItalic"):
            style_el = embedded.find(qn(f"p:{style_tag}"))
            if style_el is None:
                continue
            rid = style_el.get(qn("r:id"))
            if not rid:
                continue
            try:
                font_part = part.related_part(rid)
            except KeyError:
                continue
            data = font_part.blob
            base_name = f"{_slug(family)}-{style_tag}"
            payload = _parse_eot_font_data(data)
            ext = _font_extension(payload) if payload else None
            if ext:
                out_path = fonts_dir / f"{base_name}.{ext}"
                if not out_path.exists():
                    out_path.write_bytes(payload)
                if family not in mapping or style_tag == "regular":
                    mapping[family] = f"fonts/{out_path.name}"
            else:
                out_path = fonts_dir / f"{base_name}.fntdata"
                if not out_path.exists():
                    out_path.write_bytes(data)
    return mapping
