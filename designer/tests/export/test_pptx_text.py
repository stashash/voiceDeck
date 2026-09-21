"""Тесты set_text: оформление образца, абзацы, очистка фигуры, кегль подгонки. Задача T-08, кегль — T-20."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from designer.export.pptx_text import set_text

SAMPLE_FONT = "PT Sans"
SAMPLE_SIZE = Pt(22)
SAMPLE_COLOR = "0077FF"


def _textbox():
    """Фигура с образцом: первый фрагмент оформлен, дальше идут чужие фрагменты и абзац."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(914400), Emu(914400), Emu(3657600), Emu(1828800))
    frame = box.text_frame
    first = frame.paragraphs[0]
    run = first.add_run()
    run.text = "Заголовок образца"
    run.font.name = SAMPLE_FONT
    run.font.size = SAMPLE_SIZE
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(SAMPLE_COLOR)
    tail = first.add_run()
    tail.text = " хвост"
    tail.font.size = Pt(9)
    second = frame.add_paragraph()
    second_run = second.add_run()
    second_run.text = "Второй абзац образца"
    second_run.font.size = Pt(11)
    return prs, box


def _reopen(prs, tmp_path):
    path = tmp_path / "text.pptx"
    prs.save(str(path))
    return Presentation(str(path)).slides[0].shapes[0]


def test_set_text_keeps_first_run_style(tmp_path):
    prs, box = _textbox()

    set_text(box, "Новый заголовок")

    shape = _reopen(prs, tmp_path)
    run = shape.text_frame.paragraphs[0].runs[0]
    assert shape.text_frame.text == "Новый заголовок"
    assert run.font.name == SAMPLE_FONT
    assert run.font.size == SAMPLE_SIZE
    assert run.font.bold is True
    assert str(run.font.color.rgb) == SAMPLE_COLOR


def test_set_text_drops_other_runs_and_paragraphs():
    _, box = _textbox()

    set_text(box, "Одна строка")

    paragraphs = box.text_frame.paragraphs
    assert len(paragraphs) == 1
    assert len(paragraphs[0].runs) == 1
    assert "хвост" not in box.text_frame.text
    assert "Второй абзац" not in box.text_frame.text


def test_set_text_newline_gives_paragraph_with_same_style(tmp_path):
    prs, box = _textbox()

    set_text(box, "Первая\nВторая\nТретья")

    shape = _reopen(prs, tmp_path)
    paragraphs = shape.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Первая", "Вторая", "Третья"]
    for paragraph in paragraphs:
        assert paragraph.runs[0].font.size == SAMPLE_SIZE
        assert paragraph.runs[0].font.name == SAMPLE_FONT


def test_set_text_empty_clears_shape_but_keeps_it(tmp_path):
    prs, box = _textbox()

    set_text(box, "")

    shape = _reopen(prs, tmp_path)
    assert shape.text_frame.text == ""
    assert len(shape.text_frame.paragraphs) == 1


def test_set_text_accepts_shape_xml():
    _, box = _textbox()

    set_text(box._element, "Текст по XML фигуры")

    assert box.text_frame.text == "Текст по XML фигуры"


def test_set_text_leaves_shapes_without_text_frame_alone():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture_free_shape = slide.shapes.add_table(2, 2, Emu(0), Emu(0), Emu(914400), Emu(914400))

    set_text(picture_free_shape, "мимо")

    assert picture_free_shape.has_table


def test_set_text_with_size_applies_it_to_every_fragment(tmp_path):
    prs, box = _textbox()

    set_text(box, "Первая\nВторая", size_pt=14)

    shape = _reopen(prs, tmp_path)
    paragraphs = shape.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Первая", "Вторая"]
    for paragraph in paragraphs:
        assert paragraph.runs[0].font.size == Pt(14)


def test_set_text_with_size_disables_autofit(tmp_path):
    prs, box = _textbox()

    set_text(box, "Подогнанный текст", size_pt=12)

    shape = _reopen(prs, tmp_path)
    body_pr = shape.text_frame._txBody.find(qn("a:bodyPr"))
    assert body_pr.find(qn("a:noAutofit")) is not None
    assert body_pr.find(qn("a:normAutofit")) is None


def test_set_text_without_size_keeps_sample_size(tmp_path):
    prs, box = _textbox()

    set_text(box, "Без подгонки")

    shape = _reopen(prs, tmp_path)
    assert shape.text_frame.paragraphs[0].runs[0].font.size == SAMPLE_SIZE
