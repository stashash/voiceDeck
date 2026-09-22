"""Разбор выданных шаблонов в паттерны. Факты сняты с самих файлов."""
import time
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu

from designer.contracts import SlideKind
from designer.parse import geometry as geo
from designer.parse.patterns import extract_layouts, extract_patterns

_parsed: dict[Path, list] = {}


def parse(path: Path) -> list:
    if path not in _parsed:
        _parsed[path] = extract_patterns(path)
    return _parsed[path]


def pick(templates: list[Path], part: str) -> Path:
    for path in templates:
        if part.lower() in path.stem.lower():
            return path
    pytest.skip(f"среди шаблонов нет файла с «{part}» в имени")


def slide_of(templates: list[Path], part: str, number: int):
    return next(p for p in parse(pick(templates, part)) if p.source_slide == number)


def shape_ids(shapes, found: set[int]) -> set[int]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            shape_ids(shape.shapes, found)
        else:
            found.add(shape.shape_id)
    return found


# ---------- факты выданных шаблонов ----------

def test_cards_grid_three_by_two(templates):
    pattern = slide_of(templates, "VK Tech", 17)
    assert [s.role for s in pattern.slots] == ["title"]
    assert len(pattern.groups) == 1
    group = pattern.groups[0]
    assert (group.direction, group.cols, group.rows) == ("grid", 3, 2)
    assert len(group.units) == 6
    assert group.max_units == 6
    assert len(group.unit_slots) == 2
    heading, body = group.unit_slots
    assert heading.role == "heading" and body.role == "body"
    assert heading.sample_text and body.sample_text
    assert len(heading.sample_text) < len(body.sample_text)
    assert pattern.kind is SlideKind.cards


def test_row_of_three_blocks(templates):
    pattern = slide_of(templates, "WorkSpace", 6)
    group = max(pattern.groups, key=lambda g: len(g.units))
    assert group.direction == "row"
    assert len(group.units) == 3
    assert group.step[0] == pytest.approx(0.30, abs=0.02)
    assert group.step[1] == 0.0


def test_column_of_four_blocks(templates):
    pattern = slide_of(templates, "WorkSpace", 22)
    group = max(pattern.groups, key=lambda g: len(g.units))
    assert group.direction == "column"
    assert len(group.units) == 4
    assert group.step[1] == pytest.approx(0.20, abs=0.02)
    assert group.step[0] == 0.0


def test_dark_template(templates):
    patterns = parse(pick(templates, "WorkSpace"))
    dark = sum(1 for p in patterns if p.theme == "dark")
    assert dark / len(patterns) >= 0.8


@pytest.mark.parametrize("part", ["Education", "VK Tech"])
def test_light_templates(templates, part):
    patterns = parse(pick(templates, part))
    light = sum(1 for p in patterns if p.theme == "light")
    assert light > len(patterns) / 2


def test_pattern_per_slide_with_real_shape_ids(templates):
    for path in templates:
        presentation = Presentation(str(path))
        patterns = parse(path)
        assert len(patterns) == len(presentation.slides)
        for slide, pattern in zip(presentation.slides, patterns):
            real = shape_ids(slide.shapes, set())
            used = set(pattern.decor_shape_ids)
            used |= {s.shape_id for s in pattern.slots}
            used |= {a.shape_id for a in pattern.areas if a.shape_id is not None}
            for group in pattern.groups:
                used |= {s.shape_id for s in group.unit_slots}
                for unit in group.units:
                    used |= set(unit.shape_ids)
            assert used == real, f"{path.stem}, слайд {pattern.source_slide}"


def test_most_slides_get_a_kind(templates):
    for path in templates:
        patterns = parse(path)
        known = [p for p in patterns if p.kind is not SlideKind.other]
        assert len(known) >= len(patterns) / 2, path.stem


def test_slots_and_groups_are_well_formed(templates):
    for path in templates:
        for pattern in parse(path):
            assert 0.0 <= pattern.kind_confidence <= 1.0
            for slot in pattern.slots:
                assert slot.max_chars >= 1 and slot.max_lines >= 1
                assert slot.style.size_pt and slot.style.size_pt > 0
            for group in pattern.groups:
                assert len(group.units) == group.cols * group.rows
                assert group.max_units >= len(group.units)
                assert group.unit_size[0] > 0 and group.unit_size[1] > 0
                assert group.unit_slots or group.unit_areas


@pytest.mark.parametrize("number", [21, 22, 23])
def test_numbers_and_captions_under_them_are_linked_groups(templates, number):
    pattern = slide_of(templates, "Education", number)
    numbers = [g for g in pattern.groups if any(s.role == "number" for s in g.unit_slots)]
    captions = [g for g in pattern.groups if g.id not in {n.id for n in numbers}]
    assert len(numbers) == 1 and len(captions) == 1
    circles, texts = numbers[0], captions[0]
    assert len(circles.units) == len(texts.units)
    assert circles.linked_group_ids == [texts.id]
    assert texts.linked_group_ids == [circles.id]
    assert pattern.primary_group_id == texts.id
    assert pattern.kind is SlideKind.steps


@pytest.mark.parametrize("number", [12, 13])
def test_team_slides_keep_places_for_photos(templates, number):
    pattern = slide_of(templates, "VK Tech", number)
    assert pattern.kind is SlideKind.team
    assert pattern.needs_images
    main = next(g for g in pattern.groups if g.id == pattern.primary_group_id)
    holders = [a for a in main.unit_areas if a.placeholder]
    assert holders and all(a.kind == "image" for a in holders)


@pytest.mark.parametrize("part,number", [("Education", 18), ("VK Tech", 33), ("WorkSpace", 17)])
def test_slide_built_around_one_number(templates, part, number):
    pattern = slide_of(templates, part, number)
    assert pattern.kind is SlideKind.big_number
    numbers = [s for s in pattern.slots if s.role == "number"]
    assert numbers
    assert max(s.style.size_pt for s in numbers) >= 40


def test_pattern_with_groups_names_its_main_group(templates):
    for path in templates:
        for pattern in parse(path):
            if not pattern.groups:
                continue
            ids = {g.id for g in pattern.groups}
            assert pattern.primary_group_id in ids, f"{path.stem}, слайд {pattern.source_slide}"


def test_linked_groups_answer_each_other(templates):
    for path in templates:
        for pattern in parse(path):
            by_id = {g.id: g for g in pattern.groups}
            for group in pattern.groups:
                for other_id in group.linked_group_ids:
                    other = by_id[other_id]
                    assert group.id in other.linked_group_ids
                    assert len(other.units) == len(group.units)


# ---------- дефекты живой колоды: оформление с рамкой и образец диаграммы из фигур ----------

def test_decor_carries_a_frame_for_every_shape(templates):
    """По рамке оформления вёрстка решает, мешает ли оно содержимому слайда."""
    for path in templates:
        for pattern in parse(path):
            where = f"{path.stem}, слайд {pattern.source_slide}"
            assert {d.shape_id for d in pattern.decor} == set(pattern.decor_shape_ids), where
            for shape in pattern.decor:
                long_side = max(shape.box[2], shape.box[3])
                assert long_side > 0, where
                flat = min(shape.box[2], shape.box[3]) <= 0.01 < long_side
                assert not flat or shape.kind in ("line", "image"), where


def test_row_of_bars_of_the_template_is_a_place_for_a_chart(templates):
    """Полосы разной длины на общей оси это образец диаграммы, а не пять картинок."""
    pattern = slide_of(templates, "VK Tech", 52)
    charts = [area for area in pattern.areas if area.kind == "chart"]
    assert len(charts) == 1
    box = charts[0].box
    assert geo.area(box) >= 0.2
    left = [area for area in pattern.areas
            if area.kind == "image" and geo.covered(area.box, box) > 0.5]
    assert not left, [area.id for area in left]


def test_slide_without_bars_gets_no_chart_area(templates):
    pattern = slide_of(templates, "VK Tech", 17)
    assert not [area for area in pattern.areas if area.kind == "chart"]


def test_layouts_carry_placeholders(templates):
    layouts = extract_layouts(pick(templates, "VK Tech"))
    assert layouts
    assert any(layout.placeholders for layout in layouts)
    assert {layout.master_index for layout in layouts} == {0, 1}


def test_one_template_parses_under_ninety_seconds(templates):
    started = time.time()
    extract_patterns(templates[0])
    assert time.time() - started < 90


def test_foreign_pptx_parses(foreign_pptx):
    patterns = extract_patterns(foreign_pptx)
    assert patterns
    assert [p.source_slide for p in patterns] == list(range(1, len(patterns) + 1))
    assert all(p.theme in ("light", "dark") for p in patterns)


# ---------- случаи, собранные прямо в тесте ----------

def _blank_deck() -> tuple:
    presentation = Presentation()
    presentation.slide_width = Emu(9144000)
    presentation.slide_height = Emu(5143500)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation, slide


def test_dark_picture_on_the_whole_frame_makes_theme_dark(tmp_path):
    picture = tmp_path / "fon.png"
    Image.new("RGB", (64, 36), (18, 20, 24)).save(picture)
    presentation, slide = _blank_deck()
    slide.shapes.add_picture(
        str(picture), 0, 0, presentation.slide_width, presentation.slide_height
    )
    deck = tmp_path / "fon.pptx"
    presentation.save(str(deck))

    pattern = extract_patterns(deck)[0]
    assert pattern.theme == "dark"
    assert pattern.background_asset


def test_group_coordinates_are_recalculated(tmp_path):
    presentation, slide = _blank_deck()
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Emu(1000000), Emu(500000), Emu(2000000), Emu(1000000))
    box.text_frame.text = "текст в группе"

    xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    for tag, x, y in (("a:off", 0, 0), ("a:chOff", 0, 0)):
        node = xfrm.find(qn(tag))
        node.set("x", str(x))
        node.set("y", str(y))
    for tag, cx, cy in (("a:ext", 8000000, 4000000), ("a:chExt", 4000000, 2000000)):
        node = xfrm.find(qn(tag))
        node.set("cx", str(cx))
        node.set("cy", str(cy))
    deck = tmp_path / "gruppa.pptx"
    presentation.save(str(deck))

    pattern = extract_patterns(deck)[0]
    slot = next(s for s in pattern.slots if s.sample_text)
    assert slot.box == pytest.approx(
        (2000000 / 9144000, 1000000 / 5143500, 4000000 / 9144000, 2000000 / 5143500), abs=0.002
    )


def test_gray_round_shape_is_a_place_for_a_photo(tmp_path):
    presentation, slide = _blank_deck()
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Emu(1000000), Emu(800000), Emu(1500000), Emu(1500000)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xC4, 0xC4, 0xC4)
    deck = tmp_path / "zaglushka.pptx"
    presentation.save(str(deck))

    pattern = extract_patterns(deck)[0]
    area = next(a for a in pattern.areas if a.shape_id == shape.shape_id)
    assert (area.kind, area.placeholder) == ("image", True)
    assert not pattern.slots


def test_colored_circle_with_a_number_stays_text(tmp_path):
    presentation, slide = _blank_deck()
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Emu(1000000), Emu(800000), Emu(700000), Emu(700000)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x00, 0x77, 0xFF)
    shape.text_frame.text = "1"
    deck = tmp_path / "kruzhok.pptx"
    presentation.save(str(deck))

    pattern = extract_patterns(deck)[0]
    assert pattern.areas == []
    assert [s.shape_id for s in pattern.slots] == [shape.shape_id]


def test_large_flat_gray_picture_asks_for_an_own_photo(tmp_path):
    picture = tmp_path / "seroe.png"
    Image.new("RGB", (64, 64), (196, 196, 196)).save(picture)
    presentation, slide = _blank_deck()
    shape = slide.shapes.add_picture(
        str(picture), Emu(500000), Emu(500000), Emu(4000000), Emu(3000000)
    )
    deck = tmp_path / "foto.pptx"
    presentation.save(str(deck))

    pattern = extract_patterns(deck)[0]
    area = next(a for a in pattern.areas if a.shape_id == shape.shape_id)
    assert area.placeholder
    assert pattern.needs_images


def test_flat_shape_of_the_decor_is_a_line(tmp_path):
    presentation, slide = _blank_deck()
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(500000), Emu(2000000), Emu(6000000), Emu(12700)
    )
    deck = tmp_path / "liniya.pptx"
    presentation.save(str(deck))

    decor = extract_patterns(deck)[0].decor
    assert [(d.shape_id, d.kind) for d in decor] == [(shape.shape_id, "line")]


def test_font_size_falls_back_to_default(tmp_path):
    presentation, slide = _blank_deck()
    box = slide.shapes.add_textbox(Emu(500000), Emu(500000), Emu(3000000), Emu(600000))
    box.text_frame.text = "без заданного кегля"
    assert box.text_frame.paragraphs[0].runs[0].font.size is None
    deck = tmp_path / "kegel.pptx"
    presentation.save(str(deck))

    slot = extract_patterns(deck)[0].slots[0]
    assert slot.style.size_pt == 18.0
    assert slot.max_lines == 2


def test_card_header_plate_moves_with_its_card(templates):
    pattern = slide_of(templates, "WorkSpace", 16)
    units = next(g for g in pattern.groups if g.id == pattern.primary_group_id).units
    # Плашки 692 и 696 есть у первых двух карточек, третья залита целиком.
    assert 692 in units[0].shape_ids and 696 in units[1].shape_ids
    assert not {692, 696} & {shape.shape_id for shape in pattern.decor}
