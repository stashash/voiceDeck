"""Stage-2 chart contracts, source adapters and both export representations."""
from io import BytesIO
import json
import math
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from typing import get_args

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pydantic import ValidationError

from designer import store
from designer.api.app import app
from designer.audit.checks_integrity import check_chart_labels
from designer.contracts import ChartSpec, Deck, DeckPlan, DesignSystem, Element, Margins, Scene, Series, Tokens
from designer.export.html import render_deck
from designer.export.svg_chart import render_chart
from designer.insights.plan import plan_table
from designer.insights.suggest import Suggestion, _Pick, insight_slide
from designer.insights.table import read_tables
from designer.viz.pptx_native import _CHART_TYPES, add_chart


def tokens():
    return Tokens(colors=[], fonts=[], type_scale=[], margins=Margins(left=0, top=0, right=0, bottom=0))


def scatter():
    return ChartSpec(type="scatter", x_axis_title="Spend, USD", y_axis_title="Leads", unit="leads",
                     series=[Series(name="Leads", x_values=[-1, 0, 8], values=[-2, 0, 4])])


def test_declared_chart_types_match_native_export_types_and_legacy_model_cannot_retype_them():
    assert set(get_args(ChartSpec.model_fields["type"].annotation)) == set(_CHART_TYPES)
    assert not {"scatter", "grouped_column", "stacked_column"} & set(get_args(_Pick.model_fields["visual"].annotation))


@pytest.mark.parametrize("patch", [
    {"categories": ["fake"]}, {"series": []},
    {"series": [{"name": "a", "values": [1]}]},
    {"series": [{"name": "a", "values": [1, 2], "x_values": [1]}]},
    {"series": [{"name": "a", "values": [math.inf], "x_values": [1]}]},
    {"series": [{"name": "a", "values": [1], "x_values": [math.nan]}]},
])
def test_scatter_contract_rejects_unpaired_or_nonfinite_data(patch):
    with pytest.raises(ValidationError):
        ChartSpec.model_validate(scatter().model_dump() | patch)


@pytest.mark.parametrize("kind", ["grouped_column", "stacked_column"])
def test_column_contract_does_not_impute_missing_points(kind):
    with pytest.raises(ValidationError, match="match the categories"):
        ChartSpec(type=kind, categories=["A", "B"], series=[Series(name="a", values=[1])])


def test_stacked_contract_rejects_negative_or_overflowing_composition():
    for values in [[-1], [1e308]]:
        with pytest.raises(ValidationError, match="nonnegative"):
            ChartSpec(type="stacked_column", categories=["A"],
                      series=[Series(name="a", values=values), Series(name="b", values=values)])


def test_scatter_native_roundtrip_preserves_numeric_xy_axes_and_embedded_workbook(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_chart(slide, scatter(), (0.1, 0.1, 0.8, 0.7), (prs.slide_width, prs.slide_height), tokens())
    path = tmp_path / "scatter.pptx"
    prs.save(path)
    chart = Presentation(path).slides[0].shapes[0].chart
    assert chart.chart_type == XL_CHART_TYPE.XY_SCATTER
    assert tuple(chart.series[0].values) == (-2, 0, 4)
    assert tuple(chart.series[0].iter_values()) == (-2, 0, 4)
    assert chart.category_axis.axis_title.text_frame.text == "Spend, USD"
    assert chart.value_axis.axis_title.text_frame.text == "Leads"
    with ZipFile(path) as archive:
        xml = ET.fromstring(archive.read("ppt/charts/chart1.xml"))
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        assert [float(v.text) for v in xml.findall(".//c:xVal//c:pt/c:v", ns)] == [-1, 0, 8]
        assert len(xml.findall(".//c:valAx", ns)) == 2
        assert not xml.findall(".//c:catAx", ns)
        book = next(n for n in archive.namelist() if n.endswith(".xlsx"))
        rows = list(load_workbook(BytesIO(archive.read(book)), data_only=True).active.values)
        assert [tuple(r[:2]) for r in rows[1:]] == [(-1, -2), (0, 0), (8, 4)]


def test_scatter_svg_uses_numeric_spacing_no_connecting_line_and_escaped_axis_labels():
    spec = scatter().model_copy(update={"x_axis_title": "<Spend>"})
    svg = render_chart(spec, [])
    root = ET.fromstring(svg)
    points = root.findall("circle")
    positions = [float(p.attrib["cx"]) for p in points]
    assert len(points) == 3
    assert (positions[2] - positions[1]) / (positions[1] - positions[0]) == pytest.approx(8, abs=0.02)
    assert 'chart-line' not in svg and "&lt;Spend&gt;" in svg
    assert all(20 <= float(p.attrib["cy"]) <= 310 for p in points)


@pytest.mark.parametrize("kind", ["grouped_column", "stacked_column"])
def test_column_svg_geometry_retains_separate_series_and_correct_stacking(kind):
    spec = ChartSpec(type=kind, categories=["A", "B"],
                     series=[Series(name="part 1", values=[2, 3]), Series(name="part 2", values=[3, 5])])
    root = ET.fromstring(render_chart(spec, []))
    bars = [e for e in root.findall("rect") if e.attrib.get("class") == "chart-bar"]
    assert len(bars) == 4
    a, b = bars[:2]
    assert (a.attrib["x"] == b.attrib["x"]) == (kind == "stacked_column")
    if kind == "stacked_column":
        assert float(b.attrib["y"]) + float(b.attrib["height"]) == pytest.approx(float(a.attrib["y"]), abs=0.1)
        assert float(a.attrib["height"]) / float(b.attrib["height"]) == pytest.approx(2 / 3, abs=0.002)


def test_grouped_negative_values_have_positive_geometry_on_opposite_sides_of_zero():
    spec = ChartSpec(type="grouped_column", categories=["A"],
                     series=[Series(name="a", values=[-2]), Series(name="b", values=[3])])
    root = ET.fromstring(render_chart(spec, []))
    bars = [e for e in root.findall("rect") if e.attrib.get("class") == "chart-bar"]
    zero = float(root.find("line").attrib["y1"])
    assert float(bars[0].attrib["y"]) == zero
    assert float(bars[1].attrib["y"]) < zero
    assert all(float(b.attrib["height"]) > 0 for b in bars)


def test_export_fails_closed_for_undeclared_chart_type():
    with pytest.raises(ValueError, match="Unsupported"):
        render_chart(scatter().model_copy(update={"type": "radar"}), [])


@pytest.mark.parametrize("kind,goal,data,fields", [
    ("scatter", "relationship", b"x,y\n0,0\n1,2\n2,\n3,6\n4,8\n5,10\n", ["x", "y"]),
    ("grouped_column", "compare", b"category,actual,target\nA,0,1\nB,2,3\nC,3,4\n", ["category", "actual", "target"]),
    ("stacked_column", "stacked composition", b"category,online_usd,offline_usd\nA,0,2\nB,3,4\nC,5,\n", ["category", "online_usd", "offline_usd"]),
])
def test_plan_materialization_and_html_deck_export_preserve_source_rows(tmp_path, monkeypatch, kind, goal, data, fields):
    monkeypatch.setenv("DESIGNER_DATA_DIR", str(tmp_path / "data"))
    response = TestClient(app).post("/insights/plan/materialize", files={"file": ("source.csv", data, "text/csv")},
                                    data={"goal": goal, "candidate_kind": kind, "required_fields": json.dumps(fields)})
    assert response.status_code == 200, response.text
    saved = store.load_insights(response.json()["id"])
    suggestion = Suggestion.model_validate(saved["suggestions"][0])
    intent = insight_slide(suggestion)
    assert intent.chart.type == kind
    assert intent.chart.series[0].values[0] == 0
    assert len(intent.provenance.excluded_row_lines) == (0 if kind == "grouped_column" else 1)
    assert intent.provenance.columns == fields
    assert intent.provenance.visualization == kind
    if kind == "scatter":
        assert intent.chart.series[0].x_values == [0, 1, 3, 4, 5]
        assert intent.chart.series[0].values == [0, 2, 6, 8, 10]
    else:
        assert len(intent.chart.series) == 2
        assert all(len(s.values) == (3 if kind == "grouped_column" else 2) for s in intent.chart.series)
    element = Element(id="chart", type="chart", box=(0.1, 0.1, 0.8, 0.8), chart=intent.chart)
    scene = Scene(slide_id="s1", pattern_id="p1", elements=[element])
    ds = DesignSystem(id="ds", source_file="", slide_size_emu=(9144000, 5143500), tokens=tokens())
    deck = Deck(id="deck", design_system_id="ds", variant="a", plan=DeckPlan(title="Test", purpose="test", slides=[intent]), scenes=[scene])
    assert f"chart-{kind}" in render_deck(deck, ds, tmp_path)


@pytest.mark.parametrize("goal,headers,rows", [
    ("shares", "category,a_usd,b_usd", "A,1,2\nB,3,4"),
    ("stacked", "category,a_usd,b_eur", "A,1,2\nB,3,4"),
    ("stacked", "category,a,b", "A,1,2\nB,3,4"),
    ("stacked", "category,a_usd,b_usd", "A,-1,2\nB,3,4"),
    ("stacked", "category,actual_usd,target_usd", "A,1,2\nB,3,4"),
])
def test_stacked_is_not_inferred_for_unsafe_components(goal, headers, rows):
    table = read_tables(f"{headers}\n{rows}\n".encode(), "data.csv")[0]
    assert "stacked_column" not in {c.kind for c in plan_table(table, goal).candidates}


def test_scatter_audit_requires_axis_names_instead_of_categories():
    scene = Scene(slide_id="s", pattern_id="p", elements=[Element(id="c", type="chart", box=(0, 0, 1, 1), chart=scatter())])
    ds = DesignSystem(id="ds", source_file="", slide_size_emu=(9144000, 5143500), tokens=tokens())
    assert check_chart_labels([scene], ds) == []


def test_precision_loss_keeps_extended_candidate_unrenderable():
    table = read_tables(b"category,actual,target\nA,9007199254740993,2\nB,9007199254740995,4\n", "data.csv")[0]
    candidate = next(c for c in plan_table(table, "compare").candidates if c.kind == "grouped_column")
    assert not candidate.renderable
    assert "диапазон чисел" in candidate.reason
