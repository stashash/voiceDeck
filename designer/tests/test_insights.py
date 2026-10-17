"""Подсказки по таблице: разбор csv и xlsx, кандидаты точной арифметикой, сверка чисел модели, API и план."""
from __future__ import annotations

import io
import json
from decimal import Decimal

import httpx
import pytest
from fastapi.testclient import TestClient

from designer import store
from designer.api.app import app, get_llm_client
from designer.contracts import DeckPlan, SlideIntent, SlideKind
from designer.insights.facts import find_facts
from designer.insights.suggest import insert_insight_slides, numbers_in, suggest
from designer.insights.table import TableError, parse_number, read_table
from designer.llm.client import LlmClient

QUARTERS = ("Квартал;Выручка, млн руб.;Клиенты;Доля онлайн, %\n"
            "I кв 2026;120;1 250;31,5\nII кв 2026;135;1 410;33\n"
            "III кв 2026;150;1 380;36,2\nIV кв 2026;177;1 690;40,1\n").encode("cp1251")
REGIONS = ("Регион,Q3 2026,Q4 2026\nСеверо-Запад,48,61\nЦентр,90,95\nУрал,30,28\n"
           "Сибирь,22,31\nЮг,40,40\nИтого,230,255\n").encode()


def _llm(payload: dict) -> LlmClient:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(payload, ensure_ascii=False)}}]})
    return LlmClient("http://test/v1", "model", transport=httpx.MockTransport(handler))


@pytest.mark.parametrize("text,value", [("1 234,5", "1234.5"), ("-12.3", "-12.3"), ("18 %", "18"),
                                        ("−5", "-5"), ("120 ₽", "120"), ("1 250", "1250")])
def test_numbers_in_russian_formats(text, value):
    assert parse_number(text) == Decimal(value)


@pytest.mark.parametrize("text", ["=SUM(A1:A3)", "IV кв", "12.05.2026", "", "1,2,3"])
def test_not_numbers(text):
    assert parse_number(text) is None


def test_cp1251_semicolon_csv_reads():
    table = read_table(QUARTERS, "итоги.csv")
    assert table.headers[1] == "Выручка, млн руб."
    assert table.numbers[0][2] == Decimal(1250)


def test_xlsx_first_sheet_with_cached_formula_values():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.append(["Квартал", "Выручка"])
    sheet.append(["Q1 2026", 100])
    sheet.append(["Q2 2026", 125.5])
    buffer = io.BytesIO()
    book.save(buffer)
    table = read_table(buffer.getvalue(), "q.xlsx")
    assert table.rows == [["Q1 2026", "100"], ["Q2 2026", "125.5"]]


@pytest.mark.parametrize("data,name,message", [
    (b"a,b\n1,2\n", "x.csv", "две строки"),
    (b",b\n1,2\n3,4\n", "x.csv", "Шапка"),
    (b"a,b\n1,2\n3,4\n", "x.pdf", "xlsx или csv"),
    (b"x" * (1024 * 1024 + 1), "x.csv", "1 МБ"),
], ids=["one-row", "empty-header", "pdf", "too-big"])
def test_bad_tables_say_what_to_fix(data, name, message):
    with pytest.raises(TableError, match=message):
        read_table(data, name)


def test_quarter_facts_are_exact():
    facts = {f.kind + f.statement[:6]: f for f in find_facts(read_table(QUARTERS, "q.csv"))}
    revenue = next(f for f in facts.values() if f.kind == "change_last" and f.statement.startswith("Выручка"))
    assert revenue.numbers == ["+18 %", "150", "177"]
    share = next(f for f in facts.values() if f.kind == "change_last" and f.statement.startswith("Доля"))
    assert share.numbers[0] == "+3,9 п. п."  # доля меняется в процентных пунктах


def test_total_row_is_not_a_region():
    facts = find_facts(read_table(REGIONS, "r.csv"))
    grew = next(f for f in facts if f.kind == "rows_grew")
    assert grew.numbers == ["3", "5"]
    total = next(f for f in facts if f.kind == "total_change")
    assert total.numbers == ["+10,9 %", "230", "255"]


def test_model_number_not_from_facts_is_dropped():
    table = read_table(REGIONS, "r.csv")
    facts = find_facts(table)
    grew = next(f for f in facts if f.kind == "rows_grew")
    answer = {"suggestions": [
        {"fact_ids": [grew.id], "headline": "В Q4 выросли 3 из 5 регионов", "number": "3 из 5", "visual": "column",
         "why": "видно, где рост"},
        {"fact_ids": [grew.id], "headline": "Выручка выросла на 25 %", "number": "+25 %", "visual": "big_number",
         "why": "придуманное число"},
    ]}
    suggestions, note = suggest(table, facts, "итоги", "директора", _llm(answer))
    assert note == ""
    assert suggestions[0].headline == "В Q4 выросли 3 из 5 регионов" and suggestions[0].by_model
    assert all("25" not in numbers_in(s.headline + " " + s.number) for s in suggestions)
    assert len(suggestions) >= 3  # добрано кандидатами кода
    assert not suggestions[-1].by_model


def test_model_failure_falls_back_to_code_facts():
    def handler(request):
        return httpx.Response(200, json={"choices": [{"message": {"content": "не json"}}]})
    table = read_table(QUARTERS, "q.csv")
    suggestions, note = suggest(table, find_facts(table), "", "", LlmClient("http://test/v1", "m", transport=httpx.MockTransport(handler)))
    assert note and len(suggestions) == 3 and not any(s.by_model for s in suggestions)


def test_model_headline_must_keep_period_sign_and_metric():
    """Ред-тим 17.10: «за 2026 год» при сравнении кварталов, спад при росте, чужая метрика не проходят."""
    table = read_table(QUARTERS, "q.csv")
    facts = find_facts(table)
    total = next(f for f in facts if f.kind == "change_total" and f.topic == "Выручка")
    last = next(f for f in facts if f.kind == "change_last" and f.topic == "Клиенты")
    answer = {"suggestions": [
        {"fact_ids": [total.id], "headline": "Выручка выросла на 47,5 % за 2026 год", "number": "+47,5 %",
         "visual": "line", "why": "рост на 99 % нужен совету"},
        {"fact_ids": [last.id], "headline": "Клиентов в IV квартале на 22,5 % меньше", "number": "−22,5 %",
         "visual": "column", "why": ""},
        {"fact_ids": [last.id], "headline": "Выручка в IV квартале выросла на 22,5 %", "number": "+22,5 %",
         "visual": "column", "why": ""},
    ]}
    suggestions, _ = suggest(table, facts, "", "", _llm(answer))
    by_model = [s for s in suggestions if s.by_model]
    # «За 2026 год» у сравнения кварталов чинится точным периодом, спад при росте и чужая метрика отбрасываются.
    assert [s.headline for s in by_model] == ["Выручка выросла на 47,5 % с I кв 2026 по IV кв 2026"]
    assert by_model[0].why == ""  # число в «зачем показывать» не из таблицы


def test_newest_first_table_keeps_time_order():
    reversed_quarters = ("Квартал;Выручка\nIV кв 2026;177\nIII кв 2026;150\nII кв 2026;135\nI кв 2026;120\n").encode()
    facts = find_facts(read_table(reversed_quarters, "r.csv"))
    change = next(f for f in facts if f.kind == "change_last")
    assert change.numbers == ["+18 %", "150", "177"]
    assert next(f for f in facts if f.kind == "growth_count").numbers == ["3", "3"]


def test_year_column_is_a_label_not_a_metric():
    years = "Год;Выручка\n2023;90\n2024;120\n2025;177\n".encode()
    facts = find_facts(read_table(years, "y.csv"))
    assert all(f.topic == "Выручка" for f in facts)
    assert any("2024 → 2025" in f.statement for f in facts)


def test_rows_keep_file_line_numbers_with_blank_lines():
    table = read_table("Квартал;Выручка\n\nI кв 2026;120\n\nII кв 2026;135\n".encode(), "b.csv")
    assert table.lines == [3, 5]


def test_xlsx_errors_name_the_sheet():
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = "Обложка"
    book.active.append(["Отчёт за год"])
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(TableError, match="Лист «Обложка»"):
        read_table(buffer.getvalue(), "book.xlsx")


def test_broken_csv_quote_is_explained_in_russian():
    with pytest.raises(TableError, match=r"Строка \d+: кавычка"):
        read_table(b'a;b\n"1;2\n3;4\n', "q.csv")


def test_rejected_key_falls_back_with_reason():
    def handler(request):
        return httpx.Response(401, json={"Code": "UNAUTHORIZED"})
    table = read_table(REGIONS, "r.csv")
    suggestions, note = suggest(table, find_facts(table), "", "", LlmClient("http://test/v1", "m", transport=httpx.MockTransport(handler)))
    assert "ключ" in note and len(suggestions) == 3


def test_numbers_in_ignores_sign_and_groups():
    assert numbers_in("+18,0 % (1 380 → 1690)") == {"18", "1380", "1690"}


def test_insight_slides_go_after_title_and_agenda():
    table = read_table(QUARTERS, "q.csv")
    suggestions, _ = suggest(table, find_facts(table), "", "", None)
    plan = DeckPlan(title="t", purpose="p", slides=[
        SlideIntent(id="", kind=SlideKind.title, title="Итоги"),
        SlideIntent(id="", kind=SlideKind.agenda, title="План"),
        SlideIntent(id="", kind=SlideKind.thanks, title="Спасибо")])
    insert_insight_slides(plan, suggestions)
    assert [s.kind for s in plan.slides[2:2 + len(suggestions)]] != []
    assert plan.slides[2].title == suggestions[0].headline
    assert plan.slides[2].chart is not None and plan.slides[2].chart.categories[-1] == "IV кв 2026"
    assert [s.id for s in plan.slides] == [f"s{i}" for i in range(1, len(plan.slides) + 1)]
    assert "Источник: столбец" in plan.slides[2].notes


def test_plan_repeats_of_insights_are_dropped():
    """Прогон 17.10: план повторил четыре вывода своими слайдами и подписал столбцы «До конца 2025»."""
    from designer.contracts import ChartSpec, Item, Series
    table = read_table(QUARTERS, "q.csv")
    suggestions, _ = suggest(table, find_facts(table), "", "", None)
    revenue = next(s for s in suggestions if s.headline.startswith("Выручка"))
    plan = DeckPlan(title="t", purpose="p", slides=[
        SlideIntent(id="", kind=SlideKind.title, title="Итоги 2026 года"),
        SlideIntent(id="", kind=SlideKind.chart, title="Выручка выросла за год",
                    chart=ChartSpec(type="column", categories=["До конца 2025", "Конец 2026"],
                                    series=[Series(name="Выручка", values=[120, 177])])),
        SlideIntent(id="", kind=SlideKind.section, title=revenue.headline),
        SlideIntent(id="", kind=SlideKind.big_number, title="Рост за 2026 год",
                    items=[Item(heading="рост", number="2026")]),
        SlideIntent(id="", kind=SlideKind.cards, title="Куда вкладываться в 2027 году"),
        SlideIntent(id="", kind=SlideKind.thanks, title="Спасибо")])
    insert_insight_slides(plan, suggestions)
    titles = [s.title for s in plan.slides]
    assert "Выручка выросла за год" not in titles  # числа вывода на чужой диаграмме
    assert titles.count(revenue.headline) == 1  # раздел с тем же заголовком убран
    assert "Рост за 2026 год" in titles  # год не число вывода
    assert "Куда вкладываться в 2027 году" in titles


def test_insight_slide_keeps_model_reading_out_of_slide_text():
    table = read_table(QUARTERS, "q.csv")
    facts = find_facts(table)
    share = next(f for f in facts if f.kind == "change_last" and f.statement.startswith("Доля"))
    answer = {"suggestions": [{"fact_ids": [share.id], "headline": "Доля онлайна в IV квартале выросла до 40,1 %", "number": "40,1 %",
                               "visual": "big_number", "why": "свидетельствует о росте лояльности"}]}
    suggestions, _ = suggest(table, facts, "", "", _llm(answer))
    plan = DeckPlan(title="t", purpose="p", slides=[SlideIntent(id="", kind=SlideKind.title, title="t")])
    insert_insight_slides(plan, suggestions[:1])
    made = plan.slides[1]
    text = made.title + made.key_message + " ".join(i.heading for i in made.items)
    # Причинное толкование модели не идёт ни на слайд, ни в заметки: причин в таблице нет.
    assert "свидетельствует" not in text and "свидетельствует" not in made.notes
    assert made.items[0].heading == share.statement


def test_api_insights_then_deck_request_finds_them(tmp_path, monkeypatch):
    monkeypatch.setenv(store.DATA_DIR_ENV, str(tmp_path / "data"))
    app.dependency_overrides[get_llm_client] = lambda: _llm({"suggestions": []})
    try:
        with TestClient(app) as client:
            response = client.post("/insights", files={"file": ("r.csv", REGIONS, "text/csv")},
                                   data={"brief": "итоги", "audience": "директора"})
            assert response.status_code == 200, response.text
            body = response.json()
            assert body["rows"] == 6 and body["columns"] == 3 and len(body["suggestions"]) == 3
            assert body["suggestions"][0]["source"].startswith("столб")
            saved = store.load_insights(body["id"])
            assert saved["suggestions"][0]["chart"]["categories"]

            bad = client.post("/insights", files={"file": ("x.csv", b"a,b\n1,2\n", "text/csv")})
            assert bad.status_code == 422 and "две строки" in bad.json()["detail"]

            missing = client.post("/decks", json={"design_system_id": "x", "brief": "b",
                                                  "insights_id": "nope", "insight_ids": ["s1"]})
            assert missing.status_code == 404
    finally:
        app.dependency_overrides.clear()
