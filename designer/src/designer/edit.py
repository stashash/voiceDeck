"""Правка варианта колоды: текст на слайде, образец, лента слайдов, просьба агенту, история.

Каждая правка снимает прежний deck.json варианта в history/<n>.json (store.snapshot_deck_history),
пересобирает затронутые слайды, гоняет детерминированный аудит по всем сценам и переиздаёт
файлы варианта (pptx, html, pdf, картинки слайдов) — так же, как это делает
pipeline._build_and_export_variant после генерации. Контекстный (по картинке) аудит здесь
не пересчитывается: он дорогой, его запускает POST /decks/{id}/{v}/audit-contextual по запросу,
как и раньше; уже найденные контекстные замечания, кроме починенного, переносятся как есть.

Клиент модели для «попросить агента» и «переписать замечание» берётся через
agent_hook._client_for("deck") — тот же приём, что и в pipeline.py.
"""
from __future__ import annotations

import math
import re
import uuid
from pathlib import Path

from designer import store
from designer.edit_runtime import serialized, schedule, lock
from designer.agent_hook import _client_for
from designer.audit.deterministic import run_checks
from designer.contracts import (
    Deck,
    DeckPlan,
    DesignSystem,
    Element,
    ElementPosition,
    Finding,
    Item,
    Pattern,
    Scene,
    SlideIntent,
    SlideKind,
    SlideSpec,
    TableSpec,
    TextStyle,
)
from designer.export import convert, render
from designer.export.html import render_deck
from designer.export.html_image import render_deck_images
from designer.export.pptx_deck import export_pptx
from designer.layout.capacity import main_group, unit_count
from designer.layout.compose import compose
from designer.layout.match import choose_pattern, fits, score_pattern
from designer.layout.scene import build_scene
from designer.parse.package import load_package
from designer.plan.writer import fill_slots

_MAX_PATTERN_OPTIONS = 6
_UNIT_ID_RE = re.compile(r"^u(\d+)s(\d+)$")
_SLIDE_ACTIONS = ("add", "copy", "delete", "move")


class DeckNotFound(ValueError):
    """Колода или вариант не найдены."""


class SlideNotFound(ValueError):
    """Номер слайда (или позиция to) вне диапазона плана варианта."""


class ElementNotFound(ValueError):
    """id элемента нет на сцене слайда."""


class PatternNotFound(ValueError):
    """Образец с этим id не найден в дизайн-системе."""


class FindingNotFound(ValueError):
    """Замечание с этим id не найдено среди findings варианта."""


class NoHistory(ValueError):
    """Истории правок нет: откатывать нечего."""


# ---------- состояние варианта ----------

class _State:
    """Разобранное deck.json варианта; raw — исходный словарь для снимка истории."""

    def __init__(self, deck_id: str, variant: str, raw: dict, ds_id: str, package_dir: Path, ds: DesignSystem):
        self.deck_id = deck_id
        self.variant = variant
        self.raw = raw
        self.ds_id = ds_id
        self.package_dir = package_dir
        self.ds = ds
        self.plan = DeckPlan.model_validate(raw["plan"])
        self.specs: list[SlideSpec] = [SlideSpec.model_validate(item) for item in raw["specs"]]
        self.scenes: list[Scene] = [Scene.model_validate(item) for item in raw["scenes"]]
        self.findings: list[Finding] = [Finding.model_validate(item) for item in raw["findings"]]


def _load(deck_id: str, variant: str) -> _State:
    raw = store.load_deck_state(deck_id, variant)
    if raw is None:
        raise DeckNotFound(f"колода не найдена: {deck_id}/{variant}")
    ds_id = raw["design_system_id"]
    package_dir = store.design_system_dir(ds_id)
    ds = load_package(package_dir)
    return _State(deck_id, variant, raw, ds_id, package_dir, ds)


def _slide_index(state: _State, slide_number: int) -> int:
    index = slide_number - 1
    if index < 0 or index >= len(state.plan.slides):
        raise SlideNotFound(f"слайда {slide_number} нет: в варианте {len(state.plan.slides)} слайдов")
    return index


def _snapshot(state: _State) -> None:
    store.snapshot_deck_history(state.deck_id, state.variant, state.raw)


def _checks(state: _State, touched: set[str] | None = None, drop_ids: set[str] | None = None) -> list[Finding]:
    """Проверка по правилам заново плюс прежние замечания модели по слайдам, которых правка не касалась.

    Замечание модели считалось по картинке слайда: у изменённого слайда оно устарело, у остальных нет.
    """
    present = {spec.slide_id for spec in state.specs}
    touched = touched or set()
    drop_ids = drop_ids or set()
    kept = [f for f in state.findings if f.kind == "contextual" and f.slide_id in present
            and f.slide_id not in touched and f.id not in drop_ids]
    return run_checks(state.scenes, state.ds) + kept


def _persist(state: _State, findings: list[Finding]) -> None:
    """Commit the document immediately; derived files are exported separately."""
    deck = Deck(id=state.deck_id, design_system_id=state.ds_id, variant=state.variant,
                plan=state.plan, specs=state.specs, scenes=state.scenes)
    store.save_deck_result(state.deck_id, deck, findings, variant=state.variant)
    raw = store.load_deck_state(state.deck_id, state.variant)
    raw["revision"] = uuid.uuid4().hex
    store._write_json(store.deck_variant_state_path(state.deck_id, state.variant), raw)
    schedule(state.deck_id, state.variant, raw["revision"])


def live_slide(deck_id: str, variant: str, number: int) -> str:
    state = _load(deck_id, variant)
    index = _slide_index(state, number)
    scene = state.scenes[index].model_copy(deep=True)
    if scene.theme == 'dark':
        for element in scene.elements:
            if element.type == 'text' and (element.style is None or element.style.color is None):
                element.style = (element.style or TextStyle()).model_copy(update={'color': 'FFFFFF'})
    deck = Deck(id=deck_id, design_system_id=state.ds_id, variant=variant,
                plan=state.plan, specs=[state.specs[index]], scenes=[scene])
    return render_deck(deck, state.ds, state.package_dir).split("<script>")[0] + "</body></html>"


def _export_files(state: _State, files_dir: Path) -> None:
    deck = Deck(id=state.deck_id, design_system_id=state.ds_id, variant=state.variant,
                plan=state.plan, specs=state.specs, scenes=state.scenes)
    pptx_path = files_dir / "deck.pptx"
    # Native tables can have content-fitted heights, so their editor geometry is
    # applied after export instead of the template-shape proximity matcher.
    export_specs = [spec.model_copy(deep=True) for spec in state.specs]
    for spec in export_specs:
        if spec.table is not None:
            spec.element_positions.pop(spec.viz_area_id or "viz", None)
    export_pptx(export_specs, state.ds, state.package_dir, pptx_path)
    _export_element_overrides(state, pptx_path)
    markup_html = render_deck(deck, state.ds, state.package_dir)
    (files_dir / "deck.markup.html").write_text(markup_html, encoding="utf-8")
    (files_dir / "deck.html").write_text(markup_html, encoding="utf-8")

    if convert.available():
        try:
            convert.to_pdf(pptx_path, files_dir / "deck.pdf")
        except convert.ConverterUnavailable:
            pass
        try:
            pngs = render.render_slides(pptx_path, len(state.specs)) if state.specs else []
        except convert.ConverterUnavailable:
            pngs = []
        if pngs:
            slides_dir = files_dir / "slides"
            slides_dir.mkdir(exist_ok=True)
            for index, png in enumerate(pngs, start=1):
                (slides_dir / f"slide-{index:03d}.png").write_bytes(png)
            html_text = render_deck_images(deck, state.ds, pngs)
            (files_dir / "deck.html").write_text(html_text, encoding="utf-8")


def _export_element_overrides(state: _State, pptx_path: Path) -> None:
    """Complete editor-only properties without changing the generation exporter."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN
    from pptx.shapes.shapetree import BaseShapeFactory

    from designer.export.pptx_deck import _prune_groups, _set_box, _walk
    from designer.viz.pptx_native import add_table

    def table_box(shape, box):
        sw, sh = state.ds.slide_size_emu
        shape.left, shape.top, shape.width = round(box[0] * sw), round(box[1] * sh), round(box[2] * sw)
        height = round(box[3] * sh)
        rows = list(shape.table.rows)
        for i, row in enumerate(rows):
            row.height = height // len(rows) + (1 if i < height % len(rows) else 0)

    prs = Presentation(str(pptx_path))
    for slide, spec, scene in zip(prs.slides, state.specs, state.scenes):
        tree = slide.shapes._spTree
        layers = []
        native_added = [el for el in spec.added_elements if el.type in ("text", "shape")]
        added_shapes = list(slide.shapes)[-len(native_added):] if native_added else []
        added_ids = {shape.shape_id for shape in added_shapes}
        placed = _walk(slide.shapes, state.ds.slide_size_emu)
        used = set(added_ids)

        def text_style(shape, style):
            if style is None or not shape.has_text_frame:
                return
            for paragraph in shape.text_frame.paragraphs:
                if style.align is not None:
                    paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                                           "right": PP_ALIGN.RIGHT}[style.align]
                for run in paragraph.runs:
                    run.font.italic = style.italic

        for element_id, position in spec.element_positions.items():
            is_table = spec.table is not None and element_id == (spec.viz_area_id or "viz")
            if is_table:
                tables = [shape for shape in slide.shapes if shape.has_table and shape.shape_id not in used]
                if not tables:
                    raise ValueError("Edited table was not found in PowerPoint")
                shape = tables[-1]
                item = placed[shape.shape_id]
                if position.deleted:
                    used.add(shape.shape_id)
                    tree.remove(shape._element)
                    continue
                table_box(shape, position.box)
            else:
                if position.deleted or (position.style is None and position.fill is None and position.z is None):
                    continue
                candidates = [(sid, item) for sid, item in placed.items()
                              if sid not in used and item.element.getparent() is not None]
                candidates.sort(key=lambda pair: (
                    sum(abs(a - b) for a, b in zip(pair[1].box, position.box)),
                    pair[0] != position.source_shape_id))
                if not candidates or sum(abs(a - b) for a, b in zip(candidates[0][1].box, position.box)) > .04:
                    raise ValueError("Edited element was not found in PowerPoint")
                sid, item = candidates[0]
                shape = BaseShapeFactory(item.element, slide.shapes)
            used.add(shape.shape_id)
            text_style(shape, position.style)
            if position.fill is not None:
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(position.fill)
            if position.z is not None:
                # A selected child may cross group boundaries when brought forward.
                if item.element.getparent() is not tree:
                    item.element.getparent().remove(item.element)
                    tree.insert_element_before(item.element, "p:extLst")
                    item.transform = (1.0, 1.0, 0.0, 0.0)
                    _set_box(item, position.box, state.ds.slide_size_emu)
                layers.append((position.z, item.element))

        for element, shape in zip(native_added, added_shapes):
            text_style(shape, element.style)
            layers.append((element.z, shape._element))
        for element in spec.added_elements:
            if element.type == "table":
                shape = add_table(slide, element.table, element.box, state.ds.slide_size_emu,
                                  state.ds.tokens, scene.theme)
                table_box(shape, element.box)
                layers.append((element.z, shape._element))

        background = None
        if spec.background_color is not None:
            pattern = next((p for p in state.ds.patterns if p.id == spec.pattern_id), None)
            for item in _walk(slide.shapes, state.ds.slide_size_emu).values():
                shape = BaseShapeFactory(item.element, slide.shapes)
                if (pattern and pattern.background_asset and shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                        and shape.image.sha1.startswith(pattern.background_asset)
                        and item.box[2] >= .95 and item.box[3] >= .95):
                    item.element.getparent().remove(item.element)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(spec.background_color)
            # This also covers inherited master/layout pictures behind slide content.
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
            shape.name = "VoiceDeck background"
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(spec.background_color)
            shape.line.fill.background()
            background = shape._element
            tree.remove(background)
            tree.insert(2, background)

        for _, node in sorted((pair for pair in layers if pair[0] < 0), key=lambda pair: pair[0], reverse=True):
            tree.remove(node)
            tree.insert(3 if background is not None else 2, node)
        for _, node in sorted((pair for pair in layers if pair[0] >= 0), key=lambda pair: pair[0]):
            tree.remove(node)
            tree.insert_element_before(node, "p:extLst")
        _prune_groups(slide)
    prs.save(str(pptx_path))



def _recompose(state: _State, index: int, intent: SlideIntent, pattern: Pattern) -> None:
    spec = compose(intent, pattern, state.ds)
    scene = build_scene(spec, pattern, state.ds, state.package_dir)
    state.plan.slides[index] = intent
    state.specs[index] = spec
    state.scenes[index] = scene


# ---------- текст элемента на слайде ----------

def _slot_ref(el: Element, pattern: Pattern | None):
    """Слот паттерна для элемента сцены: верхнего уровня либо слот повторяющегося блока.

    Дублирует небольшую часть audit.fixes._slot_ref: тот модуль потоку 1 не принадлежит,
    трогать его нельзя, а связь «элемент сцены -> ключ SlideSpec» нужна и здесь.
    """
    if pattern is None:
        return None, None, None
    for slot in pattern.slots:
        if slot.id == el.id:
            return slot, None, None
    match = _UNIT_ID_RE.match(el.id)
    if not match or el.source_shape_id is None:
        return None, None, None
    index = int(match.group(1))
    for group in pattern.groups:
        for slot in group.unit_slots:
            if slot.shape_id == el.source_shape_id:
                return slot, group, index
    return None, None, None


def _set_slot_text(spec: SlideSpec, slot, group, index: int | None, text: str) -> None:
    if group is None:
        spec.slot_text[slot.id] = text
        return
    target = spec.linked_unit_text[group.id] if group.id in spec.linked_unit_text else spec.unit_text
    if index is not None and 0 <= index < len(target):
        target[index][slot.id] = text


def _set_intent_text(intent: SlideIntent, slot, group, unit_index: int | None, text: str) -> None:
    """Правка автора пишется и в намерение слайда: смена образца и просьба агенту пересобирают
    слайд из намерения, и без этого ручная правка терялась (проверено 2026-09-24 в интерфейсе)."""
    if group is None:
        if slot.role == "title":
            intent.title = text
        elif slot.role in ("subtitle", "body"):
            intent.key_message = text
        return
    if unit_index is None or not 0 <= unit_index < len(intent.items):
        return
    item = intent.items[unit_index]
    if slot.role in ("heading", "title", "label"):
        item.heading = text
    elif slot.role == "number":
        item.number = text
    else:
        item.body = text


@serialized
def move_element(deck_id: str, variant: str, slide_number: int, element_id: str,
                 dx: float = 0, dy: float = 0, align: str | None = None) -> None:
    state = _load(deck_id, variant)
    index = _slide_index(state, slide_number)
    el = next((e for e in state.scenes[index].elements if e.id == element_id), None)
    if el is None:
        raise ElementNotFound("Выбранный элемент не найден на слайде")
    x, y, w, h = el.box
    x, y = x + dx, y + dy
    if align == "left": x = 0.02
    if align == "right": x = 0.98 - w
    if align == "top": y = 0.02
    if align == "bottom": y = 0.98 - h
    if align == "center": x, y = (1-w)/2, (1-h)/2
    box = (max(0, min(1-w, x)), max(0, min(1-h, y)), w, h)
    _snapshot(state)
    previous = state.specs[index].element_positions.get(el.id)
    added = next((e for e in state.specs[index].added_elements if e.id == el.id), None)
    if added is not None:
        added.box = box
    else:
        state.specs[index].element_positions[el.id] = (previous.model_copy(update={"box": box}) if previous else
            ElementPosition(original_box=el.box, box=box, source_shape_id=el.source_shape_id))
    el.box = box
    _persist(state, _checks(state, {state.specs[index].slide_id}))


@serialized
def set_slide_text(deck_id: str, variant: str, slide_number: int, element_id: str, text: str) -> None:
    state = _load(deck_id, variant)
    index = _slide_index(state, slide_number)
    scene = state.scenes[index]
    el = next((e for e in scene.elements if e.id == element_id), None)
    if el is None:
        raise ElementNotFound(f"элемент {element_id} не найден на слайде {slide_number}")
    if el.type != "text":
        raise ValueError("Выбранный объект не является текстом")

    _snapshot(state)
    el.text = text
    pattern = next((p for p in state.ds.patterns if p.id == scene.pattern_id), None)
    slot, group, unit_index = _slot_ref(el, pattern)
    if slot is not None:
        _set_slot_text(state.specs[index], slot, group, unit_index, text)
        _set_intent_text(state.plan.slides[index], slot, group, unit_index, text)
    added = next((e for e in state.specs[index].added_elements if e.id == el.id), None)
    if added is not None:
        added.text = text
    elif slot is None:
        override = state.specs[index].element_positions.setdefault(el.id,
            ElementPosition(original_box=el.box, box=el.box, source_shape_id=el.source_shape_id))
        override.text = text

    findings = _checks(state, {state.specs[index].slide_id})
    _persist(state, findings)


# ---------- образец слайда ----------

def _element_box(box: tuple[float, float, float, float]) -> tuple[float, float, float, float]:
    x, y, w, h = box
    if (not all(math.isfinite(value) for value in box) or min(x, y) < 0 or min(w, h) <= 0
            or x + w > 1 + 1e-9 or y + h > 1 + 1e-9):
        raise ValueError("Element dimensions must be positive and fit inside the slide")
    return box


def _element_override(spec: SlideSpec, el: Element) -> ElementPosition:
    return spec.element_positions.setdefault(el.id, ElementPosition(
        original_box=el.box, box=el.box, source_shape_id=el.source_shape_id))


def _save_element(scene: Scene, spec: SlideSpec, old: Element, new: Element) -> None:
    for i, added in enumerate(spec.added_elements):
        if added.id == old.id:
            spec.added_elements[i] = new.model_copy(deep=True)
            break
    else:
        if any(getattr(old, field) != getattr(new, field) for field in ("box", "style", "fill", "z")):
            override = _element_override(spec, old)
            override.box = new.box
            if new.style != old.style:
                override.style = new.style.model_copy(deep=True) if new.style else None
            if new.fill != old.fill:
                override.fill = new.fill
            if new.z != old.z:
                override.z = new.z
    scene.elements = [new if item.id == old.id else item for item in scene.elements]


def _edited_table(table: TableSpec, request) -> TableSpec:
    from designer.api.schemas import ElementTableRequest

    result = ElementTableRequest.model_validate(table.model_dump()).model_copy(deep=True)
    row, column = request.row, request.column
    if request.action == "table_cell":
        if row > len(result.rows) or column > len(result.columns):
            raise ValueError("Table cell is outside the table")
        cells = result.columns if row == 0 else result.rows[row - 1]
        cells[column - 1] = request.text
    elif request.action == "table_row_add":
        row = row if row is not None else len(result.rows) + 1
        values = request.values if request.values is not None else [""] * len(result.columns)
        if row > len(result.rows) + 1 or len(values) != len(result.columns):
            raise ValueError("Invalid row insertion index or number of cells")
        result.rows.insert(row - 1, list(values))
    elif request.action == "table_row_delete":
        if row > len(result.rows):
            raise ValueError("Table row is outside the table")
        del result.rows[row - 1]
    elif request.action == "table_column_add":
        column = column if column is not None else len(result.columns) + 1
        values = request.values if request.values is not None else [""] * len(result.rows)
        if column > len(result.columns) + 1 or len(values) != len(result.rows):
            raise ValueError("Invalid column insertion index or number of cells")
        result.columns.insert(column - 1, request.text)
        for cells, value in zip(result.rows, values):
            cells.insert(column - 1, value)
    elif request.action == "table_column_delete":
        if column > len(result.columns) or len(result.columns) == 1:
            raise ValueError("Column must exist and the table must retain at least one column")
        del result.columns[column - 1]
        for cells in result.rows:
            del cells[column - 1]
    return TableSpec.model_validate(ElementTableRequest.model_validate(result.model_dump()).model_dump())


@serialized
def element_action(deck_id: str, variant: str, number: int, action: str,
                   element_id: str | None = None, **properties) -> None:
    """Validate a deterministic edit completely before writing history or state."""
    from designer.api.schemas import ElementActionRequest, ElementTableRequest
    from designer.viz.palette import contrast_text_color

    target = {"element_id": element_id} if element_id is not None else {}
    request = ElementActionRequest(action=action, **target, **properties)
    supplied = request.model_fields_set
    state = _load(deck_id, variant)
    index = _slide_index(state, number)
    scene, spec = state.scenes[index], state.specs[index]
    el = next((e for e in scene.elements if e.id == element_id), None)
    if action not in ("add", "background") and el is None:
        raise ElementNotFound("Selected element was not found on this slide")
    if el is not None and el.source_shape_id is not None and action in ("style", "delete", "z_order"):
        if any(other.id != el.id and other.source_shape_id == el.source_shape_id
               and (other.id == el.id + "c" or el.id == other.id + "c") for other in scene.elements):
            raise ValueError("This number and caption share one native shape; edit a separate text element")

    if action in ("add", "duplicate"):
        if action == "duplicate":
            if el.type not in ("text", "shape", "table") or (el.type == "table" and el.table is None):
                raise ValueError("Duplicating this element type is not supported")
            new = el.model_copy(deep=True)
            if new.table is not None:
                new.table = ElementTableRequest.model_validate(new.table.model_dump())
            x, y, w, h = _element_box(new.box)
            new.box = _element_box((min(1 - w, x + .02), min(1 - h, y + .02), w, h))
            additions = [new]
        else:
            kind = request.element_type
            default_height = .4 if kind == "table" else .3 if kind == "card" else .15
            box = _element_box((.1, .15, request.width or .5, request.height or default_height))
            color = next((e.style.color for e in scene.elements if e.type == "text" and e.style and e.style.color),
                         next((c.hex for c in state.ds.tokens.colors if c.role == "text"),
                              "FFFFFF" if scene.theme == "dark" else "222222"))
            style = TextStyle(size_pt=32 if kind == "title" else 24,
                              family=next((f.family for f in state.ds.tokens.fonts), None), color=color)
            if kind == "table":
                table = request.table or TableSpec(columns=["Column 1", "Column 2"], rows=[["", ""], ["", ""]])
                additions = [Element(id="new", type="table", role="viz", box=box, table=table)]
            elif kind in ("shape", "card"):
                additions = [Element(id="new", type="shape", role="card" if kind == "card" else "other",
                                     box=box, fill=(request.fill or "D9D9D9").upper())]
                if kind == "card":
                    x, y, w, h = box
                    padding = min(w, h) * .08
                    additions.append(Element(id="new", type="text", role="body", text=request.text,
                                             box=(x + padding, y + padding, w - 2 * padding, h - 2 * padding),
                                             style=style.model_copy(update={
                                                 "color": contrast_text_color(additions[0].fill)})))
            else:
                additions = [Element(id="new", type="text", role="title" if kind == "title" else "body",
                                     box=box, text=request.text, style=style)]
        z = max((e.z for e in scene.elements), default=0)
        for offset, new in enumerate(additions, start=1):
            new.id = "custom-" + uuid.uuid4().hex
            new.source_shape_id = None
            new.z = z + offset
            spec.added_elements.append(new.model_copy(deep=True))
            scene.elements.append(new)
    elif action == "delete":
        if any(e.id == el.id for e in spec.added_elements):
            spec.added_elements = [e for e in spec.added_elements if e.id != el.id]
        else:
            _element_override(spec, el).deleted = True
        scene.elements = [e for e in scene.elements if e.id != el.id]
    elif action == "style":
        text_fields = supplied & {"size_pt", "scale", "color", "bold", "italic", "text_align"}
        if text_fields and el.type != "text":
            raise ValueError("Font properties are supported only for text elements")
        if request.fill is not None and el.type != "shape":
            raise ValueError("fill is supported only for shape elements")
        new = el.model_copy(deep=True)
        if request.width is not None or request.height is not None:
            x, y, w, h = new.box
            new.box = _element_box((x, y, request.width or w, request.height or h))
        if text_fields:
            style = (el.style or TextStyle()).model_copy(deep=True)
            if "size_pt" in supplied or "scale" in supplied:
                size = request.size_pt if request.size_pt is not None else (style.size_pt or 24) * request.scale
                if not math.isfinite(size) or not 6 <= size <= 144:
                    raise ValueError("Resulting font size must be between 6 and 144 pt")
                style.size_pt = size
            for field in ("bold", "italic"):
                if field in supplied:
                    setattr(style, field, getattr(request, field))
            if request.color is not None:
                style.color = request.color.upper()
            if request.text_align is not None:
                style.align = request.text_align
            new.style = style
        if request.fill is not None:
            new.fill = request.fill.upper()
        _save_element(scene, spec, el, new)
    elif action == "background":
        spec.background_color = request.color.upper()
        scene.background_color = spec.background_color
        scene.background_asset = None
    elif action == "z_order":
        new = el.model_copy(deep=True)
        levels = [0, *(item.z for item in scene.elements)]
        new.z = max(levels) + 1 if request.order == "front" else min(levels) - 1
        _save_element(scene, spec, el, new)
    elif action.startswith("table_"):
        if el.type != "table" or el.table is None:
            raise ValueError("Select an editable table")
        added = any(item.id == el.id for item in spec.added_elements)
        if not added and (spec.table is None or el.id != (spec.viz_area_id or "viz")):
            raise ValueError("This template table is not editable")
        new = el.model_copy(deep=True)
        new.table = _edited_table(el.table, request)
        if not added:
            spec.table = new.table.model_copy(deep=True)
            state.plan.slides[index].table = new.table.model_copy(deep=True)
        _save_element(scene, spec, el, new)

    findings = _checks(state, {spec.slide_id})
    _snapshot(state)
    _persist(state, findings)


@serialized
def set_slide_pattern(deck_id: str, variant: str, slide_number: int, pattern_id: str) -> None:
    state = _load(deck_id, variant)
    index = _slide_index(state, slide_number)
    pattern = next((p for p in state.ds.patterns if p.id == pattern_id), None)
    if pattern is None:
        raise PatternNotFound(f"образец {pattern_id} не найден в дизайн-системе")

    _snapshot(state)
    _recompose(state, index, state.plan.slides[index], pattern)

    findings = _checks(state, {state.specs[index].slide_id})
    _persist(state, findings)


def list_slide_patterns(deck_id: str, variant: str, slide_number: int) -> list[dict]:
    """До шести образцов, которые годятся под намерение слайда; текущий — первым."""
    state = _load(deck_id, variant)
    index = _slide_index(state, slide_number)
    intent = state.plan.slides[index]
    current_id = state.specs[index].pattern_id
    used = [spec.pattern_id for i, spec in enumerate(state.specs) if i != index]

    candidates = [p for p in state.ds.patterns if p.id == current_id or fits(intent, p, state.ds)]
    ranked = sorted(candidates, key=lambda p: (p.id != current_id, -score_pattern(intent, p, used, state.ds), p.id))
    chosen = ranked[:_MAX_PATTERN_OPTIONS]

    return [{
        "pattern_id": p.id,
        "kind": p.kind.value,
        "preview": f"/design-systems/{state.ds_id}/previews/{p.id}.png" if p.preview else None,
        "current": p.id == current_id,
    } for p in chosen]


# ---------- агент переписывает содержание слайда ----------

def _draft_schema() -> dict:
    item_schema = {
        "type": "object",
        "properties": {
            "heading": {"type": "string", "maxLength": 200},
            "body": {"type": "string", "maxLength": 400},
            "number": {"type": "string", "maxLength": 20},
        },
        "required": ["heading", "body"],
    }
    return {
        "type": "object",
        "properties": {
            "title": {"type": "string", "maxLength": 200},
            "key_message": {"type": "string", "maxLength": 400},
            "items": {"type": "array", "items": item_schema},
            "notes": {"type": "string", "maxLength": 600},
        },
        "required": ["title", "key_message", "items", "notes"],
    }


def _draft_system(instruction: str) -> str:
    return (
        "Ты редактируешь слайд презентации по просьбе автора.\n"
        "Перепиши заголовок, ключевое сообщение, пункты и заметки докладчика по просьбе ниже; "
        "то, чего просьба не касается, оставь как есть по смыслу. Числа и факты бери только "
        "из текущего содержания слайда, не выдумывай новые. Ответ на языке текущего содержания.\n\n"
        f"Просьба автора: {instruction}\n\n"
        "Ответ — только JSON по схеме, без пояснений и без рассуждений вслух."
    )


def _intent_material(intent: SlideIntent) -> str:
    lines = [f"Заголовок: {intent.title}"]
    if intent.key_message:
        lines.append(f"Ключевое сообщение: {intent.key_message}")
    for i, item in enumerate(intent.items, start=1):
        piece = item.heading
        if item.body:
            piece = f"{piece} — {item.body}" if piece else item.body
        if item.number:
            piece = f"{item.number} {piece}".strip()
        lines.append(f"Пункт {i}: {piece}")
    if intent.notes:
        lines.append(f"Заметки докладчика: {intent.notes}")
    return "\n".join(lines)


def _draft_intent(intent: SlideIntent, data: dict) -> SlideIntent:
    items = [
        Item(heading=str(raw.get("heading", "")), body=str(raw.get("body", "")),
             number=(str(raw["number"]) if raw.get("number") else None))
        for raw in data.get("items", []) or []
    ]
    return intent.model_copy(update={
        "title": str(data.get("title", intent.title)),
        "key_message": str(data.get("key_message", intent.key_message)),
        "items": items,
        "notes": str(data.get("notes", intent.notes)),
    })


def _fit_to_pattern(intent: SlideIntent, pattern: Pattern, ds: DesignSystem, client) -> SlideIntent:
    from designer import pipeline  # локальный импорт: те же лимиты слотов, что при генерации

    group = main_group(pattern)
    n_units = unit_count(group, len(intent.items)) if group is not None else 0
    limits, unit_limits = pipeline._role_limits(pattern, intent, n_units, ds.tokens.type_scale)
    if not limits and not unit_limits:
        return intent
    return fill_slots(intent, limits, unit_limits, n_units, client)


def _rewrite_intent(intent: SlideIntent, pattern: Pattern, instruction: str, ds: DesignSystem) -> SlideIntent:
    client = _client_for("deck")
    try:
        data = client.complete_json(system=_draft_system(instruction), user=_intent_material(intent),
                                     schema=_draft_schema())
        draft = _draft_intent(intent, data)
        return _fit_to_pattern(draft, pattern, ds, client)
    finally:
        client.close()


def ask_agent_rewrite(deck_id: str, variant: str, slide_number: int, instruction: str) -> None:
    state = _load(deck_id, variant)
    index = _slide_index(state, slide_number)
    intent = state.plan.slides[index]
    pattern = next((p for p in state.ds.patterns if p.id == state.specs[index].pattern_id), None)
    if pattern is None:
        pattern = choose_pattern(intent, state.ds, [s.pattern_id for s in state.specs])

    new_intent = _rewrite_intent(intent, pattern, instruction, state.ds)
    # Model latency must not lock the document. Recovery or a newer edit wins.
    with lock(deck_id, variant):
        if store.load_deck_state(deck_id, variant) != state.raw:
            raise ValueError("Документ изменился или команда остановлена. Ответ агента не применён.")
        _snapshot(state)
        _recompose(state, index, new_intent, pattern)
        findings = _checks(state, {state.specs[index].slide_id})
        _persist(state, findings)


@serialized
def recover_editor(deck_id: str, variant: str) -> None:
    state = _load(deck_id, variant)
    state.raw["recovery_token"] = uuid.uuid4().hex
    store._write_json(store.deck_variant_state_path(deck_id, variant), state.raw)


def rewrite_from_finding(deck_id: str, variant: str, finding_id: str) -> None:
    """Замечание модели: слайд переписывается с текстом замечания как просьбой автора."""
    state = _load(deck_id, variant)
    finding = next((f for f in state.findings if f.id == finding_id), None)
    if finding is None:
        raise FindingNotFound(f"замечание {finding_id} не найдено")
    index = next((i for i, spec in enumerate(state.specs) if spec.slide_id == finding.slide_id), None)
    if index is None:
        raise SlideNotFound(f"слайд замечания {finding_id} не найден")

    intent = state.plan.slides[index]
    pattern = next((p for p in state.ds.patterns if p.id == state.specs[index].pattern_id), None)
    if pattern is None:
        raise PatternNotFound("образец слайда не найден в дизайн-системе")

    ask_agent_rewrite(deck_id, variant, index + 1, finding.message)


# ---------- заметки докладчика ----------

@serialized
def set_slide_notes(deck_id: str, variant: str, slide_number: int, notes: str) -> None:
    state = _load(deck_id, variant)
    index = _slide_index(state, slide_number)

    _snapshot(state)
    state.plan.slides[index].notes = notes
    state.specs[index].notes = notes

    findings = _checks(state)
    _persist(state, findings)


# ---------- лента слайдов: добавить, копия, удалить, переставить ----------

@serialized
def apply_slide_action(deck_id: str, variant: str, action: str, index: int, to: int | None = None) -> None:
    if action not in _SLIDE_ACTIONS:
        raise ValueError(f"неизвестное действие: {action}")
    state = _load(deck_id, variant)
    n = len(state.plan.slides)

    if action == "delete":
        if n <= 1:
            raise ValueError("нельзя удалить последний слайд варианта")
        pos = _slide_index(state, index)
        _snapshot(state)
        del state.plan.slides[pos]
        del state.specs[pos]
        del state.scenes[pos]

    elif action == "move":
        if to is None:
            raise ValueError("move: нужен параметр to")
        pos = _slide_index(state, index)
        target = to - 1
        if target < 0 or target >= n:
            raise SlideNotFound(f"позиции {to} нет: в варианте {n} слайдов")
        _snapshot(state)
        intent, spec, scene = state.plan.slides.pop(pos), state.specs.pop(pos), state.scenes.pop(pos)
        state.plan.slides.insert(target, intent)
        state.specs.insert(target, spec)
        state.scenes.insert(target, scene)

    else:  # add | copy
        if action == "add" and index == 0:
            pos = -1
        else:
            pos = _slide_index(state, index)
        _snapshot(state)
        used = [s.pattern_id for s in state.specs]
        if action == "copy":
            source = state.plan.slides[pos]
            new_intent = source.model_copy(update={"id": uuid.uuid4().hex})
            pattern = (next((p for p in state.ds.patterns if p.id == state.specs[pos].pattern_id), None)
                       or choose_pattern(new_intent, state.ds, used))
        else:
            new_intent = SlideIntent(id=uuid.uuid4().hex, kind=SlideKind.bullets, title="Новый слайд")
            pattern = choose_pattern(new_intent, state.ds, used)
        if action == "copy":
            spec = state.specs[pos].model_copy(deep=True, update={"slide_id": new_intent.id})
            scene = state.scenes[pos].model_copy(deep=True, update={"slide_id": new_intent.id})
        else:
            spec = compose(new_intent, pattern, state.ds)
            scene = build_scene(spec, pattern, state.ds, state.package_dir)
        insert_at = pos + 1
        state.plan.slides.insert(insert_at, new_intent)
        state.specs.insert(insert_at, spec)
        state.scenes.insert(insert_at, scene)

    findings = _checks(state)
    _persist(state, findings)


# ---------- вернуть последнюю правку ----------

@serialized
def revert(deck_id: str, variant: str) -> None:
    history_dir = store.deck_variant_history_dir(deck_id, variant)
    numbers = sorted((int(p.stem) for p in history_dir.glob("*.json") if p.stem.isdigit()), reverse=True)
    if not numbers:
        raise NoHistory("истории правок нет")
    path = history_dir / f"{numbers[0]}.json"
    raw = store.read_json_file(path)

    ds_id = raw["design_system_id"]
    package_dir = store.design_system_dir(ds_id)
    ds = load_package(package_dir)
    state = _State(deck_id, variant, raw, ds_id, package_dir, ds)

    _persist(state, state.findings)
    path.unlink()
