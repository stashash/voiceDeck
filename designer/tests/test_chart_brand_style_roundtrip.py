from pptx import Presentation
from pptx.enum.chart import XL_LEGEND_POSITION
from designer.contracts import ChartSpec, ChartStyle, ColorToken, FontToken, Margins, Series, Tokens
from designer.viz.pptx_native import add_chart


def test_single_line_and_pie_use_explicit_brand_style_in_native_file(tmp_path):
    tokens = Tokens(colors=[ColorToken(hex="713CDA",role="accent",share=1,source="theme"),
                            ColorToken(hex="162032",role="text",share=1,source="theme")],
                    fonts=[FontToken(family="Play",role="body",share=1)],
                    chart_style=ChartStyle(legend_position="bottom",label_size_pt=12),type_scale=[],
                    margins=Margins(left=.06,right=.06,top=.06,bottom=.06))
    prs=Presentation()
    for kind in ["line","pie"]:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        add_chart(slide,ChartSpec(type=kind,categories=["A","B"],series=[Series(name="Metric",values=[1,2])]),
                  (.1,.2,.8,.6),(prs.slide_width,prs.slide_height),tokens)
    path=tmp_path/"brand.pptx"
    prs.save(path)
    restored=Presentation(path)
    line=restored.slides[0].shapes[0].chart
    assert str(line.series[0].format.line.color.rgb)=="713CDA"
    for slide in restored.slides:
        chart=slide.shapes[0].chart
        assert chart.legend.position==XL_LEGEND_POSITION.BOTTOM
        assert chart.legend.font.name=="Play"
        assert chart.font.name=="Play"
        assert chart.font.size.pt==12
