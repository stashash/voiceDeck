"""Тесты клонирования слайда-образца. Задача T-08."""
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from designer.export.pptx_clone import clone_slide

REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _flat(shapes, out=None):
    """Фигуры слайда с раскрытием групп."""
    if out is None:
        out = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _flat(shape.shapes, out)
            continue
        out.append(shape)
    return out


def _pictures(slide):
    return [s for s in _flat(slide.shapes) if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _texts(slide):
    return [
        s.text_frame.text
        for s in _flat(slide.shapes)
        if s.has_text_frame and s.text_frame.text.strip()
    ]


def _richest_slide(prs) -> int:
    """Номер слайда с наибольшим числом картинок: на нём проверяем перенос связей."""
    counts = [(len(_pictures(slide)), index) for index, slide in enumerate(prs.slides)]
    best = max(counts)
    if best[0] == 0:
        pytest.skip("в шаблоне нет слайдов с картинками")
    return best[1]


@pytest.fixture(scope="module")
def template(templates):
    return templates[0]


def _saved(prs, tmp_path):
    path = tmp_path / "clone.pptx"
    prs.save(str(path))
    return Presentation(str(path))


def test_clone_repeats_pictures_and_texts(template, tmp_path):
    prs = Presentation(str(template))
    index = _richest_slide(prs)
    source = prs.slides[index]
    expected_pictures = len(_pictures(source))
    expected_texts = sorted(_texts(source))

    clone_slide(prs, index)

    reopened = _saved(prs, tmp_path)
    clone = reopened.slides[len(reopened.slides) - 1]
    assert len(_pictures(clone)) == expected_pictures
    assert sorted(_texts(clone)) == expected_texts


def test_clone_keeps_image_bytes(template, tmp_path):
    prs = Presentation(str(template))
    index = _richest_slide(prs)
    expected = sorted(picture.image.sha1 for picture in _pictures(prs.slides[index]))

    clone_slide(prs, index)

    reopened = _saved(prs, tmp_path)
    clone = reopened.slides[len(reopened.slides) - 1]
    assert sorted(picture.image.sha1 for picture in _pictures(clone)) == expected


def test_clone_relationships_resolve(template, tmp_path):
    prs = Presentation(str(template))
    index = _richest_slide(prs)

    clone_slide(prs, index)

    reopened = _saved(prs, tmp_path)
    clone = reopened.slides[len(reopened.slides) - 1]
    used = {
        value
        for node in clone._element.iter()
        for name, value in node.attrib.items()
        if name.startswith(REL_NS) and value
    }
    assert used
    assert all(rel_id in clone.part.rels for rel_id in used)


def test_clone_keeps_layout_and_background(template, tmp_path):
    prs = Presentation(str(template))
    index = _richest_slide(prs)
    source = prs.slides[index]
    has_background = source._element.find(qn("p:cSld")).find(qn("p:bg")) is not None

    clone_slide(prs, index)

    reopened = _saved(prs, tmp_path)
    clone = reopened.slides[len(reopened.slides) - 1]
    assert clone.slide_layout.name == source.slide_layout.name
    assert (clone._element.find(qn("p:cSld")).find(qn("p:bg")) is not None) == has_background


def test_clone_has_no_notes(template):
    prs = Presentation(str(template))

    clone = clone_slide(prs, _richest_slide(prs))

    assert clone.has_notes_slide is False


def test_two_clones_of_one_sample_are_independent(template, tmp_path):
    prs = Presentation(str(template))
    index = _richest_slide(prs)

    first = clone_slide(prs, index)
    clone_slide(prs, index)
    text_shape = next((s for s in _flat(first.shapes) if s.has_text_frame), None)
    if text_shape is None:
        pytest.skip("на слайде нет текстовых фигур")
    text_shape.text_frame.text = "Правка первого клона"

    reopened = _saved(prs, tmp_path)
    total = len(reopened.slides)
    changed = _texts(reopened.slides[total - 2])
    untouched = _texts(reopened.slides[total - 1])
    assert "Правка первого клона" in changed
    assert "Правка первого клона" not in untouched
