"""Regression for white series/pie sectors disappearing on a near-white slide."""
from itertools import combinations
from types import SimpleNamespace
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor

from designer.contracts import (ChartSpec, ColorToken, Deck, DeckPlan, DesignSystem,
                                Element, Margins, Scene, Series, SlideSpec, Tokens)
from designer.export.html import render_deck
from designer.export.svg_chart import chart_colors, render_chart
from designer.viz.palette import (MIN_CHART_CONTRAST, MIN_CHART_DISTANCE, chart_color_distance,
                                  chart_series_colors, contrast_ratio, contrasting_chart_colors)
from designer.viz.pptx_native import add_chart


def tokens(*colors, background="FFFFFF"):
    return Tokens(colors=[
        ColorToken(hex=background, role="background", share=.6, source="theme"),
        *[ColorToken(hex=color, role="accent" if index == 0 else "accent_alt",
                     share=.1, source="usage") for index, color in enumerate(colors)]
    ], fonts=[], type_scale=[], margins=Margins(left=0, top=0, right=0, bottom=0))


@pytest.mark.parametrize("background", ["FFFFFF", "FAFCFF", "F8F8F4", "000000", "101522", "777777"])
@pytest.mark.parametrize("count", [2, 3, 5, 12])
def test_chart_palette_guarantees_visible_and_distinct_colors(background, count):
    palette = contrasting_chart_colors(["0077FF", "FFFFFF", "00BCD4"], background, count)
    assert len(palette) == count
    assert all(contrast_ratio(color, background) >= MIN_CHART_CONTRAST for color in palette)
    assert all(chart_color_distance(a, b) >= MIN_CHART_DISTANCE for a, b in combinations(palette, 2))
    assert contrasting_chart_colors(palette, background, count) == palette


def test_readable_originals_are_reserved_before_replacing_an_unsafe_color():
    source = ["0077FF", "FFFFFF", "FF00AA"]
    result = chart_series_colors(tokens(*source), background="FFFFFF", count=3)
    assert result[0] == source[0] and result[2] == source[2]
    assert result[1] != "FFFFFF"
    assert source == ["0077FF", "FFFFFF", "FF00AA"]
    assert chart_series_colors(tokens("0077FF", "FFFFFF", "00BCD4"), "dark", "000000", 3) == [
        "0077FF", "FFFFFF", "00BCD4"]


def test_actual_background_override_beats_misleading_theme_tokens():
    ts = tokens("0077FF", "FFFFFF", background="000000")
    assert chart_colors(ts, "dark", "FAFCFF", 2)[1] != "FFFFFF"
    assert chart_colors(ts, "light", "000000", 2)[1] == "FFFFFF"


@pytest.mark.parametrize("override,pattern_background,expected_background", [
    (None, "FAFCFF", "FAFCFF"), ("101522", "FAFCFF", "101522"),
])
def test_native_deck_exporter_passes_actual_pattern_or_editor_background(override, pattern_background, expected_background):
    from designer.export.pptx_deck import _fill_viz
    ts = tokens("0077FF", "FFFFFF", background="000000")
    chart = ChartSpec(type="grouped_column", categories=["A"], series=[
        Series(name="Original", values=[10]), Series(name="Target", values=[12])])
    spec = SlideSpec(slide_id="s", pattern_id="p", chart=chart, viz_box=(.1, .1, .8, .7),
                     background_color=override)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_viz(slide, {}, spec, SimpleNamespace(theme="light", background_color=pattern_background),
              SimpleNamespace(slide_size_emu=(prs.slide_width, prs.slide_height), tokens=ts))
    fills = [str(series.format.fill.fore_color.rgb) for series in slide.shapes[0].chart.series]
    assert fills == chart_series_colors(ts, background=expected_background, count=2)


@pytest.mark.parametrize("kind", ["grouped_column", "stacked_column", "pie", "donut", "line", "scatter"])
@pytest.mark.parametrize("background", ["FAFCFF", "101522"])
def test_active_native_xml_fills_and_strokes_match_svg_without_changing_data(tmp_path, kind, background):
    ts = tokens("0077FF", "FFFFFF", "00BCD4")
    theme = "dark" if background == "101522" else "light"
    pie = kind in {"pie", "donut"}
    spec = ChartSpec(type=kind, categories=[] if kind == "scatter" else ["Alpha", "Beta", "Gamma"],
                     x_axis_title="Input", y_axis_title="Output", unit="USD",
                     series=[Series(name=name, values=[10, 0, 15],
                                    x_values=[1, 2, 3] if kind == "scatter" else None)
                             for name in (["Primary"] if pie else ["Primary", "Target"])])
    original = spec.model_dump()
    count = 3 if pie else 2
    expected = chart_series_colors(ts, theme, background, count)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
    add_chart(slide, spec, (.1, .1, .8, .7), (prs.slide_width, prs.slide_height), ts,
              theme=theme, background=background)
    path = tmp_path / "chart.pptx"
    prs.save(path)
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("ppt/charts/chart1.xml"))
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    chart_node = {"grouped_column": "barChart", "stacked_column": "barChart", "pie": "pieChart",
                  "donut": "doughnutChart", "line": "lineChart", "scatter": "scatterChart"}[kind]
    query = (f".//c:{chart_node}/c:ser/c:dPt/c:spPr/a:solidFill/a:srgbClr" if pie else
             f".//c:{chart_node}/c:ser/c:spPr/a:ln/a:solidFill/a:srgbClr" if kind == "line" else
             f".//c:{chart_node}/c:ser/c:marker/c:spPr/a:solidFill/a:srgbClr" if kind == "scatter" else
             f".//c:{chart_node}/c:ser/c:spPr/a:solidFill/a:srgbClr")
    active_fills = [node.attrib["val"] for node in root.findall(query, ns)]
    assert active_fills == expected
    if pie:
        data_label_colors = [node.attrib["val"] for node in root.findall(
            f".//c:{chart_node}/c:ser/c:dLbls/c:dLbl/c:txPr/a:p/a:pPr/a:defRPr/a:solidFill/a:srgbClr", ns)]
        assert len(data_label_colors) == len(expected)
        assert all(contrast_ratio(ink, fill) >= 4.5 for ink, fill in zip(data_label_colors, expected))
    svg = render_chart(spec, ["0077FF", "FFFFFF", "00BCD4"], background=background)
    for color in expected:
        assert f'#{color}' in svg
        assert contrast_ratio(color, background) >= MIN_CHART_CONTRAST
    reopened = Presentation(path).slides[0].shapes[0].chart
    assert [list(series.values) for series in reopened.series] == [[10, 0, 15]] * len(spec.series)
    assert [series.name for series in reopened.series] == [series.name for series in spec.series]
    if not pie:
        assert reopened.category_axis.axis_title.text_frame.text == "Input"
        assert reopened.value_axis.axis_title.text_frame.text == "Output"
    assert spec.model_dump() == original


def test_html_uses_each_scene_background_instead_of_one_deck_palette(tmp_path):
    ts = tokens("0077FF", "FFFFFF", background="000000")
    chart = ChartSpec(type="grouped_column", categories=["A"], series=[
        Series(name="Original", values=[10]), Series(name="Target", values=[12])])
    scenes = [Scene(slide_id=f"s{i}", pattern_id="p", background_color=background,
                    theme="dark" if i else "light",
                    elements=[Element(id="viz", type="chart", box=(.1, .1, .8, .7), chart=chart)])
              for i, background in enumerate(["FAFCFF", "000000"])]
    ds = DesignSystem(id="ds", source_file="source.pptx", name="Test", slide_size_emu=(9144000, 5143500),
                      tokens=ts)
    deck = Deck(id="deck", design_system_id="ds", variant="a",
                plan=DeckPlan(title="Test", purpose="Test", slides=[]), specs=[], scenes=scenes)
    markup = render_deck(deck, ds, tmp_path)
    sections = re.findall(r'<section\b.*?</section>', markup, flags=re.S)
    assert len(sections) == 2
    for section, background in zip(sections, ["FAFCFF", "000000"]):
        expected = chart_series_colors(ts, background=background, count=2)
        fills = re.findall(r'class="chart-bar"[^>]+fill="#([A-F0-9]+)"', section)
        assert fills == expected
    assert 'fill="#FFFFFF"' not in sections[0]
    assert 'fill="#FFFFFF"' in sections[1]
