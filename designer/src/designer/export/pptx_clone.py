"""Клон слайда-образца внутри той же презентации. Владелец: задача T-08.

python-pptx копировать слайды не умеет. Новый слайд заводится на макете образца,
дерево фигур переносится целиком, а ссылки на части пакета (картинки, гиперссылки,
встроенные объекты) переписываются на связи нового слайда. Заметки не копируются:
слайд ссылается на них не из своего XML, а через связь, которую мы не переносим.
"""
from __future__ import annotations

import copy

from pptx.oxml.ns import qn

_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

_NOT_A_SHAPE = (qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst"))


def clone_slide(prs, source_index: int):
    """Копия слайда source_index в конце презентации. Идентификаторы фигур сохраняются."""
    source = prs.slides[source_index]
    slide = prs.slides.add_slide(source.slide_layout)
    _copy_shapes(source.shapes._spTree, slide.shapes._spTree)
    _copy_look(source._element, slide._element)
    _carry_rels(source.part, slide.part, slide._element)
    return slide


def _copy_shapes(source_tree, target_tree) -> None:
    """Заполнители, добавленные макетом, уходят: на слайде живут фигуры образца."""
    for child in list(target_tree):
        if child.tag not in _NOT_A_SHAPE:
            target_tree.remove(child)
    for child in source_tree:
        if child.tag not in _NOT_A_SHAPE:
            target_tree.append(copy.deepcopy(child))


def _copy_look(source_sld, target_sld) -> None:
    """Фон слайда и карта цветов: без них клон взял бы вид макета, а не образца."""
    source_csld = source_sld.find(qn("p:cSld"))
    target_csld = target_sld.find(qn("p:cSld"))
    if source_csld is not None and target_csld is not None:
        _replace(target_csld, source_csld.find(qn("p:bg")), qn("p:bg"), 0)
    _replace(target_sld, source_sld.find(qn("p:clrMapOvr")), qn("p:clrMapOvr"), 1)


def _replace(parent, node, tag: str, index: int) -> None:
    old = parent.find(tag)
    if old is not None:
        parent.remove(old)
    if node is not None:
        parent.insert(index, copy.deepcopy(node))


def _carry_rels(source_part, target_part, element) -> None:
    """Переписывает ссылки на части пакета: у нового слайда свои идентификаторы связей."""
    renamed: dict[str, str] = {}
    for node in element.iter():
        for name, value in list(node.attrib.items()):
            if not name.startswith(_REL_NS) or not value:
                continue
            new_id = renamed.get(value)
            if new_id is None:
                new_id = _relate(source_part, target_part, value)
                if new_id is None:
                    continue
                renamed[value] = new_id
            node.set(name, new_id)


def _relate(source_part, target_part, rel_id: str) -> str | None:
    """Заводит у нового слайда такую же связь и возвращает её идентификатор."""
    rel = source_part.rels.get(rel_id)
    if rel is None:
        return None
    if rel.is_external:
        return target_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
    return target_part.relate_to(rel.target_part, rel.reltype)
