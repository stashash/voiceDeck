"""Deterministic generated-deck edits, isolated persistence and native export."""
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

from designer import edit, store
from designer.api.app import app
from designer.api.schemas import ElementActionRequest
from designer.contracts import (
    Deck, DeckPlan, DesignSystem, Element, ElementPosition, Pattern, Scene, SlideIntent,
    SlideKind, SlideSpec, Slot, TableSpec, TextStyle, Tokens,
)


@pytest.fixture
def deck(tmp_path, monkeypatch):
    monkeypatch.setenv(store.DATA_DIR_ENV, str(tmp_path / "data"))
    scheduled = []
    monkeypatch.setattr(edit, "schedule", lambda *args: scheduled.append(args))
    monkeypatch.setattr(edit.convert, "available", lambda: False)
    package = store.design_system_dir("voice-v2")
    package.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    title_box, shape_box = (.1, .1, .5, .1), (.1, .3, .3, .2)

    def emu(box):
        return tuple(int(value * (sw if i % 2 == 0 else sh)) for i, value in enumerate(box))

    title = slide.shapes.add_textbox(*emu(title_box))
    title.text = "Original title"
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *emu(shape_box))
    prs.save(str(package / "source.pptx"))
    style = TextStyle(size_pt=24, color="111111", family="Arial")
    pattern = Pattern(id="p1", source_slide=1, layout_name="Blank", kind=SlideKind.title,
                      kind_confidence=1, theme="light", background_color="FFFFFF",
                      slots=[Slot(id="title", role="title", shape_id=title.shape_id,
                                  box=title_box, style=style, max_chars=100, max_lines=2)],
                      decor_shape_ids=[shape.shape_id])
    ds = DesignSystem(id="voice-v2", source_file="source.pptx", slide_size_emu=(sw, sh),
                      tokens=Tokens(colors=[], fonts=[], type_scale=[],
                                    margins={"left": .02, "right": .02, "top": .02, "bottom": .02}),
                      patterns=[pattern])
    store._write_json(package / "manifest.json", ds.model_dump(mode="json"))
    scene = Scene(slide_id="s1", pattern_id="p1", background_color="FFFFFF", elements=[
        Element(id="title", type="text", role="title", box=title_box, z=0,
                text="Original title", style=style, source_shape_id=title.shape_id),
        Element(id="shape", type="shape", box=shape_box, z=1, fill="D9D9D9",
                source_shape_id=shape.shape_id),
    ])
    spec = SlideSpec(slide_id="s1", pattern_id="p1", slot_text={"title": "Original title"})
    plan = DeckPlan(title="Voice test", purpose="test", slides=[
        SlideIntent(id="s1", kind=SlideKind.title, title="Original title")])
    value = Deck(id="voice-v2", design_system_id=ds.id, variant="a", plan=plan, specs=[spec], scenes=[scene])
    store.save_deck_result(value.id, value, [])
    return SimpleNamespace(id=value.id, client=TestClient(app), scheduled=scheduled, tmp=tmp_path,
                           url=f"/decks/{value.id}/a/slides/1/elements")


def saved(deck):
    return store.load_deck_state(deck.id, "a")


def post(deck, **payload):
    response = deck.client.post(deck.url, json=payload)
    assert response.status_code == 200, response.text
    assert response.json()["revision"] == saved(deck)["revision"]
    return response.json()


def element(raw, element_id):
    return next(item for item in raw["scenes"][0]["elements"] if item["id"] == element_id)


def add(deck, kind, **properties):
    before = {item["id"] for item in saved(deck)["scenes"][0]["elements"]}
    raw = post(deck, action="add", element_type=kind, **properties)
    return [item for item in raw["scenes"][0]["elements"] if item["id"] not in before]


def export(deck):
    directory = deck.tmp / "export"
    directory.mkdir(exist_ok=True)
    edit._export_files(edit._load(deck.id, "a"), directory)
    return Presentation(str(directory / "deck.pptx")), (directory / "deck.markup.html").read_text(encoding="utf-8")


def test_text_style_persists_and_undo_restores_original(deck):
    before = saved(deck)
    raw = post(deck, action="style", element_id="title", bold=True, italic=True,
               text_align="right", color="aabbcc", size_pt=32, width=.7, height=.15)
    title = element(raw, "title")
    assert title["style"] == dict(family="Arial", size_pt=32, bold=True, italic=True,
                                  color="AABBCC", align="right")
    assert title["box"] == pytest.approx([.1, .1, .7, .15])
    override = raw["specs"][0]["element_positions"]["title"]
    assert override["original_box"] == [.1, .1, .5, .1]
    assert len(deck.scheduled) == 1
    prs, markup = export(deck)
    native = next(shape for shape in prs.slides[0].shapes if shape.has_text_frame and shape.text == "Original title")
    paragraph = native.text_frame.paragraphs[0]
    assert paragraph.alignment == PP_ALIGN.RIGHT
    assert paragraph.runs[0].font.bold and paragraph.runs[0].font.italic
    assert paragraph.runs[0].font.size.pt == 32
    assert "font-style:italic;" in markup and "text-align:right;" in markup
    raw = post(deck, action="style", element_id="title", bold=False, italic=False)
    assert not element(raw, "title")["style"]["bold"]
    assert not element(raw, "title")["style"]["italic"]
    edit.revert(deck.id, "a")
    assert element(saved(deck), "title") == title
    edit.revert(deck.id, "a")
    after = saved(deck)
    for key in ("scenes", "specs", "plan"):
        assert after[key] == before[key]
    assert not list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))


@pytest.mark.parametrize("payload", [
    {"action": "fly", "element_id": "title"},
    {"action": "style", "element_id": "title"},
    {"action": "style", "element_id": "title", "bold": "true"},
    {"action": "style", "element_id": "title", "bold": None},
    {"action": "style", "element_id": "title", "text_align": "justify"},
    {"action": "style", "element_id": "title", "size_pt": 500},
    {"action": "style", "element_id": "title", "scale": 1, "size_pt": 24},
    {"action": "style", "element_id": "title", "color": "#ffffff"},
    {"action": "style", "element_id": "title", "color": "red"},
    {"action": "style", "element_id": "title", "fill": "123456"},
    {"action": "style", "element_id": "shape", "bold": True, "fill": "123456"},
    {"action": "style", "element_id": "shape", "italic": False},
    {"action": "style", "element_id": "shape", "text_align": "center"},
    {"action": "style", "element_id": "shape", "width": 1},
    {"action": "style", "element_id": "title", "width": 0},
    {"action": "style", "element_id": "title", "height": -.1},
    {"action": "style", "element_id": "title", "width": True},
    {"action": "style", "element_id": "title", "width": ".2"},
    {"action": "style", "bold": True},
    {"action": "delete", "element_id": "title", "text": "ignored"},
    {"action": "duplicate", "element_id": "shape", "width": .2},
    {"action": "background"},
    {"action": "background", "element_id": "title", "color": "AABBCC"},
    {"action": "background", "color": "AABBCC", "unknown": True},
    {"action": "z_order", "element_id": "title", "order": "sideways"},
    {"action": "z_order", "element_id": "title"},
    {"action": "add", "element_type": "video"},
    {"action": "add", "element_type": "shape", "text": "invisible"},
    {"action": "add", "element_type": "text", "fill": "AABBCC"},
    {"action": "add", "element_type": "card", "width": .95},
    {"action": "add", "element_type": "table", "table": {"columns": [], "rows": []}},
    {"action": "add", "element_type": "table", "table": {"columns": ["A"], "rows": [["1", "2"]]}},
    {"action": "add", "element_type": "table", "table": {"columns": ["A"], "rows": [[1]]}},
    {"action": "add", "element_type": "table", "table": {"columns": ["A"], "rows": [], "extra": 1}},
    {"action": "table_cell", "element_id": "shape", "row": 1, "column": 1, "text": "x"},
])
def test_invalid_request_does_not_save_or_consume_undo(deck, payload):
    path = store.deck_variant_state_path(deck.id, "a")
    before = path.read_bytes()
    response = deck.client.post(deck.url, json=payload)
    assert response.status_code == 422, response.text
    assert path.read_bytes() == before
    assert not list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))
    assert not deck.scheduled


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -float("inf")])
def test_nonfinite_dimensions_fail_before_loading(monkeypatch, value):
    monkeypatch.setattr(edit, "_load", lambda *args: pytest.fail("Must validate before loading"))
    with pytest.raises(ValueError):
        edit.element_action("unused", "a", 1, "style", "title", width=value)


def test_missing_target_never_falls_back_to_another_element(deck):
    before = saved(deck)
    response = deck.client.post(deck.url, json={"action": "style", "element_id": "absent", "bold": True})
    assert response.status_code == 404
    assert saved(deck) == before
    assert not deck.scheduled


def test_font_scaling_rejects_out_of_range_result_without_snapshot(deck):
    post(deck, action="style", element_id="title", size_pt=100)
    before = saved(deck)
    response = deck.client.post(deck.url, json={"action": "style", "element_id": "title", "scale": 2})
    assert response.status_code == 422
    assert saved(deck) == before
    assert len(deck.scheduled) == 1
    assert len(list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))) == 1


def test_shape_fill_background_and_layers_export(deck):
    post(deck, action="style", element_id="shape", fill="abcdef", width=.4, height=.25)
    post(deck, action="background", color="102030")
    custom = add(deck, "text", text="Added text")[0]
    post(deck, action="style", element_id=custom["id"], bold=True, italic=True, text_align="center")
    post(deck, action="z_order", element_id=custom["id"], order="back")
    raw = post(deck, action="z_order", element_id="title", order="front")
    levels = {item["id"]: item["z"] for item in raw["scenes"][0]["elements"]}
    assert levels[custom["id"]] < levels["shape"] < levels["title"]
    assert raw["specs"][0]["background_color"] == "102030"
    assert raw["scenes"][0]["background_asset"] is None
    prs, markup = export(deck)
    shapes = list(prs.slides[0].shapes)
    assert shapes[0].name == "VoiceDeck background"
    assert shapes[1].text == "Added text"
    assert shapes[-1].text == "Original title"
    native_shape = shapes[2]
    assert str(native_shape.fill.fore_color.rgb) == "ABCDEF"
    assert native_shape.width / prs.slide_width == pytest.approx(.4)
    assert str(prs.slides[0].background.fill.fore_color.rgb) == "102030"
    assert shapes[1].text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert shapes[1].text_frame.paragraphs[0].runs[0].font.italic
    assert "background-color:#102030;" in markup and "background:#ABCDEF;" in markup


def test_card_is_atomic_shape_and_editable_literal_text(deck):
    before = saved(deck)
    card, text = add(deck, "card", text='Delete this slide? No, keep "Case"!', fill="334455", width=.6, height=.4)
    assert card["type"] == "shape" and card["role"] == "card" and card["fill"] == "334455"
    assert text["type"] == "text" and text["z"] > card["z"]
    assert text["style"]["color"] == "FFFFFF"
    assert text["text"] == 'Delete this slide? No, keep "Case"!'
    assert card["box"][0] < text["box"][0]
    assert len(list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))) == 1
    prs, markup = export(deck)
    assert text["text"] in [shape.text for shape in prs.slides[0].shapes if shape.has_text_frame]
    assert "keep &quot;Case&quot;!" in markup
    edit.revert(deck.id, "a")
    assert saved(deck)["scenes"] == before["scenes"]


def test_table_operations_persist_export_and_undo(deck):
    table = add(deck, "table", table={"columns": ["A", "B"], "rows": [["1", "2"]]})[0]
    target = table["id"]
    post(deck, action="table_cell", element_id=target, row=0, column=1, text="Header")
    post(deck, action="table_cell", element_id=target, row=1, column=2, text="Case: KEEP!")
    post(deck, action="table_row_add", element_id=target, row=1, values=["first", "row"])
    post(deck, action="table_column_add", element_id=target, column=2, text="Middle", values=["x", "y"])
    post(deck, action="table_row_delete", element_id=target, row=2)
    raw = post(deck, action="table_column_delete", element_id=target, column=3)
    expected = {"columns": ["Header", "Middle"], "rows": [["first", "x"]]}
    assert element(raw, target)["table"] == expected
    assert raw["specs"][0]["added_elements"][0]["table"] == expected
    prs, markup = export(deck)
    native = next(shape.table for shape in prs.slides[0].shapes if shape.has_table)
    assert [[cell.text for cell in row.cells] for row in native.rows] == [["Header", "Middle"], ["first", "x"]]
    assert "<th>Middle</th>" in markup and "<td>first</td>" in markup
    for _ in range(6):
        edit.revert(deck.id, "a")
    assert element(saved(deck), target)["table"] == table["table"]
    edit.revert(deck.id, "a")
    assert not saved(deck)["specs"][0]["added_elements"]


@pytest.mark.parametrize("properties", [
    {"action": "table_cell", "row": 2, "column": 1, "text": "x"},
    {"action": "table_cell", "row": -1, "column": 1, "text": "x"},
    {"action": "table_cell", "row": 1, "column": 0, "text": "x"},
    {"action": "table_cell", "row": True, "column": 1, "text": "x"},
    {"action": "table_cell", "row": 1, "column": 1},
    {"action": "table_row_add", "row": 3},
    {"action": "table_row_add", "values": ["too", "many"]},
    {"action": "table_row_delete", "row": 0},
    {"action": "table_row_delete", "row": 2},
    {"action": "table_column_add", "column": 3},
    {"action": "table_column_add", "values": []},
    {"action": "table_column_delete", "column": 1},
])
def test_invalid_table_edits_preserve_state_and_history(deck, properties):
    target = add(deck, "table", table={"columns": ["A"], "rows": [["1"]]})[0]["id"]
    before = saved(deck)
    response = deck.client.post(deck.url, json={"element_id": target, **properties})
    assert response.status_code == 422, response.text
    assert saved(deck) == before
    assert len(list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))) == 1


def test_table_capacity_default_insertions_duplicate_and_delete(deck):
    target = add(deck, "table")[0]["id"]
    for _ in range(5):
        post(deck, action="table_row_add", element_id=target)
    for _ in range(3):
        post(deck, action="table_column_add", element_id=target)
    before = saved(deck)
    for action in ("table_row_add", "table_column_add"):
        response = deck.client.post(deck.url, json={"action": action, "element_id": target})
        assert response.status_code == 422
        assert saved(deck) == before
    raw = post(deck, action="duplicate", element_id=target)
    duplicate = raw["specs"][0]["added_elements"][-1]
    assert duplicate["id"] != target
    assert duplicate["table"] == element(raw, target)["table"]
    post(deck, action="table_cell", element_id=duplicate["id"], row=1, column=1, text="Copy only")
    assert element(saved(deck), target)["table"]["rows"][0][0] == ""
    post(deck, action="delete", element_id=duplicate["id"])
    assert len(saved(deck)["specs"][0]["added_elements"]) == 1


def test_existing_table_updates_plan_and_resizes_exports(deck):
    raw = saved(deck)
    table = {"columns": ["A", "B"], "rows": [["1", "2"]]}
    raw["specs"][0].update(table=table, viz_box=[.5, .5, .4, .3])
    raw["plan"]["slides"][0]["table"] = table
    raw["scenes"][0]["elements"].append(Element(id="viz", type="table", role="viz", box=(.5, .5, .4, .3),
                                                table=TableSpec(**table), z=2).model_dump(mode="json"))
    store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    post(deck, action="table_cell", element_id="viz", row=1, column=1, text="Changed")
    post(deck, action="style", element_id="viz", width=.3, height=.2)
    post(deck, action="z_order", element_id="viz", order="back")
    current = saved(deck)
    assert current["plan"]["slides"][0]["table"]["rows"][0][0] == "Changed"
    prs, _ = export(deck)
    native = next(shape for shape in prs.slides[0].shapes if shape.has_table)
    assert native.table.cell(1, 0).text == "Changed"
    assert native.width / prs.slide_width == pytest.approx(.3)
    assert native.height / prs.slide_height == pytest.approx(.2)
    post(deck, action="delete", element_id="viz")
    prs, _ = export(deck)
    assert not any(shape.has_table for shape in prs.slides[0].shapes)


def test_table_dimensions_remain_exact_with_template_type_scale(deck):
    manifest = store.design_system_dir("voice-v2") / "manifest.json"
    ds = store.read_json_file(manifest)
    ds["tokens"]["type_scale"] = [{"size_pt": 12, "role": "body", "share": 1}]
    store._write_json(manifest, ds)
    table = add(deck, "table", height=.5)[0]
    prs, _ = export(deck)
    native = next(shape for shape in prs.slides[0].shapes if shape.has_table)
    assert native.height / prs.slide_height == pytest.approx(table["box"][3])
    assert sum(row.height for row in native.table.rows) == native.height


@pytest.mark.parametrize("align", [None, "left", "center", "right"])
def test_grouped_element_can_move_in_front_of_top_level_content(deck, align):
    path = store.design_system_dir("voice-v2") / "source.pptx"
    prs = Presentation(str(path))
    slide = prs.slides[0]
    title = slide.shapes[0]
    group = slide.shapes.add_group_shape([title])
    # A scaled parent exercises conversion back to absolute slide coordinates.
    group.width *= 2
    prs.save(str(path))
    raw = saved(deck)
    raw["scenes"][0]["elements"][0]["box"] = [.1, .1, 1, .1]
    raw["scenes"][0]["elements"][0]["style"]["align"] = align
    manifest = path.parent / "manifest.json"
    ds = store.read_json_file(manifest)
    ds["patterns"][0]["slots"][0]["box"] = [.1, .1, 1, .1]
    ds["patterns"][0]["slots"][0]["style"]["align"] = align
    store._write_json(manifest, ds)
    store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    post(deck, action="style", element_id="title", width=.6, italic=True)
    post(deck, action="z_order", element_id="title", order="front")
    prs, _ = export(deck)
    title = prs.slides[0].shapes[-1]
    assert title.has_text_frame and title.text == "Original title"
    assert title.width / prs.slide_width == pytest.approx(.6)
    assert title.left / prs.slide_width == pytest.approx(.1)
    assert title.text_frame.paragraphs[0].runs[0].font.italic


@pytest.mark.parametrize("scale", [1.3, 1.7, 2.3])
@pytest.mark.parametrize("align", [None, "left", "center", "right"])
def test_margin_normalization_never_deletes_overlapping_decoy(deck, scale, align):
    from designer.export.pptx_deck import _fill_slide, _walk

    state = edit._load(deck.id, "a")
    prs = Presentation(str(state.package_dir / "source.pptx"))
    slide = prs.slides[0]
    target = slide.shapes[0]
    target.width = round(.8 * prs.slide_width)
    group = slide.shapes.add_group_shape([target])
    group.width = round(group.width * scale)
    original = _walk(slide.shapes, state.ds.slide_size_emu)[target.shape_id].box
    pattern = state.ds.patterns[0]
    pattern.slots[0].box = original
    pattern.slots[0].style.align = align
    decoy = slide.shapes.add_textbox(round(.1 * prs.slide_width), round(.1 * prs.slide_height),
                                    round(.88 * prs.slide_width), round(.1 * prs.slide_height))
    decoy.text = "Do not delete"
    pattern.decor_shape_ids.append(decoy.shape_id)
    spec = state.specs[0]
    spec.element_positions["title"] = ElementPosition(original_box=original, box=original,
                                                      source_shape_id=target.shape_id, deleted=True)
    _fill_slide(slide, spec, pattern, state.ds)
    remaining = _walk(slide.shapes, state.ds.slide_size_emu)
    assert target.shape_id not in remaining
    assert decoy.shape_id in remaining
    assert decoy.text == "Do not delete"


@pytest.mark.parametrize("grouped", [False, True])
def test_background_replaces_full_bleed_asset_and_undo_restores_it(deck, grouped):
    from PIL import Image

    path = store.design_system_dir("voice-v2") / "source.pptx"
    picture_path = deck.tmp / "background.png"
    Image.new("RGB", (8, 8), color="blue").save(picture_path)
    prs = Presentation(str(path))
    slide = prs.slides[0]
    picture = slide.shapes.add_picture(str(picture_path), 0, 0, prs.slide_width, prs.slide_height)
    asset_id = picture.image.sha1[:12]
    if grouped:
        slide.shapes.add_group_shape([picture])
    prs.save(str(path))
    manifest = path.parent / "manifest.json"
    ds = store.read_json_file(manifest)
    ds["patterns"][0]["background_asset"] = asset_id
    store._write_json(manifest, ds)
    raw = saved(deck)
    raw["scenes"][0]["background_asset"] = asset_id
    store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    post(deck, action="background", color="334455")
    prs, _ = export(deck)
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in prs.slides[0].shapes)
    edit.revert(deck.id, "a")
    assert saved(deck)["scenes"][0]["background_asset"] == asset_id
    prs, _ = export(deck)
    from designer.export.pptx_deck import _walk
    assert any(item.kind == "image" for item in _walk(prs.slides[0].shapes, (prs.slide_width, prs.slide_height)).values())


@pytest.mark.parametrize("payload", [
    {"action": "style", "bold": True}, {"action": "delete"}, {"action": "z_order", "order": "front"},
])
def test_split_number_caption_cannot_silently_mutate_shared_native_shape(deck, payload):
    raw = saved(deck)
    title = element(raw, "title")
    caption = dict(title, id="titlec", role="caption", text="Caption", box=[.1, .2, .5, .1])
    raw["scenes"][0]["elements"].append(caption)
    store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    response = deck.client.post(deck.url, json={"element_id": "title", **payload})
    assert response.status_code == 422
    assert saved(deck) == raw
    assert not deck.scheduled


def test_schema_preserves_original_minimal_payloads():
    for payload in ({"action": "add"}, {"action": "add", "element_type": "title", "text": "Title"},
                    {"action": "style", "element_id": "x", "scale": 1.2},
                    {"action": "delete", "element_id": "x"}, {"action": "duplicate", "element_id": "x"}):
        assert ElementActionRequest.model_validate(payload).model_dump(exclude_unset=True) == payload
