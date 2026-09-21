"""Тесты add_table: текст ячеек, заливка шапки, рамка, лимиты строк и колонок. Владелец: T-05, тема — T-20."""
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.util import Pt

from designer.contracts import ColorToken, FontToken, Margins, TableSpec, Tokens, TypeStep
from designer.viz import palette
from designer.viz.pptx_native import add_table

SLIDE_SIZE = (9144000, 5143500)


def _tokens() -> Tokens:
    return Tokens(
        colors=[
            ColorToken(hex="0077FF", role="accent", share=0.3, source="theme"),
            ColorToken(hex="EEEEEE", role="surface", share=0.4, source="usage"),
        ],
        fonts=[FontToken(family="PT Sans", role="body", share=0.8)],
        type_scale=[TypeStep(size_pt=12, role="body", share=0.5)],
        margins=Margins(left=0.03, top=0.06, right=0.03, bottom=0.06),
    )


def _blank_slide():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_SIZE
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide


def _roundtrip(prs: Presentation, tmp_path: Path) -> Presentation:
    path = tmp_path / "out.pptx"
    prs.save(str(path))
    return Presentation(str(path))


def _spec(rows: int = 2, cols: int = 3) -> TableSpec:
    columns = [f"стлб{i}" for i in range(cols)]
    body = [[f"{r}-{c}" for c in range(cols)] for r in range(rows)]
    return TableSpec(columns=columns, rows=body)


def test_add_table_cell_text_and_header_fill(tmp_path):
    tokens = _tokens()
    prs, slide = _blank_slide()
    spec = _spec(rows=2, cols=3)

    add_table(slide, spec, (0.1, 0.1, 0.6, 0.3), SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    table = reopened.slides[0].shapes[0].table
    assert table.cell(0, 0).text == "стлб0"
    assert table.cell(1, 1).text == "0-1"
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "0077FF"


def test_add_table_box_matches_within_one_percent(tmp_path):
    tokens = _tokens()
    prs, slide = _blank_slide()
    spec = _spec(rows=2, cols=2)
    box = (0.1, 0.2, 0.5, 0.3)

    add_table(slide, spec, box, SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    shape = reopened.slides[0].shapes[0]
    slide_w, slide_h = SLIDE_SIZE
    assert shape.left == pytest.approx(box[0] * slide_w, rel=0.01)
    assert shape.top == pytest.approx(box[1] * slide_h, rel=0.01)
    assert shape.width == pytest.approx(box[2] * slide_w, rel=0.01)
    assert shape.height == pytest.approx(box[3] * slide_h, rel=0.01)


def test_add_table_more_than_seven_rows_raises():
    tokens = _tokens()
    _, slide = _blank_slide()
    spec = _spec(rows=8, cols=2)

    with pytest.raises(ValueError):
        add_table(slide, spec, (0.1, 0.1, 0.5, 0.3), SLIDE_SIZE, tokens)


def test_add_table_more_than_five_columns_raises():
    tokens = _tokens()
    _, slide = _blank_slide()
    spec = _spec(rows=2, cols=6)

    with pytest.raises(ValueError):
        add_table(slide, spec, (0.1, 0.1, 0.5, 0.3), SLIDE_SIZE, tokens)


def test_add_table_header_row_not_taller_than_one_and_half_data_rows(tmp_path):
    tokens = _tokens()
    prs, slide = _blank_slide()
    spec = _spec(rows=4, cols=2)

    add_table(slide, spec, (0.1, 0.1, 0.5, 0.4), SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    rows = reopened.slides[0].shapes[0].table.rows
    assert rows[0].height <= 1.5 * rows[1].height


# ---------- тема: T-20 ----------

def _dark_tokens(surface_hex: str) -> Tokens:
    return Tokens(
        colors=[
            ColorToken(hex="0B1220", role="background", share=0.5, source="theme"),
            ColorToken(hex="66CCFF", role="accent", share=0.3, source="theme"),
            ColorToken(hex="F2F2F2", role="text", share=0.2, source="usage"),
            ColorToken(hex=surface_hex, role="surface", share=0.2, source="usage"),
        ],
        fonts=[FontToken(family="PT Sans", role="body", share=0.8)],
        type_scale=[TypeStep(size_pt=12, role="body", share=0.5)],
        margins=Margins(left=0.03, top=0.06, right=0.03, bottom=0.06),
    )


def _body_cell_color(table, row_idx: int, col_idx: int = 0) -> str:
    run = table.cell(row_idx, col_idx).text_frame.paragraphs[0].runs[0]
    return str(run.font.color.rgb)


def test_add_table_dark_theme_body_text_meets_contrast(tmp_path):
    tokens = _dark_tokens("1A2233")  # тёмная surface
    prs, slide = _blank_slide()
    spec = _spec(rows=2, cols=2)

    add_table(slide, spec, (0.1, 0.1, 0.6, 0.3), SLIDE_SIZE, tokens, theme="dark")

    reopened = _roundtrip(prs, tmp_path)
    table = reopened.slides[0].shapes[0].table
    for row_idx in range(2):
        color = _body_cell_color(table, row_idx + 1)
        assert palette.contrast_ratio(color, "0B1220") >= 4.5


def test_add_table_dark_theme_shades_rows_with_a_dark_surface(tmp_path):
    tokens = _dark_tokens("1A2233")  # тёмная surface: заливка строк остаётся
    prs, slide = _blank_slide()
    spec = _spec(rows=2, cols=2)

    add_table(slide, spec, (0.1, 0.1, 0.6, 0.3), SLIDE_SIZE, tokens, theme="dark")

    reopened = _roundtrip(prs, tmp_path)
    table = reopened.slides[0].shapes[0].table
    assert table.cell(1, 0).fill.type == MSO_FILL_TYPE.SOLID
    assert str(table.cell(1, 0).fill.fore_color.rgb) == "1A2233"


def test_add_table_dark_theme_skips_shading_with_a_light_surface(tmp_path):
    tokens = _dark_tokens("EEEEEE")  # светлая surface: светлым текстом на ней не читается
    prs, slide = _blank_slide()
    spec = _spec(rows=2, cols=2)

    add_table(slide, spec, (0.1, 0.1, 0.6, 0.3), SLIDE_SIZE, tokens, theme="dark")

    reopened = _roundtrip(prs, tmp_path)
    table = reopened.slides[0].shapes[0].table
    assert table.cell(1, 0).fill.type == MSO_FILL_TYPE.BACKGROUND


def test_add_table_light_theme_leaves_body_text_color_unset(tmp_path):
    tokens = _dark_tokens("1A2233")
    prs, slide = _blank_slide()
    spec = _spec(rows=2, cols=2)

    add_table(slide, spec, (0.1, 0.1, 0.6, 0.3), SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    table = reopened.slides[0].shapes[0].table
    run = table.cell(1, 0).text_frame.paragraphs[0].runs[0]
    assert run.font.color.type is None


# ---------- кегль текста по высоте рамки: T-20 ----------

def _tokens_with_caption_step() -> Tokens:
    tokens = _tokens()
    tokens.type_scale.append(TypeStep(size_pt=8, role="caption", share=0.3))
    return tokens


def test_add_table_text_size_uses_caption_step_at_or_below_half_slide(tmp_path):
    tokens = _tokens_with_caption_step()
    prs, slide = _blank_slide()
    spec = _spec(rows=1, cols=2)

    add_table(slide, spec, (0.1, 0.1, 0.5, 0.3), SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    run = reopened.slides[0].shapes[0].table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    assert run.font.size == Pt(8)
