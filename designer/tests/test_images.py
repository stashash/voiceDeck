"""Картинки слайдов: главное фото образца, граф ComfyUI, шаг иллюстраций, подмена фото в pptx и html."""
from __future__ import annotations

import io
import json
from pathlib import Path

import httpx
import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from designer import store
from designer.contracts import DeckPlan, Item, SlideImage, SlideIntent, SlideKind
from designer.export.pptx_deck import export_pptx
from designer.images.engine import ComfyEngine, image_size
from designer.images.illustrate import IMAGE_LIMIT, candidates, illustrate
from designer.layout.compose import compose
from designer.layout.match import choose_pattern, hero_area
from designer.layout.scene import build_scene
from designer.llm.client import LlmClient
from designer.parse.package import load_package
from designer.plan.planner import plan_schema

DEMO = Path(__file__).resolve().parents[2] / "demo" / "design-systems" / "shablon-prezentacii-vk-education"
GRAPH = json.loads((Path(__file__).resolve().parents[1] / "skills" / "illustrate" / "graph.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def ds():
    if not (DEMO / "manifest.json").is_file():
        pytest.skip("нет демо-пакета VK Education")
    return load_package(DEMO)


def _png(color=(200, 40, 40), size=(320, 240)) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, "PNG")
    return buffer.getvalue()


def _intent(path: str = "") -> SlideIntent:
    return SlideIntent(id="s2", kind=SlideKind.image_text, title="Курьер у двери офиса вечером",
                       key_message="Онлайн-заказы приходят и после закрытия магазинов",
                       image=SlideImage(prompt="courier", path=path) if path else None)


def test_hero_is_the_single_big_photo(ds):
    by_id = {p.id: p for p in ds.patterns}
    assert hero_area(by_id["p034"]) is not None  # одно фото на половину слайда
    assert hero_area(by_id["p031"]) is None  # экран телефона и скриншот: два крупных фото


def test_slide_with_image_gets_photo_pattern(ds):
    pattern = choose_pattern(_intent("decks/x/images/s2.png"), ds)
    assert hero_area(pattern) is not None
    plain = choose_pattern(_intent(), ds)
    assert not plain.needs_images  # без картинки образцы с фото по-прежнему не берутся


def test_plan_schema_has_no_image_field():
    assert "image" not in plan_schema()["$defs"]["SlideIntent"]["properties"]


def test_image_size_keeps_aspect_and_step():
    width, height = image_size(16 / 9)
    assert width % 16 == 0 and height % 16 == 0 and abs(width / height - 16 / 9) < 0.03


def test_engine_fills_graph_and_returns_png():
    seen = {}
    png = _png()

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/prompt":
            seen["graph"] = json.loads(request.content)["prompt"]
            return httpx.Response(200, json={"prompt_id": "p1"})
        if request.url.path == "/history/p1":
            return httpx.Response(200, json={"p1": {"status": {"completed": True}, "outputs": {
                "13": {"images": [{"filename": "voicedeck_0001.png", "subfolder": "", "type": "output"}]}}}})
        if request.url.path == "/view":
            return httpx.Response(200, content=png)
        return httpx.Response(404)

    engine = ComfyEngine("http://comfy:8189", GRAPH, transport=httpx.MockTransport(handler))
    assert engine.generate("a courier at the door", 1024, 768, 7) == png
    graph = seen["graph"]
    assert graph["4"]["inputs"]["text"] == "a courier at the door"
    assert graph["7"]["inputs"]["width"] == 1024 and graph["8"]["inputs"]["height"] == 768
    assert graph["10"]["inputs"]["noise_seed"] == 7


def test_illustrate_sets_images_only_where_they_fit(ds, tmp_path, monkeypatch):
    monkeypatch.setenv(store.DATA_DIR_ENV, str(tmp_path / "data"))

    def llm(request):
        content = json.dumps({"prompt": "evening office door, courier with bags, warm side light", "alt": "курьер"})
        return httpx.Response(200, json={"choices": [{"message": {"content": content}}]})

    class Engine:
        calls = 0

        def generate(self, prompt, width, height, seed):
            Engine.calls += 1
            return _png(size=(width // 4, height // 4))

    slides = [SlideIntent(id="s1", kind=SlideKind.title, title="Итоги года")]
    slides += [SlideIntent(id=f"s{i}", kind=SlideKind.image_text, title=f"Сцена {i}", key_message="мысль")
               for i in range(2, 8)]
    slides += [SlideIntent(id="s8", kind=SlideKind.chart, title="Данные"),
               SlideIntent(id="s9", kind=SlideKind.thanks, title="Спасибо")]
    plan = DeckPlan(title="Итоги", purpose="отчёт", slides=slides)
    client = LlmClient("http://test/v1", "m", transport=httpx.MockTransport(llm))
    done = illustrate(plan, ds, client, Engine(), "deck1")
    assert len(done) == IMAGE_LIMIT == Engine.calls
    for slide in plan.slides:
        if slide.id in done:
            assert (store.data_dir() / slide.image.path).is_file()
            assert slide.image.alt == "курьер"
        else:
            assert slide.image is None
    assert plan.slides[0].image is None and plan.slides[-2].image is None


def test_pptx_and_html_show_the_slide_image(ds, tmp_path, monkeypatch):
    monkeypatch.setenv(store.DATA_DIR_ENV, str(tmp_path / "data"))
    rel = "decks/x/images/s2.png"
    target = store.data_dir() / rel
    target.parent.mkdir(parents=True)
    png = _png(color=(10, 200, 90), size=(800, 450))
    target.write_bytes(png)

    intent = _intent(rel)
    pattern = choose_pattern(intent, ds)
    spec = compose(intent, pattern, ds)
    hero = hero_area(pattern)
    assert spec.image_area_id == hero.id and hero.shape_id not in spec.remove_shape_ids

    out = export_pptx([spec], ds, DEMO, tmp_path / "deck.pptx")
    slide = Presentation(str(out)).slides[0]
    blobs = [shape.image.blob for shape in slide.shapes if shape.shape_type == 13]
    assert png in blobs
    pic = next(shape for shape in slide.shapes if shape.shape_type == 13 and shape.image.blob == png)
    assert pic._element.find(".//" + qn("a:srcRect")) is not None

    scene = build_scene(spec, pattern, ds, DEMO)
    assert any(el.image == rel for el in scene.elements)
