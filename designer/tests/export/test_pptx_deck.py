"""Тесты сборки pptx: клон образца, блоки, диаграмма, чистый файл. Задача T-08."""
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from designer.contracts import ChartSpec, DesignSystem, Pattern, RepeatGroup, Series, SlideSpec, TableSpec
from designer.export.pptx_deck import export_pptx
from designer.layout.units import place_units
from designer.parse.package import build_package
from designer.parse.patterns import ThemeInfo, flatten_shapes

REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
BOX_TOLERANCE = 0.005
SMALL_UNITS = 3
BIG_UNITS = 5
FULL_FRAME = 0.9


# ---------- разбор собранного файла ----------

def _infos(slide, slide_size):
    """Фигуры слайда с раскрытием групп и рамками в долях слайда."""
    return flatten_shapes(slide.shapes, slide_size, ThemeInfo())


def _texts(slide, slide_size):
    return [info.text for info in _infos(slide, slide_size) if info.text]


def _font_parts(pptx_path) -> int:
    """Сколько встроенных шрифтов лежит в пакете."""
    with zipfile.ZipFile(str(pptx_path)) as package:
        return len([name for name in package.namelist() if "/fonts/" in name])


def _bounds(boxes):
    left = min(b[0] for b in boxes)
    top = min(b[1] for b in boxes)
    right = max(b[0] + b[2] for b in boxes)
    bottom = max(b[1] + b[3] for b in boxes)
    return (left, top, right - left, bottom - top)


# ---------- план колоды ----------

@dataclass
class _Plan:
    specs: list = field(default_factory=list)
    small: tuple = ()
    big: tuple = ()
    chart_pattern: Pattern | None = None
    table_pattern: Pattern | None = None


def _slot_text(pattern: Pattern, tag: str) -> dict[str, str]:
    return {slot.id: f"{tag} поле {i}" for i, slot in enumerate(pattern.slots)}


def _unit_text(group: RepeatGroup, count: int) -> list[dict[str, str]]:
    return [
        {slot.id: f"блок {i} поле {k}" for k, slot in enumerate(group.unit_slots)}
        for i in range(count)
    ]


def _pick_group(ds: DesignSystem, used: set[str], count: int):
    """Паттерн с одним блоком под нужное число блоков: в блоке только текст, по слоту на фигуру."""
    for pattern in ds.patterns:
        if pattern.id in used or len(pattern.groups) != 1:
            continue
        for group in pattern.groups:
            per_unit = len(group.unit_slots)
            if not per_unit or group.unit_areas:
                continue
            if any(len(unit.shape_ids) != per_unit for unit in group.units):
                continue
            if not group.min_units <= count <= group.max_units or len(group.units) == count:
                continue
            used.add(pattern.id)
            return pattern, group
    return None, None


def _pick_plain(ds: DesignSystem, used: set[str], with_area: bool):
    for pattern in ds.patterns:
        if pattern.id in used or pattern.groups or not pattern.slots:
            continue
        if with_area and not pattern.areas:
            continue
        used.add(pattern.id)
        return pattern
    return None


def _plan(ds: DesignSystem) -> _Plan:
    """Пять слайдов: блок на три, блок на пять, диаграмма, таблица и обычный слайд."""
    used: set[str] = set()
    small = _pick_group(ds, used, SMALL_UNITS)
    big = _pick_group(ds, used, BIG_UNITS)
    chart_pattern = _pick_plain(ds, used, with_area=True)
    table_pattern = _pick_plain(ds, used, with_area=False)
    plain = _pick_plain(ds, used, with_area=False)
    if small[0] is None or big[0] is None or None in (chart_pattern, table_pattern, plain):
        pytest.skip("в шаблоне нет паттернов под такую колоду")

    plan = _Plan(small=small, big=big, chart_pattern=chart_pattern, table_pattern=table_pattern)
    for (pattern, group), count, tag in ((small, SMALL_UNITS, "малый"), (big, BIG_UNITS, "большой")):
        plan.specs.append(SlideSpec(
            slide_id=tag,
            pattern_id=pattern.id,
            slot_text=_slot_text(pattern, tag),
            group_id=group.id,
            unit_text=_unit_text(group, count),
            notes=f"Заметки докладчика: {tag}",
        ))
    plan.specs.append(SlideSpec(
        slide_id="диаграмма",
        pattern_id=chart_pattern.id,
        slot_text=_slot_text(chart_pattern, "диаграмма"),
        chart=ChartSpec(
            type="column",
            categories=["раз", "два", "три"],
            series=[Series(name="выручка", values=[3.0, 5.0, 8.0])],
            unit="млн",
        ),
        viz_area_id=chart_pattern.areas[0].id,
    ))
    plan.specs.append(SlideSpec(
        slide_id="таблица",
        pattern_id=table_pattern.id,
        slot_text=_slot_text(table_pattern, "таблица"),
        table=TableSpec(columns=["этап", "срок"], rows=[["разбор", "день"], ["сборка", "час"]]),
    ))
    plan.specs.append(SlideSpec(
        slide_id="обычный",
        pattern_id=plain.id,
        slot_text=_slot_text(plain, "обычный"),
    ))
    return plan


def _generic_specs(ds: DesignSystem, limit: int) -> list[SlideSpec]:
    """Инструкции по первым паттернам шаблона: заполняются все слоты и все блоки."""
    specs = []
    for index, pattern in enumerate(ds.patterns[:limit]):
        group = pattern.groups[0] if pattern.groups else None
        count = 0
        if group is not None:
            count = max(group.min_units, min(len(group.units), group.max_units))
        specs.append(SlideSpec(
            slide_id=f"s{index}",
            pattern_id=pattern.id,
            slot_text=_slot_text(pattern, f"слайд {index}"),
            group_id=group.id if group else None,
            unit_text=_unit_text(group, count) if group else [],
        ))
    return specs


# ---------- фикстуры ----------

@dataclass
class _Built:
    ds: DesignSystem
    plan: _Plan
    path: Path
    prs: Presentation


@pytest.fixture(scope="session")
def packages(tmp_path_factory):
    """Пакет дизайн-системы по шаблону, один раз на прогон."""
    cache: dict[str, tuple[Path, DesignSystem]] = {}

    def build(pptx_path: Path):
        key = str(pptx_path)
        if key not in cache:
            out_dir = tmp_path_factory.mktemp("package")
            cache[key] = (out_dir, build_package(Path(pptx_path), out_dir))
        return cache[key]

    return build


@pytest.fixture(scope="session")
def decks(packages, templates, tmp_path_factory):
    """Собранная колода по шаблону с этим номером."""
    cache: dict[int, _Built] = {}

    def build(index: int) -> _Built:
        if index >= len(templates):
            pytest.skip(f"выдано шаблонов меньше {index + 1}")
        if index not in cache:
            package_dir, ds = packages(templates[index])
            plan = _plan(ds)
            out_path = tmp_path_factory.mktemp("deck") / "deck.pptx"
            export_pptx(plan.specs, ds, package_dir, out_path)
            cache[index] = _Built(
                ds=ds, plan=plan, path=out_path, prs=Presentation(str(out_path)),
            )
        return cache[index]

    return build


# ---------- колода целиком ----------

@pytest.mark.parametrize("index", [0, 1, 2])
def test_deck_holds_only_assembled_slides(decks, index):
    built = decks(index)
    slide_size = built.ds.slide_size_emu

    assert len(built.prs.slides) == len(built.plan.specs)

    samples = set()
    for spec in built.plan.specs:
        pattern = next(p for p in built.ds.patterns if p.id == spec.pattern_id)
        for slot in pattern.slots:
            samples.add(slot.sample_text)
        for group in pattern.groups:
            for slot in group.unit_slots:
                samples.add(slot.sample_text)
    samples = {text[:40] for text in samples if len(text) >= 5}

    written = "\n".join(
        text for slide in built.prs.slides for text in _texts(slide, slide_size)
    )
    assert not [sample for sample in samples if sample in written]


@pytest.mark.parametrize("index", [0, 1, 2])
def test_no_slide_is_a_single_picture(decks, index):
    built = decks(index)
    for slide in built.prs.slides:
        infos = _infos(slide, built.ds.slide_size_emu)
        pictures = [i for i in infos if i.kind == "image"]
        covers = [i for i in pictures if i.box[2] >= FULL_FRAME and i.box[3] >= FULL_FRAME]
        assert not (len(infos) == 1 and covers)


@pytest.mark.parametrize("index", [0, 1, 2])
def test_slide_relationships_resolve(decks, index):
    built = decks(index)
    for slide in built.prs.slides:
        used = {
            value
            for node in slide._element.iter()
            for name, value in node.attrib.items()
            if name.startswith(REL_NS) and value
        }
        assert all(rel_id in slide.part.rels for rel_id in used)


@pytest.mark.parametrize("index", [0, 1, 2])
def test_shape_ids_are_unique_on_every_slide(decks, index):
    built = decks(index)
    for slide in built.prs.slides:
        ids = [node.get("id") for node in slide._element.iter(qn("p:cNvPr"))]
        assert len(ids) == len(set(ids))


def test_layouts_masters_and_embedded_fonts_survive(decks, templates):
    built = decks(0)
    source = Presentation(str(templates[0]))

    assert len(built.prs.slide_masters) == len(source.slide_masters)
    assert sum(len(m.slide_layouts) for m in built.prs.slide_masters) == sum(
        len(m.slide_layouts) for m in source.slide_masters
    )
    assert _font_parts(built.path) == _font_parts(templates[0])


# ---------- блоки ----------

def _check_units(built: _Built, slide_index: int, picked, count: int):
    pattern, group = picked
    slide = built.prs.slides[slide_index]
    infos = _infos(slide, built.ds.slide_size_emu)
    per_unit = len(group.unit_slots)

    unit_shapes = [info for info in infos if info.text.startswith("блок ")]
    assert len(unit_shapes) == count * per_unit

    expected = place_units(group, count)
    for i in range(count):
        boxes = [info.box for info in infos if info.text.startswith(f"блок {i} ")]
        assert len(boxes) == per_unit
        actual = _bounds(boxes)
        for side in range(4):
            assert abs(actual[side] - expected[i][side]) <= BOX_TOLERANCE

    others = (
        len(pattern.slots)
        + len(pattern.areas)
        + len(pattern.decor_shape_ids)
        + sum(len(u.shape_ids) for g in pattern.groups for u in g.units)
        - sum(len(u.shape_ids) for u in group.units)
    )
    assert len(infos) == others + count * per_unit


def test_group_shrinks_to_three_units(decks):
    built = decks(0)
    _check_units(built, 0, built.plan.small, SMALL_UNITS)


def test_group_fits_five_units(decks):
    built = decks(0)
    _check_units(built, 1, built.plan.big, BIG_UNITS)


def _pick_growing_group(ds: DesignSystem):
    """Блок, куда помещается больше блоков, чем на образце: проверяем копирование."""
    for pattern in ds.patterns:
        for group in pattern.groups:
            per_unit = len(group.unit_slots)
            if not per_unit or group.unit_areas or group.max_units <= len(group.units):
                continue
            if any(len(unit.shape_ids) != per_unit for unit in group.units):
                continue
            return pattern, group
    return None, None


@pytest.mark.parametrize("index", [0, 1, 2])
def test_missing_units_are_copied_from_the_sample(packages, templates, tmp_path, index):
    if index >= len(templates):
        pytest.skip(f"выдано шаблонов меньше {index + 1}")
    package_dir, ds = packages(templates[index])
    pattern, group = _pick_growing_group(ds)
    if pattern is None:
        pytest.skip("в шаблоне нет блока с запасом мест")
    count = len(group.units) + 1
    spec = SlideSpec(
        slide_id="рост",
        pattern_id=pattern.id,
        group_id=group.id,
        unit_text=_unit_text(group, count),
    )

    out_path = export_pptx([spec], ds, package_dir, tmp_path / "grown.pptx")

    slide = Presentation(str(out_path)).slides[0]
    infos = _infos(slide, ds.slide_size_emu)
    per_unit = len(group.unit_slots)
    assert len([info for info in infos if info.text.startswith("блок ")]) == count * per_unit
    expected = place_units(group, count)
    for i in range(count):
        boxes = [info.box for info in infos if info.text.startswith(f"блок {i} ")]
        assert len(boxes) == per_unit
        actual = _bounds(boxes)
        for side in range(4):
            assert abs(actual[side] - expected[i][side]) <= BOX_TOLERANCE


def test_notes_are_written(decks):
    built = decks(0)
    slide = built.prs.slides[0]
    assert slide.has_notes_slide
    assert slide.notes_slide.notes_text_frame.text == built.plan.specs[0].notes


# ---------- диаграмма и таблица ----------

def test_chart_slide_holds_native_chart(decks):
    built = decks(0)
    slide = built.prs.slides[2]
    charts = [shape for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 1
    assert [s.name for s in charts[0].chart.series] == ["выручка"]


def test_table_slide_holds_native_table(decks):
    built = decks(0)
    slide = built.prs.slides[3]
    tables = [shape for shape in slide.shapes if shape.has_table]
    assert len(tables) == 1
    assert tables[0].table.cell(0, 0).text == "этап"


# ---------- чужой шаблон и время сборки ----------

def test_foreign_template_passes_the_same_path(packages, foreign_pptx, tmp_path):
    package_dir, ds = packages(foreign_pptx)
    specs = _generic_specs(ds, 5)
    if not specs:
        pytest.skip("в стороннем файле нет слайдов")

    out_path = export_pptx(specs, ds, package_dir, tmp_path / "foreign.pptx")

    reopened = Presentation(str(out_path))
    assert len(reopened.slides) == len(specs)


def test_fifteen_slides_build_under_a_minute(packages, templates, tmp_path):
    package_dir, ds = packages(templates[0])
    specs = _generic_specs(ds, 15)
    if len(specs) < 15:
        pytest.skip("в шаблоне меньше 15 слайдов-образцов")

    started = time.monotonic()
    out_path = export_pptx(specs, ds, package_dir, tmp_path / "long.pptx")
    spent = time.monotonic() - started

    assert spent < 60
    assert len(Presentation(str(out_path)).slides) == 15
