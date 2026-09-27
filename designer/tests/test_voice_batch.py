"""Atomic voice batches against real deck persistence and native exports."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from copy import deepcopy
from threading import Barrier, Event
from unittest.mock import Mock
import json

import pytest
from pptx.enum.text import PP_ALIGN

from designer import edit, store, voice_batch
from designer.api.schemas import DeckVariantState
from designer.contracts import Element, ElementPosition, TextStyle
from designer.edit_runtime import lock
from test_voice_elements_v2 import deck, element, export, saved  # noqa: F401


def url(deck, number=1):
    return f"/decks/{deck.id}/a/slides/{number}/voice-batch"


def payload(raw, operations=None, **changes):
    return {"expected_revision": raw.get("revision") or "",
            "target_slide_id": raw["scenes"][0]["slide_id"],
            "operations": operations if operations is not None else [
                {"element_id": "title", "style": {"bold": True}}], **changes}


def disk(deck):
    root = store.deck_variant_dir(deck.id, "a")
    return {str(path.relative_to(root)): path.read_bytes() for path in root.rglob("*") if path.is_file()}


def assert_rejected(deck, request, status=422, number=1):
    before, scheduled = disk(deck), list(deck.scheduled)
    response = deck.client.post(url(deck, number), json=request)
    assert response.status_code == status, response.text
    assert disk(deck) == before
    assert deck.scheduled == scheduled


def test_batch_persists_exports_once_and_one_undo_restores_everything(deck, monkeypatch):
    before = saved(deck)
    spies = {}
    for name in ("_save_element", "_checks", "_snapshot", "_persist"):
        spies[name] = Mock(wraps=getattr(edit, name))
        monkeypatch.setattr(edit, name, spies[name])
    for name in ("element_action", "move_element", "set_slide_text"):
        monkeypatch.setattr(edit, name, Mock(side_effect=AssertionError("Per-element writes are not atomic")))
    request = payload(before, [
        {"element_id": "title", "box": [.2, .1, .6, .15],
         "style": {"size_pt": 32, "bold": True, "italic": True, "color": "aabbcc", "align": "right"}},
        {"element_id": "shape", "box": [.3, .5, .4, .25], "fill": "abcdef"},
    ])
    response = deck.client.post(url(deck), json=request)
    assert response.status_code == 200, response.text
    after = saved(deck)
    assert response.json() == DeckVariantState(**after).model_dump(mode="json")
    assert after["revision"] and after["revision"] != before.get("revision")
    title = element(after, "title")
    assert title["box"] == [.2, .1, .6, .15]
    assert title["text"] == "Original title"
    assert title["style"] == dict(family="Arial", size_pt=32, bold=True, italic=True,
                                  color="AABBCC", align="right")
    assert element(after, "shape")["fill"] == "ABCDEF"
    positions = after["specs"][0]["element_positions"]
    assert positions["title"]["original_box"] == element(before, "title")["box"]
    assert positions["title"]["style"] == title["style"]
    assert positions["shape"]["box"] == [.3, .5, .4, .25]
    assert positions["shape"]["fill"] == "ABCDEF"
    assert spies["_save_element"].call_count == 2
    for name in ("_checks", "_snapshot", "_persist"):
        spies[name].assert_called_once()
    assert spies["_checks"].call_args.args[1] == {"s1"}
    assert len(deck.scheduled) == 1
    history = list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))
    assert len(history) == 1
    assert store.read_json_file(history[0]) == before
    assert_rejected(deck, request, status=409)
    assert spies["_persist"].call_count == 1

    prs, markup = export(deck)
    native = next(shape for shape in prs.slides[0].shapes if shape.has_text_frame and shape.text == title["text"])
    assert native.left / prs.slide_width == pytest.approx(.2)
    assert native.top / prs.slide_height == pytest.approx(.1)
    assert native.width / prs.slide_width == pytest.approx(.6)
    assert native.height / prs.slide_height == pytest.approx(.15)
    paragraph = native.text_frame.paragraphs[0]
    assert paragraph.alignment == PP_ALIGN.RIGHT
    assert paragraph.runs[0].font.bold and paragraph.runs[0].font.italic
    assert paragraph.runs[0].font.size.pt == 32
    assert str(paragraph.runs[0].font.color.rgb) == "AABBCC"
    shape = next(item for item in prs.slides[0].shapes if item.shape_id == element(after, "shape")["source_shape_id"])
    assert str(shape.fill.fore_color.rgb) == "ABCDEF"
    assert shape.left / prs.slide_width == pytest.approx(.3)
    assert shape.height / prs.slide_height == pytest.approx(.25)
    assert "font-style:italic;" in markup and "background:#ABCDEF;" in markup

    edit.revert(deck.id, "a")
    undone = saved(deck)
    assert {key: value for key, value in undone.items() if key != "revision"} == before
    assert undone["revision"] != after["revision"]
    assert not list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))
    assert len(deck.scheduled) == 2


@pytest.mark.parametrize("second", [
    {"element_id": "missing", "box": [0, 0, .2, .2]},
    {"element_id": "shape", "style": {"bold": True}},
    {"element_id": "shape", "box": [.9, 0, .2, .2]},
    {"element_id": "title", "fill": "abcdef"},
    {"element_id": "title", "style": {"italic": True}},
])
def test_invalid_second_operation_leaves_disk_history_and_exports_untouched(deck, second):
    # Keep an existing undo entry too, so rejection cannot clear or replace it.
    edit.element_action(deck.id, "a", 1, "style", "title", italic=True)
    assert_rejected(deck, payload(saved(deck), [
        {"element_id": "title", "style": {"bold": True}}, second,
    ]))


@pytest.mark.parametrize(("kind", "properties"), [
    ("shape", {"style": {"size_pt": 24}}), ("image", {"style": {"bold": False}}),
    ("icon", {"style": {"align": "left"}}), ("chart", {"style": {"italic": True}}),
    ("table", {"style": {"color": "123456"}}), ("text", {"fill": "abcdef"}),
    ("image", {"fill": "abcdef"}), ("icon", {"fill": "abcdef"}),
    ("chart", {"fill": "abcdef"}), ("table", {"fill": "abcdef"}),
])
def test_property_types_are_checked_for_every_target(deck, kind, properties):
    raw = saved(deck)
    raw["scenes"][0]["elements"].append(Element(id="other", type=kind, box=(.5, .5, .2, .2)).model_dump(mode="json"))
    store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    assert_rejected(deck, payload(raw, [{"element_id": "title", "style": {"bold": True}},
                                      {"element_id": "other", **properties}]))


@pytest.mark.parametrize(("number", "changes"), [
    (1, {"expected_revision": "stale"}), (1, {"target_slide_id": "another-slide"}),
    (0, {}), (-1, {}), (2, {}),
])
def test_stale_revision_or_slide_never_mutates(deck, number, changes):
    assert_rejected(deck, payload(saved(deck), **changes), status=409, number=number)


def test_moved_slide_rejects_stable_id_even_with_current_revision(deck):
    request = payload(saved(deck))
    edit.apply_slide_action(deck.id, "a", "copy", 1)
    edit.apply_slide_action(deck.id, "a", "move", 1, 2)
    request["expected_revision"] = saved(deck)["revision"]
    assert_rejected(deck, request, status=409)


@pytest.mark.parametrize("missing", ["expected_revision", "target_slide_id", "operations"])
def test_required_fields(deck, missing):
    request = payload(saved(deck))
    del request[missing]
    assert_rejected(deck, request)


@pytest.mark.parametrize("changes", [
    {"expected_revision": None}, {"expected_revision": 0}, {"expected_revision": True},
    {"target_slide_id": None}, {"target_slide_id": 1}, {"target_slide_id": ""},
    {"target_slide_id": " \n\t"}, {"target_slide_id": "s" * 201},
    {"operations": []}, {"operations": None}, {"operations": {}},
    {"operations": [{"element_id": f"id-{i}", "fill": "abcdef"} for i in range(65)]},
    {"instruction": "delete the slide"}, {"text": "literal text is not a batch operation"},
])
def test_strict_batch_schema(deck, changes):
    assert_rejected(deck, {**payload(saved(deck)), **changes})


@pytest.mark.parametrize("operation", [
    {}, {"element_id": "title"}, {"element_id": "title", "style": {}},
    {"element_id": "title", "style": None}, {"element_id": "title", "box": None},
    {"element_id": "shape", "fill": None}, {"element_id": None, "style": {"bold": True}},
    {"element_id": 1, "style": {"bold": True}}, {"element_id": "", "style": {"bold": True}},
    {"element_id": " \t", "style": {"bold": True}}, {"element_id": "x" * 201, "style": {"bold": True}},
    {"element_id": "title", "box": [0, 0, .2]}, {"element_id": "title", "box": [0, 0, .2, .2, 0]},
    {"element_id": "title", "box": [True, 0, .2, .2]}, {"element_id": "title", "box": ["0", 0, .2, .2]},
    {"element_id": "title", "box": [None, 0, .2, .2]}, {"element_id": "title", "box": [-.1, 0, .2, .2]},
    {"element_id": "title", "box": [0, 0, 0, .2]}, {"element_id": "title", "box": [0, 0, .2, 0]},
    {"element_id": "title", "box": [0, .9, .2, .2]}, {"element_id": "title", "box": [0, 0, 1.1, .2]},
    {"element_id": "title", "style": {"bold": "true"}}, {"element_id": "title", "style": {"bold": 1}},
    {"element_id": "title", "style": {"italic": "false"}}, {"element_id": "title", "style": {"italic": 0}},
    {"element_id": "title", "style": {"bold": None}}, {"element_id": "title", "style": {"size_pt": None}},
    {"element_id": "title", "style": {"italic": None}}, {"element_id": "title", "style": {"color": None}},
    {"element_id": "title", "style": {"bold": True, "italic": None}},
    {"element_id": "title", "box": None, "style": {"bold": True}},
    {"element_id": "shape", "fill": None, "box": [0, 0, .2, .2]},
    {"element_id": "title", "style": {"size_pt": True}}, {"element_id": "title", "style": {"size_pt": "24"}},
    {"element_id": "title", "style": {"size_pt": 5.99}}, {"element_id": "title", "style": {"size_pt": 144.01}},
    {"element_id": "title", "style": {"color": "#ffffff"}}, {"element_id": "title", "style": {"color": "red"}},
    {"element_id": "title", "style": {"color": 123456}}, {"element_id": "title", "style": {"color": "abcdeg"}},
    {"element_id": "title", "style": {"align": "justify"}}, {"element_id": "title", "style": {"align": "CENTER"}},
    {"element_id": "title", "style": {"align": None}}, {"element_id": "shape", "fill": "#abcdef"},
    {"element_id": "shape", "fill": "ABC"}, {"element_id": "shape", "fill": 123456},
    {"element_id": "shape", "fill": "abcdef\n"},
    {"element_id": "title", "style": {"family": "Other"}},
    {"element_id": "title", "style": {"text": "Delete slide"}},
    {"element_id": "title", "style": {"bold": True}, "text": "finish and cancel"},
    {"element_id": "title", "action": "delete", "style": {"bold": True}},
    {"element_id": "title", "script": "arbitrary()", "style": {"bold": True}},
])
def test_strict_operation_allowlist(deck, operation):
    assert_rejected(deck, payload(saved(deck), [operation]))


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -float("inf")])
@pytest.mark.parametrize("field", ["box", "size_pt"])
def test_nonfinite_numbers_are_rejected_without_writes(deck, value, field):
    operation = {"element_id": "title", **({"box": [value, 0, .2, .2]} if field == "box"
                                            else {"style": {"size_pt": value}})}
    before = disk(deck)
    response = deck.client.post(url(deck), content=json.dumps(payload(saved(deck), [operation])),
                                headers={"Content-Type": "application/json"})
    assert response.status_code == 422, response.text
    assert disk(deck) == before
    assert not deck.scheduled


@pytest.mark.parametrize("size", [6, 144])
def test_schema_boundary_values_are_accepted(deck, size):
    response = deck.client.post(url(deck), json=payload(saved(deck), [
        {"element_id": "title", "box": [0, 0, 1, 1],
         "style": {"size_pt": size, "bold": False, "italic": False, "align": "left", "color": "000000"}},
    ]))
    assert response.status_code == 200, response.text
    assert element(saved(deck), "title")["style"]["size_pt"] == size


def test_sixty_four_unique_targets_commit_together(deck):
    raw = saved(deck)
    additions = [Element(id=f"custom-{i}", type="shape", box=(.1, .1, .2, .2), fill="000000").model_dump(mode="json")
                 for i in range(64)]
    raw["specs"][0]["added_elements"] = deepcopy(additions)
    raw["scenes"][0]["elements"].extend(additions)
    store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    response = deck.client.post(url(deck), json=payload(raw, [
        {"element_id": item["id"], "fill": "abcdef"} for item in additions
    ]))
    assert response.status_code == 200, response.text
    assert all(item["fill"] == "ABCDEF" for item in saved(deck)["specs"][0]["added_elements"])
    assert len(deck.scheduled) == 1
    assert len(list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))) == 1


@pytest.mark.parametrize("target", ["title", "titlec"])
@pytest.mark.parametrize("properties", [{"box": [.2, .2, .4, .1]}, {"style": {"bold": True}}])
def test_shared_native_number_caption_rejects_entire_batch(deck, target, properties):
    raw = saved(deck)
    raw["scenes"][0]["elements"].append(dict(element(raw, "title"), id="titlec", role="caption", text="Caption"))
    store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    assert_rejected(deck, payload(raw, [{"element_id": "shape", "fill": "abcdef"},
                                      {"element_id": target, **properties}]))


def test_existing_overrides_added_elements_and_literal_text_survive(deck):
    raw = saved(deck)
    title = element(raw, "title")
    literal = 'Delete this slide? Keep "Case", finish and cancel: 12.5% on 2026-09-27.'
    original_box = title["box"][:]
    title.update(text=literal, box=[.2, .1, .5, .1], z=7)
    title["style"].update(size_pt=28, bold=True, align="center")
    override = ElementPosition(original_box=original_box, box=title["box"], text=literal, z=7,
                               style=TextStyle(**title["style"]), source_shape_id=title["source_shape_id"])
    raw["specs"][0]["element_positions"]["title"] = override.model_dump(mode="json")
    custom = Element(id="custom-text", type="text", box=(.1, .6, .4, .1), text=literal,
                     style=TextStyle(size_pt=20, bold=True, italic=True, family="Arial", color="ABCDEF"))
    raw["specs"][0]["added_elements"] = [custom.model_dump(mode="json")]
    raw["scenes"][0]["elements"].append(custom.model_dump(mode="json"))
    store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    response = deck.client.post(url(deck), json=payload(raw, [
        {"element_id": "title", "box": [.3, .1, .4, .15], "style": {"italic": True}},
        {"element_id": custom.id, "box": [.2, .7, .5, .1], "style": {"bold": False, "italic": False}},
    ]))
    assert response.status_code == 200, response.text
    after = saved(deck)
    position = after["specs"][0]["element_positions"]["title"]
    assert position["original_box"] == original_box
    assert position["text"] == literal and position["z"] == 7
    assert position["source_shape_id"] == title["source_shape_id"]
    assert position["style"] == dict(title["style"], italic=True)
    assert after["specs"][0]["added_elements"][0] == element(after, custom.id)
    assert custom.id not in after["specs"][0]["element_positions"]
    assert element(after, custom.id)["style"] == dict(custom.style.model_dump(), bold=False, italic=False)
    assert element(after, "title")["text"] == element(after, custom.id)["text"] == literal
    assert after["plan"] == raw["plan"]
    edit.revert(deck.id, "a")
    for key in ("scenes", "specs", "plan", "findings"):
        assert saved(deck)[key] == raw[key]


@pytest.mark.parametrize("style_absent", [False, True])
def test_no_op_preserves_revision_bytes_history_and_scheduling(deck, monkeypatch, style_absent):
    edit.element_action(deck.id, "a", 1, "style", "title", italic=True)
    raw = saved(deck)
    if style_absent:
        element(raw, "title")["style"] = None
        raw["specs"][0]["element_positions"]["title"]["style"] = None
        store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    before, scheduled = disk(deck), list(deck.scheduled)
    for name in ("_save_element", "_checks", "_snapshot", "_persist"):
        monkeypatch.setattr(edit, name, Mock(side_effect=AssertionError("No-op must not save or export")))
    request = payload(raw, [
        {"element_id": "title", "box": element(raw, "title")["box"],
         "style": {"bold": False, "italic": not style_absent}},
        {"element_id": "shape", "fill": "d9d9d9"},
    ])
    for _ in range(2):
        response = deck.client.post(url(deck), json=request)
        assert response.status_code == 200, response.text
        assert response.json() == DeckVariantState(**raw).model_dump(mode="json")
    assert disk(deck) == before
    assert deck.scheduled == scheduled
    assert_rejected(deck, dict(request, expected_revision="stale"), status=409)


def test_mixed_no_op_and_change_has_one_snapshot(deck, monkeypatch):
    spy = Mock(wraps=edit._save_element)
    monkeypatch.setattr(edit, "_save_element", spy)
    response = deck.client.post(url(deck), json=payload(saved(deck), [
        {"element_id": "title", "style": {"size_pt": 24, "bold": False}},
        {"element_id": "shape", "fill": "abcdef"},
    ]))
    assert response.status_code == 200, response.text
    spy.assert_called_once()
    assert set(saved(deck)["specs"][0]["element_positions"]) == {"shape"}
    assert len(deck.scheduled) == 1
    assert len(list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))) == 1


def test_failed_checks_do_not_create_history_or_save(deck, monkeypatch):
    monkeypatch.setattr(edit, "_checks", Mock(side_effect=ValueError("Cannot validate batch")))
    assert_rejected(deck, payload(saved(deck)))


@pytest.mark.parametrize("legacy", ["missing", None, ""])
def test_legacy_revision_is_accepted_once(deck, legacy):
    raw = saved(deck)
    if legacy != "missing":
        raw["revision"] = legacy
        store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
    request = payload(raw)
    response = deck.client.post(url(deck), json=request)
    assert response.status_code == 200, response.text
    assert response.json()["revision"]
    assert_rejected(deck, request, status=409)


@pytest.mark.parametrize("guard", ["revision", "slide_id"])
def test_preconditions_are_loaded_after_acquiring_shared_edit_lock(deck, monkeypatch, guard):
    raw = saved(deck)
    request = payload(raw)
    waiting = Event()

    @contextmanager
    def observed_lock(deck_id, variant):
        waiting.set()
        with lock(deck_id, variant):
            yield

    monkeypatch.setattr(voice_batch, "lock", observed_lock)
    with ThreadPoolExecutor(max_workers=1) as pool:
        with lock(deck.id, "a"):
            pending = pool.submit(deck.client.post, url(deck), json=request)
            assert waiting.wait(5), "Request never reached the shared edit lock"
            if guard == "revision":
                raw["revision"] = "concurrent-edit"
            else:
                raw["scenes"][0]["slide_id"] = "replacement-slide"
                raw["specs"][0]["slide_id"] = "replacement-slide"
                raw["plan"]["slides"][0]["id"] = "replacement-slide"
            store._write_json(store.deck_variant_state_path(deck.id, "a"), raw)
            before = disk(deck)
        response = pending.result(timeout=5)
    assert response.status_code == 409, response.text
    assert disk(deck) == before
    assert not deck.scheduled


def test_concurrent_batches_with_same_revision_commit_only_once(deck):
    request = payload(saved(deck))
    ready = Barrier(2)

    def submit():
        ready.wait(timeout=5)
        return deck.client.post(url(deck), json=request).status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(submit) for _ in range(2)]
        assert sorted(future.result(timeout=10) for future in futures) == [200, 409]
    assert len(deck.scheduled) == 1
    assert len(list(store.deck_variant_history_dir(deck.id, "a").glob("*.json"))) == 1


def test_missing_deck_returns_404(deck):
    request = payload(saved(deck))
    response = deck.client.post("/decks/not-found/a/slides/1/voice-batch", json=request)
    assert response.status_code == 404, response.text
    assert not deck.scheduled
