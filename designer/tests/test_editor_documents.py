import io
import copy
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from fastapi import FastAPI
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR

from designer import editor_documents as docs
from designer.api.editor_documents import router
from designer.export.editor_pptx import export_document


def document():
    return docs.Document.model_validate({'pages': [{'id': 'page', 'background': '#ffffff', 'notes': 'Speaker notes', 'components': [
        {'id': 'title', 'kind': 'title', 'text': 'Launch', 'x': 60, 'y': 40, 'width': 1000, 'height': 80},
        {'id': 'table', 'kind': 'table', 'rows': [['Item', 'Value'], ['A', '42']], 'x': 60, 'y': 200, 'width': 500, 'height': 300},
        {'id': 'chart', 'kind': 'chart', 'labels': ['A', 'B'], 'values': [10, 20], 'x': 600, 'y': 200, 'width': 500, 'height': 300},
    ]}]})


def test_revisions_idempotency_and_reload(tmp_path, monkeypatch):
    monkeypatch.setenv('DESIGNER_DATA_DIR', str(tmp_path))
    doc = document()
    assert docs.save('doc', 'one', 0, doc, 'Title')['revision'] == 1
    assert docs.save('doc', 'one', 0, doc, 'Title')['revision'] == 1
    assert docs.load('doc')['document']['pages'][0]['notes'] == 'Speaker notes'
    with pytest.raises(HTTPException) as error:
        docs.save('doc', 'two', 0, doc, 'Title')
    assert error.value.status_code == 409
    with pytest.raises(HTTPException):
        docs.save('doc', 'one', 0, doc, 'Different')
    assert docs.save('doc', 'two', 1, doc, 'Updated')['revision'] == 2


def test_concurrent_writers_cannot_lose_an_update(tmp_path, monkeypatch):
    monkeypatch.setenv('DESIGNER_DATA_DIR', str(tmp_path))
    docs.save('doc', 'init', 0, document(), 'Title')
    def writer(i):
        try:
            return docs.save('doc', str(i), 1, document(), str(i))['revision']
        except HTTPException as e:
            return e.status_code
    with ThreadPoolExecutor(2) as pool:
        assert sorted(pool.map(writer, [1, 2])) == [2, 409]


def test_first_open_and_first_save_can_initialize_concurrently(tmp_path, monkeypatch):
    monkeypatch.setenv('DESIGNER_DATA_DIR', str(tmp_path))
    def open_or_save(i):
        if i % 2:
            return docs.save('doc'+str(i), 'init', 0, document(), 'Title')
        with docs.database() as db:
            return db.execute('SELECT COUNT(*) FROM documents').fetchone()[0]
    with ThreadPoolExecutor(8) as pool:
        assert len(list(pool.map(open_or_save, range(16)))) == 16


def test_editable_pptx_contains_native_text_table_chart_and_notes():
    deck = Presentation(io.BytesIO(export_document(document())))
    slide = deck.slides[0]
    assert slide.shapes[0].text == 'Launch'
    assert slide.shapes[0].text_frame.vertical_anchor == MSO_ANCHOR.TOP
    assert all(e.get('idx') == '0' for e in slide.shapes[0]._element.xpath('.//a:effectRef'))
    assert slide.shapes[1].has_table
    assert slide.shapes[1].table.cell(1, 1).text == '42'
    assert slide.shapes[2].has_chart
    assert list(slide.shapes[2].chart.series[0].values) == [10, 20]
    assert 'Speaker notes' in slide.notes_slide.notes_text_frame.text


def test_invalid_documents_and_private_image_urls_rejected():
    for url in ['http://127.0.0.1/secret', 'file:///secret', 'javascript:alert(1)', '/editor/assets/../../secret']:
        assert not docs.safe_image_url(url)
    raw = document().model_dump(exclude_none=True)
    raw['pages'][0]['components'][1]['rows'] = [['A'], ['B', 'C']]
    with pytest.raises(ValueError):
        docs.Document.model_validate(raw)
    raw = document().model_dump(exclude_none=True)
    raw['pages'][0]['components'][0]['x'] = float('nan')
    with pytest.raises(ValueError):
        docs.Document.model_validate(raw)


def test_http_save_validate_and_export(tmp_path, monkeypatch):
    monkeypatch.setenv('DESIGNER_DATA_DIR', str(tmp_path))
    app = FastAPI();app.include_router(router)
    client = TestClient(app)
    doc = document().model_dump(exclude_none=True)
    assert client.post('/editor/documents/validate', json=doc).status_code == 200
    saved = client.put('/editor/documents/test', json={'request_id': 'one', 'base_revision': 0, 'document': doc})
    assert saved.status_code == 200
    assert client.get('/editor/documents/test').json()['revision'] == 1
    assert client.post('/editor/export/pptx', json=doc).content.startswith(b'PK')
