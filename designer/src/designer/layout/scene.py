"""Готовая геометрия слайда: фигуры образца с новым текстом и переложенными блоками.

Владелец: задача T-07. Фигуры читаются из исходника, который лежит в пакете дизайн-системы:
оформление переносится как есть, слоты получают текст инструкции, блоки встают по сетке шаблона.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from designer.contracts import (
    Box,
    DesignSystem,
    Element,
    Pattern,
    RepeatGroup,
    Scene,
    SlideSpec,
    Slot,
    TextStyle,
)
from designer.layout.capacity import (
    LINE_HEIGHT,
    number_caption,
    primary_group,
    slide_pt,
    slot_size,
    split_number,
    unit_boxes,
    unit_count,
)
from designer.layout.units import map_shape_box
from designer.parse import geometry as geo
from designer.parse import patterns as parse
from designer.parse.package import SOURCE_NAME

_KEEP_SIZE = ("image", "icon")
_SAMPLE_VIZ = ("chart", "table")
_CACHE_LIMIT = 2

_cache: dict[tuple[str, int], Presentation] = {}


def _presentation(path: Path) -> Presentation:
    """Исходник пакета. Держим разобранным: сцены одной колоды идут с одного файла."""
    key = (str(path), path.stat().st_mtime_ns)
    if key not in _cache:
        if len(_cache) >= _CACHE_LIMIT:
            _cache.pop(next(iter(_cache)))
        _cache[key] = Presentation(str(path))
    return _cache[key]


def _raw_shapes(container, out: dict | None = None) -> dict:
    """Сами фигуры слайда с раскрытием групп: за ними идём только ради заливки и картинки."""
    if out is None:
        out = {}
    for shape in container:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _raw_shapes(shape.shapes, out)
        else:
            out[shape.shape_id] = shape
    return out


def _clip(box: Box) -> Box | None:
    """Рамка, обрезанная по краю слайда; None, если от неё ничего не осталось."""
    x = min(max(box[0], 0.0), 1.0)
    y = min(max(box[1], 0.0), 1.0)
    w = min(box[0] + box[2], 1.0) - x
    h = min(box[1] + box[3], 1.0) - y
    if w <= 0 or h <= 0:
        return None
    return (x, y, w, h)


def _fill_of(shape, part, theme) -> str | None:
    if shape is None:
        return None
    return parse._fill_to_color(shape._element.find(qn("p:spPr")), part, theme)[0]


def _asset_of(shape) -> str | None:
    if shape is None:
        return None
    try:
        return shape.image.sha1[:12]
    except Exception:
        return None


def _text_element(
    el_id: str,
    role: str,
    box: Box,
    text: str,
    style: TextStyle,
    z: int,
    shape_id: int | None,
    ds: DesignSystem,
    slide: tuple[float, float],
    size_pt: float | None = None,
) -> Element | None:
    """Текстовый элемент с кеглем, который выбрала вёрстка; без него кегль считается тут же."""
    clipped = _clip(box)
    if clipped is None or not text:
        return None
    size = style.size_pt or 0.0
    if size_pt:
        fitted = size_pt
    else:
        fitted = slot_size(text, clipped, size, role, ds.tokens.type_scale, slide) if size else size
    return Element(
        id=el_id,
        type="text",
        role=role,
        box=clipped,
        z=z,
        text=text,
        style=style.model_copy(update={"size_pt": fitted or style.size_pt}),
        source_shape_id=shape_id,
    )


def _picture_element(el_id: str, role: str, box: Box, z: int, shape, shape_id: int | None) -> Element:
    return Element(
        id=el_id,
        type="icon" if role == "icon" else "image",
        role=role,
        box=box,
        z=z,
        asset=_asset_of(shape),
        source_shape_id=shape_id,
    )


def _decor_element(el_id: str, info, box: Box, z: int, shape, part, theme) -> Element:
    if info.kind == "image":
        return _picture_element(el_id, "decor", box, z, shape, info.shape_id)
    return Element(
        id=el_id,
        type="shape",
        role="decor",
        box=box,
        z=z,
        fill=_fill_of(shape, part, theme),
        source_shape_id=info.shape_id,
    )


def _viz_box(spec: SlideSpec, pattern: Pattern, group: RepeatGroup | None, ds: DesignSystem) -> Box | None:
    """Где встанет диаграмма или таблица. Рамку выбирает вёрстка, здесь её только исполняют."""
    if spec.viz_box is not None:
        return spec.viz_box
    target = spec.viz_area_id
    if target:
        for area in pattern.areas:
            if area.id == target:
                return area.box
        for slot in pattern.slots:
            if slot.id == target:
                return slot.box
        if group is not None and group.id == target:
            return geo.union([unit.box for unit in group.units])
    if group is not None:
        return geo.union([unit.box for unit in group.units])
    m = ds.tokens.margins
    return (m.left, 0.35, max(1 - m.left - m.right, 0.1), max(0.6 - m.bottom, 0.1))


def _unit_elements(
    group: RepeatGroup,
    unit_text: list[dict[str, str]],
    boxes: list[Box],
    spec: SlideSpec,
    infos: dict,
    raw: dict,
    part,
    theme,
    order: dict,
    ds: DesignSystem,
    slide: tuple[float, float],
) -> list[Element]:
    """Блоки группы по готовым рамкам: лишние блоки образца в сцену не попадают."""
    old = group.units[0].box
    slots = {slot.shape_id: slot for slot in group.unit_slots}
    areas = {area.shape_id: area for area in group.unit_areas}
    removed = set(spec.remove_shape_ids)
    out: list[Element] = []
    for index, new_box in enumerate(boxes):
        texts = unit_text[index] if index < len(unit_text) else {}
        for shape_id in group.units[0].shape_ids:
            info = infos.get(shape_id)
            if info is None or shape_id in removed:
                continue
            slot = slots.get(shape_id)
            area = areas.get(shape_id)
            if area is not None and area.kind in _SAMPLE_VIZ:
                continue
            keep = area.kind in _KEEP_SIZE if area is not None else info.kind == "image"
            box = _clip(map_shape_box(old, new_box, info.box, keep_size=keep))
            if box is None:
                continue
            el_id = f"u{index}s{shape_id}"
            z = order.get(shape_id, 0)
            if slot is not None:
                element = _text_element(
                    el_id, slot.role, box, texts.get(slot.id, ""), slot.style, z, shape_id, ds,
                    slide, spec.fitted_size_pt.get(slot.id),
                )
                if element is not None:
                    out.append(element)
            elif area is not None:
                out.append(_picture_element(el_id, area.kind, box, z, raw.get(shape_id), shape_id))
            else:
                out.append(_decor_element(el_id, info, box, z, raw.get(shape_id), part, theme))
    return out


def _captioned_number(
    slot: Slot, text: str, z: int, ds: DesignSystem, slide: tuple[float, float],
    size_pt: float | None, caption_pt: float | None,
) -> list[Element]:
    """Число и подпись из одной фигуры образца: число строкой сверху, подпись под ним своим кеглем."""
    number, caption = split_number(text)
    size = size_pt or slot.style.size_pt or 0.0
    x, y, w, h = slot.box
    head = min(h, size * LINE_HEIGHT / slide[1]) if size else h / 2
    out = [
        _text_element(slot.id, slot.role, (x, y, w, head), number, slot.style, z, slot.shape_id, ds, slide, size_pt),
        _text_element(
            f"{slot.id}c", "caption", (x, y + head, w, h - head), caption,
            slot.style.model_copy(update={"size_pt": caption_pt or slot.style.size_pt, "bold": False}),
            z, slot.shape_id, ds, slide, caption_pt,
        ),
    ]
    return [element for element in out if element is not None]


def build_scene(spec: SlideSpec, pattern: Pattern, ds: DesignSystem, package_dir: Path) -> Scene:
    """Сцена слайда: каждый объект с окончательной рамкой, текстом и ссылкой на фигуру образца."""
    pres = _presentation(Path(package_dir) / SOURCE_NAME)
    slides = list(pres.slides)
    if not 1 <= pattern.source_slide <= len(slides):
        raise ValueError(f"в исходнике пакета нет слайда {pattern.source_slide}")
    slide = slides[pattern.source_slide - 1]
    size_emu = (pres.slide_width or 1, pres.slide_height or 1)
    theme = parse._theme_of(slide.slide_layout.slide_master)
    infos = {info.shape_id: info for info in parse.flatten_shapes(slide.shapes, size_emu, theme)}
    raw = _raw_shapes(slide.shapes)
    order = {shape_id: i for i, shape_id in enumerate(infos)}
    slide_size = slide_pt(ds.slide_size_emu)

    removed = set(spec.remove_shape_ids)
    group = next((g for g in pattern.groups if g.id == spec.group_id), None)
    if group is None and not spec.group_id:
        group = primary_group(pattern)
    linked = [g for g in pattern.groups if g.id in spec.linked_unit_text]
    has_viz = spec.chart is not None or spec.table is not None
    viz_id = spec.viz_area_id if has_viz else None
    viz = _clip(_viz_box(spec, pattern, group, ds)) if has_viz else None
    elements: list[Element] = []

    for shape_id in pattern.decor_shape_ids:
        info = infos.get(shape_id)
        if info is None or shape_id in removed:
            continue
        box = _clip(info.box)
        if box is None:
            continue
        shape = raw.get(shape_id)
        # Подложка на весь кадр уже лежит фоном сцены, вторым слоем она не нужна.
        if info.kind == "image" and _asset_of(shape) == pattern.background_asset:
            if box[2] >= geo.FULL_BLEED and box[3] >= geo.FULL_BLEED:
                continue
        elements.append(
            _decor_element(f"d{shape_id}", info, box, order.get(shape_id, 0), shape, slide.part, theme)
        )

    for slot in pattern.slots:
        if slot.id == viz_id or slot.shape_id in removed:
            continue
        text = spec.slot_text.get(slot.id, "")
        if number_caption(slot) and split_number(text)[1]:
            info = infos.get(slot.shape_id)
            elements.extend(_captioned_number(
                slot, text, order.get(slot.shape_id, 0), ds, slide_size,
                spec.fitted_size_pt.get(slot.id), info.tail_size_pt if info else None,
            ))
            continue
        element = _text_element(
            slot.id,
            slot.role,
            slot.box,
            text,
            slot.style,
            order.get(slot.shape_id, 0),
            slot.shape_id,
            ds,
            slide_size,
            spec.fitted_size_pt.get(slot.id),
        )
        if element is not None:
            elements.append(element)

    for area in pattern.areas:
        if area.id == viz_id or area.kind in _SAMPLE_VIZ or area.shape_id in removed:
            continue
        box = _clip(area.box)
        if box is None:
            continue
        # Место отдано диаграмме или таблице: картинка образца оттуда уходит.
        if viz is not None and geo.area(box) <= geo.area(viz) and geo.overlap(box, viz) > 0.5:
            continue
        elements.append(
            _picture_element(
                area.id, area.kind, box, order.get(area.shape_id or -1, 0), raw.get(area.shape_id), area.shape_id
            )
        )

    if group is not None and spec.unit_text and viz_id != group.id:
        n = unit_count(group, len(spec.unit_text))
        main_boxes, linked_boxes = unit_boxes(group, linked, n)
        for item in (group, *linked):
            boxes = main_boxes if item is group else linked_boxes[item.id]
            texts = spec.unit_text if item is group else spec.linked_unit_text[item.id]
            elements.extend(_unit_elements(
                item, texts, boxes, spec, infos, raw, slide.part, theme, order, ds, slide_size
            ))

    if viz is not None:
        elements.append(Element(
            id=viz_id or "viz",
            type="chart" if spec.chart is not None else "table",
            role="viz",
            box=viz,
            z=max(order.values(), default=0) + 1,
            chart=spec.chart,
            table=spec.table,
        ))

    elements.sort(key=lambda el: el.z)
    return Scene(
        slide_id=spec.slide_id,
        pattern_id=pattern.id,
        variant=spec.variant,
        theme=pattern.theme,
        background_asset=pattern.background_asset,
        background_color=pattern.background_color,
        elements=elements,
    )
