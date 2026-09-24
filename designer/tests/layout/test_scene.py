"""Сборка сцены слайда: фигуры образца, новый текст, переложенные блоки. Задачи T-07, T-19, T-22."""
import re

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from designer.audit.deterministic import run_checks
from designer.contracts import ChartSpec, Item, Series, SlideIntent, SlideKind, TableSpec
from designer.layout.capacity import (
    NUMBER_MARGIN,
    OVERLAP_MIN,
    ink_box,
    slide_pt,
    text_lines,
    title_step,
)
from designer.layout.compose import EDGE_AREA, VIZ_INSIDE, _hint_slots, compose
from designer.layout.match import SAMPLE_MIN, choose_pattern
from designer.layout.scene import build_scene
from designer.parse import geometry as geo
from designer.parse.package import build_package

_UNIT_ID = re.compile(r"^u(\d+)s\d+$")
_SAMPLE_MIN = 8
_VIZ_MIN = 0.35


def _textbox(slide, left, top, width, height, text, size_pt):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(size_pt)
    return box


def _sample_pptx(path):
    """Шаблон из трёх карточек: заголовок слайда и повторяющийся блок с заголовком и текстом."""
    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, 457200, 228600, 6000000, 700000, "Образцовый заголовок этого шаблона", 32)
    for i in range(3):
        left = 457200 + i * 2743200
        slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(1600000), Emu(2400000), Emu(1800000))
        _textbox(slide, left + 150000, 1750000, 2100000, 400000, "Образец пункта шаблона", 20)
        _textbox(slide, left + 150000, 2250000, 2100000, 1000000, "Образцовое описание пункта в карточке", 12)
    prs.save(str(path))
    return path


def _steps_pptx(path):
    """Шаблон из ряда номеров и ряда подписей под ними: две группы одной смысловой строки."""
    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, 457200, 228600, 6000000, 700000, "Образцовый заголовок этого шаблона", 32)
    for i in range(3):
        left = 457200 + i * 2743200
        _textbox(slide, left + 800000, 1500000, 800000, 600000, f"0{i + 1}", 40)
        _textbox(slide, left, 2600000, 2400000, 700000, f"Образцовая подпись шага {i + 1}", 16)
    prs.save(str(path))
    return path


@pytest.fixture(scope="module")
def sample_package(tmp_path_factory):
    root = tmp_path_factory.mktemp("obrazec")
    ds = build_package(_sample_pptx(root / "obrazec.pptx"), root / "paket")
    return ds, root / "paket"


@pytest.fixture(scope="module")
def steps_package(tmp_path_factory):
    root = tmp_path_factory.mktemp("shagi")
    ds = build_package(_steps_pptx(root / "shagi.pptx"), root / "paket")
    return ds, root / "paket"


def _plan() -> list[SlideIntent]:
    """Восемь намерений: титул, карточки, шаги, крупное число, диаграмма, таблица, список, финал."""
    return [
        SlideIntent(id="s1", kind=SlideKind.title, title="Цифровой дизайнер презентаций",
                    key_message="Колода из брифа в стиле вашего шаблона"),
        SlideIntent(id="s2", kind=SlideKind.cards, title="Три причины попробовать",
                    items=[Item(heading="Скорость", body="Колода собирается за минуты"),
                           Item(heading="Стиль", body="Всё берётся из вашего шаблона"),
                           Item(heading="Проверка", body="Аудит ловит наезды и переполнение")]),
        SlideIntent(id="s3", kind=SlideKind.steps, title="Как идёт работа",
                    items=[Item(number="1", heading="Разбор", body="Читаем шаблон"),
                           Item(number="2", heading="План", body="Делим бриф на слайды"),
                           Item(number="3", heading="Вёрстка", body="Кладём слова в слоты"),
                           Item(number="4", heading="Аудит", body="Считаем наезды")]),
        SlideIntent(id="s4", kind=SlideKind.big_number, title="Сколько это занимает",
                    key_message="Столько идёт сборка колоды целиком",
                    items=[Item(number="5", heading="минут на колоду")]),
        SlideIntent(id="s5", kind=SlideKind.chart, title="Как растёт скорость",
                    chart=ChartSpec(type="column", categories=["Руками", "Пополам", "Сервисом"],
                                    series=[Series(name="минут", values=[40, 12, 5])], unit="мин")),
        SlideIntent(id="s6", kind=SlideKind.table, title="Что смотрит аудит",
                    table=TableSpec(columns=["Проверка", "Что ловит", "Итог"],
                                    rows=[["Границы", "выход за кадр", "ошибка"],
                                          ["Наезд", "пересечение блоков", "ошибка"],
                                          ["Плотность", "слишком много пунктов", "замечание"]])),
        SlideIntent(id="s7", kind=SlideKind.bullets, title="Что дальше",
                    items=[Item(body="Три варианта одной колоды"),
                           Item(body="Живой режим на сцене"),
                           Item(body="Починка находок по выбору")]),
        SlideIntent(id="s8", kind=SlideKind.thanks, title="Спасибо за внимание",
                    key_message="Вопросы и связь"),
    ]


def _deck(ds, package_dir):
    """Колода по плану: намерение, выбранный паттерн, инструкция сборки и сцена."""
    used: list[str] = []
    rows = []
    for intent in _plan():
        pattern = choose_pattern(intent, ds, used)
        used.append(pattern.id)
        spec = compose(intent, pattern, ds)
        rows.append((intent, pattern, spec, build_scene(spec, pattern, ds, package_dir)))
    return rows


def _scenes(deck):
    return [scene for _, _, _, scene in deck]


def _samples(ds) -> set[str]:
    texts = {slot.sample_text.strip() for pattern in ds.patterns for slot in pattern.slots}
    texts |= {slot.sample_text.strip()
              for pattern in ds.patterns for group in pattern.groups for slot in group.unit_slots}
    return {text for text in texts if len(text) >= _SAMPLE_MIN}


def _center(box) -> float:
    return box[0] + box[2] / 2


def _unit_overlaps(findings) -> list:
    """Наложения между блоками одной группы: оба элемента из разных блоков."""
    out = []
    for finding in findings:
        if finding.check_id != "layout.overlap" or len(finding.element_ids) != 2:
            continue
        left, right = (_UNIT_ID.match(el) for el in finding.element_ids)
        if left and right and left.group(1) != right.group(1):
            out.append(finding)
    return out


# ---------- шаблон, собранный прямо в тесте ----------

def test_blocks_are_rebuilt_for_a_smaller_count(sample_package):
    ds, package_dir = sample_package
    intent = SlideIntent(id="s1", kind=SlideKind.cards, title="Две причины",
                         items=[Item(heading="Скорость", body="Минуты вместо дней"),
                                Item(heading="Стиль", body="Всё из вашего шаблона")])
    pattern = choose_pattern(intent, ds, [])
    scene = build_scene(compose(intent, pattern, ds), pattern, ds, package_dir)

    units = {el.id for el in scene.elements if _UNIT_ID.match(el.id)}
    assert {_UNIT_ID.match(el).group(1) for el in units} == {"0", "1"}
    cards = [el for el in scene.elements if el.type == "shape" and _UNIT_ID.match(el.id)]
    assert len(cards) == 2
    assert cards[0].box[2] > pattern.groups[0].unit_size[0]
    assert cards[0].box[0] + cards[0].box[2] <= cards[1].box[0] + 1e-6


def test_sample_text_does_not_reach_the_scene(sample_package):
    ds, package_dir = sample_package
    intent = SlideIntent(id="s1", kind=SlideKind.cards, title="Две причины",
                         items=[Item(heading="Скорость", body="Минуты вместо дней"),
                                Item(heading="Стиль", body="Всё из вашего шаблона")])
    pattern = choose_pattern(intent, ds, [])
    scene = build_scene(compose(intent, pattern, ds), pattern, ds, package_dir)

    texts = [el.text for el in scene.elements if el.type == "text"]
    assert "Две причины" in texts
    assert all("Образц" not in text and "Образец" not in text for text in texts)


def test_numbers_stand_on_the_axis_of_their_captions(steps_package):
    ds, package_dir = steps_package
    pattern = ds.patterns[0]
    heads = next(g for g in pattern.groups if any(s.role == "heading" for s in g.unit_slots))
    numbers = next(g for g in pattern.groups if any(s.role == "number" for s in g.unit_slots))
    pattern.primary_group_id = heads.id
    heads.linked_group_ids = [numbers.id]

    intent = SlideIntent(id="s3", kind=SlideKind.steps, title="Как идёт работа",
                         items=[Item(number="1", heading="Разбор"),
                                Item(number="2", heading="План")])
    spec = compose(intent, pattern, ds)
    scene = build_scene(spec, pattern, ds, package_dir)

    digits = sorted((el for el in scene.elements if el.role == "number"), key=lambda el: el.box[0])
    captions = sorted((el for el in scene.elements if el.role == "heading"), key=lambda el: el.box[0])
    assert [el.text for el in digits] == ["1", "2"]
    assert [el.text for el in captions] == ["Разбор", "План"]
    for digit, caption in zip(digits, captions):
        assert abs(_center(digit.box) - _center(caption.box)) <= 0.03


def test_blocks_leave_the_slide_when_it_gets_a_table(sample_package):
    ds, package_dir = sample_package
    intent = SlideIntent(id="s5", kind=SlideKind.table, title="Что смотрит аудит",
                         table=TableSpec(columns=["Проверка", "Что ловит"],
                                         rows=[["Границы", "выход за кадр"],
                                               ["Наезд", "пересечение блоков"]]))
    pattern = choose_pattern(intent, ds, [])
    spec = compose(intent, pattern, ds)
    scene = build_scene(spec, pattern, ds, package_dir)

    assert not [el for el in scene.elements if _UNIT_ID.match(el.id)]
    viz = [el for el in scene.elements if el.type == "table"]
    assert len(viz) == 1 and geo.area(viz[0].box) >= _VIZ_MIN
    under = [el.text for el in scene.elements
             if el.type == "text" and el.text.strip() and geo.overlap(el.box, viz[0].box) > 0]
    assert not under


def test_scene_keeps_the_link_to_the_source_shapes(sample_package):
    ds, package_dir = sample_package
    intent = SlideIntent(id="s1", kind=SlideKind.cards, title="Две причины",
                         items=[Item(heading="Скорость", body="Минуты вместо дней")])
    pattern = choose_pattern(intent, ds, [])
    scene = build_scene(compose(intent, pattern, ds), pattern, ds, package_dir)
    assert scene.pattern_id == pattern.id
    assert all(el.source_shape_id is not None for el in scene.elements)


# ---------- выданные шаблоны и сторонний файл ----------

@pytest.fixture(scope="session")
def packages(templates, foreign_pptx, tmp_path_factory):
    root = tmp_path_factory.mktemp("pakety")
    built = []
    for index, path in enumerate([*templates, foreign_pptx]):
        out = root / f"ds{index}"
        built.append((path.stem, build_package(path, out), out))
    return built


@pytest.fixture(scope="module")
def decks(packages):
    """Колода из восьми слайдов на каждом шаблоне: собирается один раз на весь разбор."""
    return [(name, ds, _deck(ds, package_dir)) for name, ds, package_dir in packages]


def test_every_intent_gives_a_scene_on_every_template(decks):
    for name, _, deck in decks:
        assert len(deck) == len(_plan()), name
        for _, _, _, scene in deck:
            assert scene.elements, f"{name}, слайд {scene.slide_id}"
            assert any(el.type in ("text", "chart", "table") for el in scene.elements), name


def test_no_sample_text_of_the_template_on_every_template(decks):
    for name, ds, deck in decks:
        samples = _samples(ds)
        for scene in _scenes(deck):
            for el in scene.elements:
                if el.type != "text":
                    continue
                found = [sample for sample in samples if sample in el.text]
                assert not found, f"{name}, слайд {scene.slide_id}: {found[:1]}"


def test_audit_finds_no_out_of_bounds_and_no_overlap_inside_groups(decks):
    for name, ds, deck in decks:
        findings = run_checks(_scenes(deck), ds)
        out_of_bounds = [f for f in findings if f.check_id == "layout.out_of_bounds"]
        assert not out_of_bounds, f"{name}: {[f.message for f in out_of_bounds[:3]]}"
        overlaps = _unit_overlaps(findings)
        assert not overlaps, f"{name}: {[f.message for f in overlaps[:3]]}"


# ---------- качество вёрстки на выданных шаблонах ----------

def test_no_slide_stands_on_a_pattern_that_holds_on_photos(decks):
    for name, _, deck in decks:
        for intent, pattern, _, _ in deck:
            assert not pattern.needs_images, f"{name}, слайд {intent.id}"


def test_title_and_final_stand_on_patterns_without_blocks_and_samples(decks):
    for name, _, deck in decks:
        for intent, pattern, _, _ in deck:
            if intent.kind not in (SlideKind.title, SlideKind.thanks):
                continue
            assert not pattern.groups, f"{name}, слайд {intent.id}: блоки образца"
            samples = [area.kind for area in pattern.areas
                       if area.kind in ("chart", "table")
                       or (area.kind == "image" and geo.area(area.box) >= SAMPLE_MIN)]
            assert not samples, f"{name}, слайд {intent.id}: {samples}"


def test_chart_and_table_get_a_big_clean_frame(decks):
    for name, ds, deck in decks:
        for intent, _, spec, scene in deck:
            if intent.chart is None and intent.table is None:
                continue
            box = spec.viz_box
            where = f"{name}, слайд {intent.id}"
            assert box is not None, where
            assert geo.area(box) >= _VIZ_MIN, f"{where}: рамка {geo.area(box):.2f}"
            margins = ds.tokens.margins
            assert box[0] >= margins.left - 1e-6 and box[1] >= margins.top - 1e-6, where
            assert geo.right(box) <= 1 - margins.right + 1e-6, where
            assert geo.bottom(box) <= 1 - margins.bottom + 1e-6, where
            viz = [el for el in scene.elements if el.type in ("chart", "table")]
            assert len(viz) == 1 and geo.overlap(viz[0].box, box) > 0.99, where
            under = [el.text for el in scene.elements
                     if el.type == "text" and el.text.strip() and geo.overlap(el.box, box) > 0]
            assert not under, f"{where}: в рамке текст {under[:2]}"


def test_no_card_loses_its_words(decks):
    for name, _, deck in decks:
        for intent, _, _, scene in deck:
            if intent.kind is not SlideKind.cards:
                continue
            said = " ".join(el.text for el in scene.elements if el.type == "text")
            for item in intent.items:
                assert item.heading in said, f"{name}, слайд {intent.id}: {item.heading}"
                assert item.body in said, f"{name}, слайд {intent.id}: {item.body}"


def test_audit_finds_no_placeholder_text_and_no_overlap_with_the_visualisation(decks):
    for name, ds, deck in decks:
        scenes = _scenes(deck)
        findings = run_checks(scenes, ds)
        stubs = [f for f in findings if f.check_id == "integrity.placeholder_text"]
        assert not stubs, f"{name}: {[f.message for f in stubs[:3]]}"
        viz_ids = {el.id for scene in scenes for el in scene.elements
                   if el.type in ("chart", "table")}
        hits = [f for f in findings
                if f.check_id == "layout.overlap" and viz_ids & set(f.element_ids)]
        assert not hits, f"{name}: {[f.message for f in hits[:3]]}"


# ---------- дефекты живой колоды: длинное слово, заголовок и то, что под ним ----------

def test_long_word_of_a_separator_does_not_break(packages):
    """Разделитель: слово целиком в рамке, кегль не ниже ступени title, переполнения нет."""
    for name, ds, package_dir in packages:
        intent = SlideIntent(id="s1", kind=SlideKind.section, title="Реструктуризация")
        pattern = choose_pattern(intent, ds, [])
        spec = compose(intent, pattern, ds)
        scene = build_scene(spec, pattern, ds, package_dir)
        overflow = [f for f in run_checks([scene], ds) if f.check_id == "layout.text_overflow"]
        assert not overflow, f"{name}: {[f.message for f in overflow[:2]]}"
        said = [el for el in scene.elements if el.type == "text" and el.text == intent.title]
        assert said, f"{name}: заголовка нет на слайде"
        floor = title_step(ds.tokens.type_scale)
        for el in said:
            assert el.style.size_pt >= floor - 1e-6, f"{name}: {el.style.size_pt} вместо {floor}"


def test_no_text_runs_into_other_text_on_every_template(decks):
    """Заголовок не наезжает на подзаголовок, крупное число на подпись под ним."""
    for name, ds, deck in decks:
        scenes = _scenes(deck)
        texts = {(scene.slide_id, el.id)
                 for scene in scenes for el in scene.elements if el.type == "text"}
        hits = [f for f in run_checks(scenes, ds)
                if f.check_id == "layout.overlap"
                and all((f.slide_id, el) in texts for el in f.element_ids)]
        assert not hits, f"{name}: {[f.message for f in hits[:3]]}"


def test_no_block_of_the_deck_stands_empty(decks):
    """Блок без текстового слота на слайде не остаётся: его нечем заполнить."""
    for name, _, deck in decks:
        for intent, pattern, spec, _ in deck:
            for group in pattern.groups:
                if group.unit_slots or group.id in spec.linked_unit_text:
                    continue
                shapes = {sid for unit in group.units for sid in unit.shape_ids}
                left = shapes - set(spec.remove_shape_ids)
                assert not left, f"{name}, слайд {intent.id}: {sorted(left)[:3]}"


# ---------- дефекты живой колоды: число, оформление под диаграммой, немые блоки ----------

def test_big_number_slide_shows_its_number(decks, templates):
    """Крупное число стоит в слоте number и набрано не мельче ступени title."""
    given = {path.stem for path in templates}
    for name, ds, deck in decks:
        if name not in given:
            continue
        for intent, _, _, scene in deck:
            if intent.kind is not SlideKind.big_number:
                continue
            value = next(item.number for item in intent.items if item.number)
            digits = [el for el in scene.elements if el.role == "number" and el.text.strip()]
            assert [el.text for el in digits] == [value], f"{name}, слайд {intent.id}"
            floor = title_step(ds.tokens.type_scale)
            assert digits[0].style.size_pt >= floor - 1e-6, f"{name}: {digits[0].style.size_pt}"


def test_no_decor_of_the_sample_stands_under_the_visualisation(decks):
    """Оформление образца под диаграммой просвечивает сквозь неё; фон и логотип остаются."""
    for name, _, deck in decks:
        for intent, _, spec, scene in deck:
            if spec.viz_box is None:
                continue
            under = [el for el in scene.elements
                     if el.role == "decor" and geo.area(el.box) > EDGE_AREA
                     and not (el.box[2] >= geo.FULL_BLEED and el.box[3] >= geo.FULL_BLEED)
                     and geo.covered(el.box, spec.viz_box) >= VIZ_INSIDE]
            assert not under, f"{name}, слайд {intent.id}: {[el.id for el in under][:3]}"


def test_no_block_of_a_scene_stands_without_words(decks):
    """Блок, которому слов не досталось, уходит со слайда вместе со своими маркерами."""
    for name, _, deck in decks:
        for intent, _, _, scene in deck:
            words: dict[str, bool] = {}
            for el in scene.elements:
                found = _UNIT_ID.match(el.id)
                if found is None:
                    continue
                index = found.group(1)
                words[index] = words.get(index, False) or bool(el.text.strip())
            mute = sorted(index for index, said in words.items() if not said)
            assert not mute, f"{name}, слайд {intent.id}: блоки {mute}"


def test_removed_shapes_do_not_reach_the_scene(decks):
    for name, _, deck in decks:
        for intent, _, spec, scene in deck:
            removed = set(spec.remove_shape_ids)
            if not removed:
                continue
            left = {el.source_shape_id for el in scene.elements} & removed
            assert not left, f"{name}, слайд {intent.id}: {sorted(left)[:3]}"


# ---------- дефекты листа 2026-09-21: фраза целиком, число одной строкой, пустые плашки ----------

LONG_TITLE = "Перевод ночных отчётов на потоковую загрузку"
"""Заголовок в 44 знака: на кегле образца он не влезал и приходил обрезанным до слова."""

VALUES = ("11:00", "12 %", "6 недель")
"""Числа с единицей: время, процент и срок. Каждое стоит в слоте целиком и одной строкой."""

LIST_HEADS = ("Потоки", "Витрины", "Контроль")


@pytest.fixture(scope="session")
def given(templates, tmp_path_factory):
    """Три выданных шаблона, разобранных в пакеты дизайн-системы."""
    root = tmp_path_factory.mktemp("vydannye")
    out = []
    for index, path in enumerate(templates):
        package_dir = root / f"ds{index}"
        out.append((path.stem, build_package(path, package_dir), package_dir))
    return out


def _scene_of(intent, ds, package_dir):
    pattern = choose_pattern(intent, ds, [])
    spec = compose(intent, pattern, ds)
    return pattern, spec, build_scene(spec, pattern, ds, package_dir)


def _texts(scene, role=None):
    return [el for el in scene.elements
            if el.type == "text" and el.text.strip() and (role is None or el.role == role)]


def test_a_long_title_reaches_the_scene_whole(given):
    """Фраза заголовка встаёт целиком: кегль ниже образца или лишняя строка, но не обрезка."""
    intent = SlideIntent(id="s1", kind=SlideKind.title, title=LONG_TITLE,
                         key_message="Предложение о переходе на потоковую загрузку")
    for name, ds, package_dir in given:
        _, _, scene = _scene_of(intent, ds, package_dir)
        heads = [el for el in _texts(scene) if el.text.strip() == LONG_TITLE]
        assert heads, f"{name}: {[el.text for el in _texts(scene)]}"
        head = heads[0]
        size = head.style.size_pt or 0.0
        ink = ink_box(head.box, head.text, size, slide_pt(ds.slide_size_emu))
        assert ink[3] <= head.box[3] + 1e-6, f"{name}: строки вышли из рамки"


def test_a_number_with_its_unit_stands_in_one_line(given):
    """Число вместе с единицей стоит в слоте целиком: «11:00» это не «11», и без переноса."""
    for name, ds, package_dir in given:
        slide = slide_pt(ds.slide_size_emu)
        for value in VALUES:
            intent = SlideIntent(id="s1", kind=SlideKind.big_number,
                                 title="Что показывает замер",
                                 items=[Item(number=value, heading="замер пилота")])
            _, _, scene = _scene_of(intent, ds, package_dir)
            digits = _texts(scene, "number")
            assert [el.text for el in digits] == [value], f"{name}, {value}"
            size = (digits[0].style.size_pt or 0.0) * NUMBER_MARGIN
            assert text_lines(value, digits[0].box, size, slide) == 1, f"{name}, {value}: перенос"


def test_a_number_named_only_in_the_title_keeps_its_unit(given):
    """Число берётся из заголовка целиком: у времени остаются минуты, у доли знак процента."""
    for name, ds, package_dir in given:
        for value, title in (("11:00", "Отчёты готовы только к 11:00 утра"),
                             ("12 %", "Рост затрат на инфраструктуру на 12 %")):
            intent = SlideIntent(id="s1", kind=SlideKind.big_number, title=title,
                                 items=[Item(heading="замер пилота", body="строка пояснения")])
            _, _, scene = _scene_of(intent, ds, package_dir)
            assert [el.text for el in _texts(scene, "number")] == [value], f"{name}, {value}"


def test_no_empty_photo_plate_and_no_empty_icon_in_the_scene(given):
    """Подсказка шаблона «вставьте фото» и её плашка на слайд не попадают: картинок у нас нет."""
    for name, ds, package_dir in given:
        for intent in _plan():
            pattern, spec, scene = _scene_of(intent, ds, package_dir)
            hints = _hint_slots(pattern)
            if hints:
                plates = set(hints.values()) | {s.shape_id for s in pattern.slots if s.id in hints}
                assert plates <= set(spec.remove_shape_ids), f"{name}, слайд {intent.id}"
            left = {el.source_shape_id for el in scene.elements} & set(spec.remove_shape_ids)
            assert not left, f"{name}, слайд {intent.id}: {sorted(left)[:3]}"
            mute = [el.id for el in scene.elements
                    if el.type in ("image", "icon") and el.asset is None]
            assert not mute, f"{name}, слайд {intent.id}: {mute[:3]}"


def test_a_list_of_three_items_shows_all_three_headings(given):
    """Намерение со списком не уходит на паттерн раздела: видны заголовки всех трёх пунктов."""
    intent = SlideIntent(
        id="s2", kind=SlideKind.bullets, title="Что меняем в загрузке данных",
        key_message="Три шага перехода",
        items=[Item(heading=LIST_HEADS[0], body="Данные идут потоком"),
               Item(heading=LIST_HEADS[1], body="Обновление без ночного окна"),
               Item(heading=LIST_HEADS[2], body="Метрики задержки на виду")])
    for name, ds, package_dir in given:
        _, _, scene = _scene_of(intent, ds, package_dir)
        said = " ".join(el.text for el in _texts(scene))
        missing = [head for head in LIST_HEADS if head not in said]
        assert not missing, f"{name}: пропали пункты {missing}"


def test_no_text_leaves_the_frame_or_crosses_the_logo(given):
    """Набранные строки стоят в кадре и не заходят на логотип шаблона."""
    for name, ds, package_dir in given:
        logos = {asset.id for asset in ds.assets if asset.kind == "logo"}
        for intent in _plan():
            _, _, scene = _scene_of(intent, ds, package_dir)
            marks = [el.box for el in scene.elements
                     if el.type in ("image", "icon") and el.asset in logos]
            for el in _texts(scene):
                ink = ink_box(el.box, el.text, el.style.size_pt or 0.0, slide_pt(ds.slide_size_emu))
                assert ink[0] >= -1e-6 and ink[1] >= -1e-6, f"{name}, слайд {intent.id}: {el.id}"
                assert geo.right(ink) <= 1 + 1e-6 and geo.bottom(ink) <= 1 + 1e-6, \
                    f"{name}, слайд {intent.id}: {el.id}"
                hits = [box for box in marks if geo.overlap(ink, box) > OVERLAP_MIN]
                assert not hits, f"{name}, слайд {intent.id}: {el.id} на логотипе"
