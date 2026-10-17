"""Таблица автора: csv или первый лист xlsx в строки и точные числа.

Пределы те же, что у разбора CSV в браузере (`frontend/src/stage/meaningCsv.ts` из ветки Стаса):
1 МиБ, 2000 строк данных, 50 столбцов, 2000 знаков в ячейке. Формулы не выполняются: из xlsx
берётся значение, которое Excel сохранил при последнем пересчёте, строка «=...» в csv остаётся текстом.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

MAX_BYTES = 1024 * 1024
MAX_ROWS = 2000
MAX_COLS = 50
MAX_CELL = 2000


class TableError(ValueError):
    """Файл не годится: сообщение говорит автору, что поправить."""


@dataclass(frozen=True)
class Table:
    name: str
    headers: list[str]
    rows: list[list[str]]
    """Текст ячеек, как его видит автор."""
    numbers: list[list[Decimal | None]]
    """Число в ячейке или None, если ячейка не число."""
    lines: list[int]
    """Номер строки данных в файле, как его видит автор в Excel: пустые строки файла не сдвигают счёт."""
    sheet: str = ""
    """Лист xlsx, с которого прочитана таблица; у csv пусто."""


_NUMBER = re.compile(r"^[+-]?(?:\d{1,3}(?:[   ]\d{3})+|\d+)(?:[.,]\d+)?$")
_MONTH_YEAR = re.compile(r"^\d{1,2}[./]\d{4}$")
"""«01.2026» это месяц и год, а не число 1,2026."""
_CURRENCY = re.compile(r"\s*(?:₽|руб\.?|р\.|\$|€|%)\s*$", re.I)


def parse_number(text: str) -> Decimal | None:
    """«1 234,5», «-12.3», «18 %», «−5», «120 ₽» в Decimal. Формулы, даты и текст не числа."""
    value = text.strip().replace("−", "-")
    if not value or value.startswith("=") or _MONTH_YEAR.match(value):
        return None
    value = _CURRENCY.sub("", value)
    if not _NUMBER.match(value):
        return None
    value = re.sub(r"[   ]", "", value).replace(",", ".")
    try:
        return Decimal(value)
    except InvalidOperation:
        return None


def read_table(data: bytes, filename: str) -> Table:
    if len(data) > MAX_BYTES:
        raise TableError("Файл больше 1 МБ: оставьте в таблице только нужный лист и строки")
    name = (filename or "таблица").replace("\\", "/").rsplit("/", 1)[-1][:200]
    lower = name.lower()
    if lower.endswith(".xlsx"):
        sheet, grid = _read_xlsx(data)
        try:
            return _to_table(name, grid, sheet)
        except TableError as exc:
            raise TableError(f"Лист «{sheet}»: {exc} Сервис читает только первый лист книги.") from exc
    if lower.endswith(".csv") or lower.endswith(".txt"):
        return _to_table(name, _read_csv(data))
    raise TableError("Нужен файл xlsx или csv")


def _read_csv(data: bytes) -> list[list[str]]:
    if b"\0" in data:
        raise TableError("Файл не похож на таблицу csv: сохраните его из Excel заново как CSV.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel на русской Windows сохраняет CSV в cp1251 с разделителем «;».
        try:
            text = data.decode("cp1251")
        except UnicodeDecodeError as exc:
            raise TableError("Не удалось прочитать кодировку: сохраните csv в UTF-8.") from exc
    first_line = text.split("\n", 1)[0]
    delimiter = max((";", ",", "\t"), key=first_line.count)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        return [row for row in reader]
    except csv.Error as exc:
        raise TableError(f"Строка {reader.line_num}: кавычка в значении не закрыта или стоит не на месте. "
                         "Сохраните таблицу из Excel заново как CSV.") from exc


def _read_xlsx(data: bytes) -> tuple[str, list[list[str]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # повреждённый zip, не книга Excel, xls под видом xlsx
        raise TableError("Файл не открывается как книга Excel xlsx") from exc
    try:
        sheet = workbook.worksheets[0]
        grid: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            if len(grid) > MAX_ROWS + 1:
                break
            grid.append([_cell_text(value) for value in row[:MAX_COLS + 1]])
        return sheet.title, grid
    finally:
        workbook.close()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, dt.datetime):
        return value.strftime("%d.%m.%Y") if value.time() == dt.time() else value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, dt.date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        return format(Decimal(repr(value)).normalize(), "f")
    return str(value)


def _to_table(name: str, grid: list[list[str]], sheet: str = "") -> Table:
    grid = [[cell.strip() for cell in row] for row in grid]
    numbered = [(index, row) for index, row in enumerate(grid, start=1) if any(row)]
    grid = [row for _, row in numbered]
    if not grid:
        raise TableError("Таблица пустая.")
    width = max(len(row) for row in grid)
    while width and all(len(row) < width or not row[width - 1] for row in grid):
        width -= 1
    if width > MAX_COLS:
        raise TableError(f"Больше {MAX_COLS} столбцов: оставьте только нужные.")
    grid = [row[:width] + [""] * (width - len(row)) for row in grid]
    headers, body = grid[0], grid[1:]
    if len(body) > MAX_ROWS:
        raise TableError(f"Больше {MAX_ROWS} строк данных: оставьте только нужные.")
    if width < 2 or any(not header for header in headers) or len(set(headers)) != len(headers):
        raise TableError("Шапка должна быть в первой строке: у каждого столбца своё непустое имя.")
    if len(body) < 2:
        raise TableError("Нужны шапка и хотя бы две строки данных.")
    if any(len(cell) > MAX_CELL for row in grid for cell in row):
        raise TableError(f"Ячейка длиннее {MAX_CELL} знаков.")
    numbers = [[parse_number(cell) for cell in row] for row in body]
    return Table(name=name, headers=headers, rows=body, numbers=numbers,
                 lines=[index for index, _ in numbered[1:]], sheet=sheet)
