"""Тесты конвертера pptx в pdf и в png и постоянной сессии движка. Задачи T-11 и T-26.

Настоящий soffice/PowerPoint по умолчанию не запускаются: без явного согласия
владельца машины (переменная DESIGNER_RUN_LIVE_CONVERTER_TESTS=1) эти проверки
пропускаются, чтобы прогон тестов не открывал окна и не трогал живой PowerPoint.
Выбор движка, сборка командной строки, протокол сессии и исключение ConverterUnavailable
проверяются всегда, без запуска настоящих программ.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from designer.export import convert
from designer.export.convert import ConverterUnavailable, available, to_pdf, to_png

_LIVE_ENV = "DESIGNER_RUN_LIVE_CONVERTER_TESTS"
PNG = b"\x89PNG\r\n\x1a\nfake"


def _sample_pptx(path: Path) -> Path:
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for i in range(3):
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.text = f"Слайд {i + 1}"
    prs.save(str(path))
    return path


@pytest.fixture
def sample_pptx(tmp_path) -> Path:
    return _sample_pptx(tmp_path / "sample.pptx")


@pytest.fixture(autouse=True)
def _clean_env(monkeypatch):
    monkeypatch.delenv("DESIGNER_CONVERTER", raising=False)
    monkeypatch.delenv("DESIGNER_SOFFICE", raising=False)


# ---------- выбор движка ----------

def test_available_empty_when_nothing_found(monkeypatch):
    monkeypatch.setattr(convert, "_soffice_path", lambda: None)
    monkeypatch.setattr(convert, "_powerpoint_registered", lambda: False)
    assert available() == []


def test_available_lists_libreoffice_when_soffice_found(monkeypatch):
    monkeypatch.setattr(convert, "_soffice_path", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(convert, "_powerpoint_registered", lambda: False)
    assert available() == ["libreoffice"]


def test_available_lists_powerpoint_when_registered(monkeypatch):
    monkeypatch.setattr(convert, "_soffice_path", lambda: None)
    monkeypatch.setattr(convert, "_powerpoint_registered", lambda: True)
    assert available() == ["powerpoint"]


def test_no_engine_raises_converter_unavailable(monkeypatch, sample_pptx, tmp_path):
    monkeypatch.setattr(convert, "_soffice_path", lambda: None)
    monkeypatch.setattr(convert, "_powerpoint_registered", lambda: False)
    with pytest.raises(ConverterUnavailable):
        to_pdf(sample_pptx, tmp_path / "out.pdf")


def test_env_var_forces_engine_choice(monkeypatch):
    monkeypatch.setattr(convert, "_soffice_path", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(convert, "_powerpoint_registered", lambda: True)
    monkeypatch.setenv("DESIGNER_CONVERTER", "powerpoint")
    assert convert._select_engine() == "powerpoint"


def test_env_var_unknown_engine_raises(monkeypatch):
    monkeypatch.setenv("DESIGNER_CONVERTER", "cloud-magic")
    with pytest.raises(ConverterUnavailable):
        convert._select_engine()


def test_env_var_requests_unavailable_engine_raises(monkeypatch):
    monkeypatch.setattr(convert, "_soffice_path", lambda: None)
    monkeypatch.setattr(convert, "_powerpoint_registered", lambda: False)
    monkeypatch.setenv("DESIGNER_CONVERTER", "libreoffice")
    with pytest.raises(ConverterUnavailable):
        convert._select_engine()


# ---------- libreoffice: командная строка и обработка результата, без настоящего soffice ----------

def test_soffice_pdf_command_is_headless_and_has_paths(tmp_path):
    cmd = convert._soffice_pdf_command("soffice", Path("in.pptx"), tmp_path / "out", tmp_path / "profile")
    assert cmd[0] == "soffice"
    assert "--headless" in cmd
    assert cmd[cmd.index("--convert-to") + 1] == "pdf"
    assert str(tmp_path / "out") in cmd
    assert str(Path("in.pptx")) in cmd
    assert any(part.startswith("-env:UserInstallation=") for part in cmd)


def test_libreoffice_to_pdf_runs_soffice_and_copies_result(monkeypatch, sample_pptx, tmp_path):
    monkeypatch.setattr(convert, "_soffice_path", lambda: "soffice")

    def fake_run(cmd, capture_output, timeout, check):
        out_dir = Path(cmd[cmd.index("--outdir") + 1])
        (out_dir / f"{sample_pptx.stem}.pdf").write_bytes(b"%PDF-1.4\n%fake")
        return subprocess.CompletedProcess(cmd, 0, stdout=b"", stderr=b"")

    monkeypatch.setattr(convert.subprocess, "run", fake_run)
    out_path = tmp_path / "result.pdf"
    result = to_pdf(sample_pptx, out_path)
    assert result == out_path
    assert result.read_bytes().startswith(b"%PDF")


def test_libreoffice_to_pdf_raises_when_soffice_produces_nothing(monkeypatch, sample_pptx, tmp_path):
    monkeypatch.setattr(convert, "_soffice_path", lambda: "soffice")

    def fake_run(cmd, capture_output, timeout, check):
        return subprocess.CompletedProcess(cmd, 1, stdout=b"", stderr="не шмогла".encode("utf-8"))

    monkeypatch.setattr(convert.subprocess, "run", fake_run)
    with pytest.raises(ConverterUnavailable, match="не шмогла"):
        to_pdf(sample_pptx, tmp_path / "result.pdf")


# ---------- растеризация pdf в png без установленной библиотеки ----------

def test_rasterize_without_library_raises_converter_unavailable(monkeypatch, tmp_path):
    monkeypatch.setattr(convert, "_pdftoppm_path", lambda: None)
    monkeypatch.setitem(sys.modules, "fitz", None)
    pdf_path = tmp_path / "in.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n%fake")
    with pytest.raises(ConverterUnavailable):
        convert._rasterize_pdf_to_png(pdf_path, tmp_path, 1280)


def test_rasterize_prefers_pdftoppm_when_it_is_in_path(monkeypatch, tmp_path):
    """pdftoppm из PATH: страницы получают наши имена slide-NNN.png. Саму программу не зовём."""
    monkeypatch.setattr(convert, "_pdftoppm_path", lambda: "pdftoppm")
    captured = {}

    def fake_run(cmd, capture_output, timeout, check):
        captured["cmd"] = cmd
        prefix = Path(cmd[-1])
        for page in (1, 2):
            (prefix.parent / f"page-{page}.png").write_bytes(PNG)
        return subprocess.CompletedProcess(cmd, 0, stdout=b"", stderr=b"")

    monkeypatch.setattr(convert.subprocess, "run", fake_run)
    pdf_path = tmp_path / "in.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n%fake")

    pages = convert._rasterize_pdf_to_png(pdf_path, tmp_path, 1280)

    assert [page.name for page in pages] == ["slide-001.png", "slide-002.png"]
    assert captured["cmd"][:2] == ["pdftoppm", "-png"]
    assert captured["cmd"][captured["cmd"].index("-scale-to-x") + 1] == "1280"


def test_pdftoppm_without_result_raises_converter_unavailable(monkeypatch, tmp_path):
    monkeypatch.setattr(convert, "_pdftoppm_path", lambda: "pdftoppm")
    monkeypatch.setattr(
        convert.subprocess, "run",
        lambda cmd, capture_output, timeout, check: subprocess.CompletedProcess(
            cmd, 1, stdout=b"", stderr="нет такой страницы".encode("utf-8")),
    )
    pdf_path = tmp_path / "in.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n%fake")

    with pytest.raises(ConverterUnavailable, match="нет такой страницы"):
        convert._rasterize_pdf_to_png(pdf_path, tmp_path, 1280)


def test_to_png_with_libreoffice_and_no_rasterizer_raises_not_empty_list(monkeypatch, sample_pptx, tmp_path):
    monkeypatch.setattr(convert, "_soffice_path", lambda: "soffice")
    monkeypatch.setattr(convert, "_powerpoint_registered", lambda: False)
    monkeypatch.setattr(convert, "_pdftoppm_path", lambda: None)
    monkeypatch.setitem(sys.modules, "fitz", None)

    def fake_libreoffice_to_pdf(pptx_path, out_path):
        out_path.write_bytes(b"%PDF-1.4\n%fake")

    monkeypatch.setattr(convert, "_libreoffice_to_pdf", fake_libreoffice_to_pdf)
    with pytest.raises(ConverterUnavailable):
        to_png(sample_pptx, tmp_path / "png")


# ---------- powerpoint: командная строка и обработка результата, без настоящего PowerPoint ----------

def test_powerpoint_to_pdf_builds_command_without_running_it(monkeypatch, sample_pptx, tmp_path):
    captured = {}

    def fake_run_script(script, args):
        captured["script"] = script
        captured["args"] = args
        out_path = Path(args[args.index("-OutputPath") + 1])
        out_path.write_bytes(b"%PDF-1.4\n%fake")
        return subprocess.CompletedProcess([], 0, stdout=b"", stderr=b"")

    monkeypatch.setattr(convert, "_run_powershell_script", fake_run_script)
    out_path = tmp_path / "out.pdf"
    convert._powerpoint_to_pdf(sample_pptx, out_path)
    assert "PowerPoint.Application" in captured["script"]
    assert "-InputPath" in captured["args"]
    assert str(sample_pptx) in captured["args"]
    assert out_path.read_bytes().startswith(b"%PDF")


def test_powerpoint_to_pdf_raises_on_nonzero_exit(monkeypatch, sample_pptx, tmp_path):
    def fake_run_script(script, args):
        return subprocess.CompletedProcess([], 1, stdout=b"", stderr=b"COM error")

    monkeypatch.setattr(convert, "_run_powershell_script", fake_run_script)
    with pytest.raises(ConverterUnavailable, match="COM error"):
        convert._powerpoint_to_pdf(sample_pptx, tmp_path / "out.pdf")


def test_run_powershell_script_refuses_outside_windows(monkeypatch):
    monkeypatch.setattr(convert.sys, "platform", "linux")
    with pytest.raises(ConverterUnavailable):
        convert._run_powershell_script("param()", [])


def test_powerpoint_process_not_closed_when_already_running_keeps_alive_flag(monkeypatch):
    """Скрипт не завершает уже запущенный PowerPoint: проверяем сам PS-текст, не запуская его."""
    assert "$alreadyRunning" in convert._PS_TO_PDF
    assert "if (-not $alreadyRunning)" in convert._PS_TO_PDF
    assert "$app.Quit()" in convert._PS_TO_PDF


# ---------- T-26: постоянная сессия движка, без запуска настоящих программ ----------

class _FakeSessionProcess:
    """Дочерний процесс сессии: читает строку JSON со stdin, пишет картинку, отвечает строкой JSON.

    answers_left = 0: процесс молчит и умирает, так проверяется повторный запуск.
    """

    def __init__(self, answers_left: int = 99, error: str | None = None) -> None:
        self.commands: list[dict] = []
        self.answers_left = answers_left
        self.error = error
        self.alive = True
        self.terminated = False
        self.stdin = self
        self.stdout = self
        self._answers: list[str] = []

    # сторона stdin
    def write(self, line: str) -> None:
        if not self.alive:
            raise OSError("процесс сессии закрыт")
        command = json.loads(line)
        self.commands.append(command)
        if command["command"] == "close":
            self.alive = False
            self._answers.append(json.dumps({"ok": True}) + "\n")
            return
        if self.answers_left <= 0:
            self.alive = False
            self._answers.append("")
            return
        self.answers_left -= 1
        if self.error:
            self._answers.append(json.dumps({"ok": False, "error": self.error}) + "\n")
            return
        Path(command["out"]).write_bytes(PNG)
        self._answers.append(json.dumps({"ok": True, "out": command["out"]}) + "\n")

    def flush(self) -> None:
        pass

    # сторона stdout
    def readline(self) -> str:
        return self._answers.pop(0) if self._answers else ""

    def poll(self):
        return None if self.alive else 0

    def terminate(self) -> None:
        self.alive = False
        self.terminated = True

    def wait(self, timeout=None) -> int:
        self.alive = False
        return 0


def _powerpoint_session(monkeypatch, answers: list[int] | None = None,
                         error: str | None = None) -> tuple[convert._PowerPointSession, list]:
    """Сессия PowerPoint с подменённым дочерним процессом. Настоящий PowerPoint не запускается."""
    session = convert._PowerPointSession()
    started: list[_FakeSessionProcess] = []
    limits = list(answers or [])

    def fake_spawn():
        proc = _FakeSessionProcess(limits.pop(0) if limits else 99, error)
        started.append(proc)
        return proc

    monkeypatch.setattr(session, "_spawn", fake_spawn)
    monkeypatch.setattr(session, "_start_watcher", lambda: None)  # сторож простоя в тестах не нужен
    return session, started


def test_session_asks_engine_by_json_command_and_reads_json_answer(monkeypatch, sample_pptx):
    session, started = _powerpoint_session(monkeypatch)

    png = session.png(sample_pptx, 1, 900)

    assert png == PNG
    command = started[0].commands[0]
    assert command["command"] == "png"
    assert command["path"] == str(sample_pptx.resolve())
    assert command["slide"] == 2  # в PowerPoint слайды считаются с единицы
    assert command["width"] == 900


def test_session_starts_engine_once_for_several_slides(monkeypatch, sample_pptx):
    session, started = _powerpoint_session(monkeypatch)

    session.png(sample_pptx, 0)
    session.png(sample_pptx, 1)

    assert len(started) == 1
    assert len(started[0].commands) == 2


def test_session_reports_engine_error_from_the_answer(monkeypatch, sample_pptx):
    session, _ = _powerpoint_session(monkeypatch, error="слайда нет")

    with pytest.raises(ConverterUnavailable, match="слайда нет"):
        session.png(sample_pptx, 0)


def test_session_starts_broken_process_again(monkeypatch, sample_pptx):
    session, started = _powerpoint_session(monkeypatch, answers=[0])  # первый процесс молчит

    png = session.png(sample_pptx, 0)

    assert png == PNG
    assert len(started) == 2
    assert started[0].terminated


def test_session_gives_up_when_second_process_also_keeps_silence(monkeypatch, sample_pptx):
    session, started = _powerpoint_session(monkeypatch, answers=[0, 0])

    with pytest.raises(ConverterUnavailable, match="не отвечает"):
        session.png(sample_pptx, 0)

    assert len(started) == 2


def test_session_closes_engine_after_idle(monkeypatch, sample_pptx):
    session, started = _powerpoint_session(monkeypatch)
    session.png(sample_pptx, 0)

    assert session._close_if_idle() is False  # им только что пользовались

    session._last_used = time.monotonic() - convert.IDLE_TIMEOUT_S - 1
    assert session._close_if_idle() is True
    assert started[0].commands[-1] == {"command": "close"}
    assert not started[0].alive


def test_session_starts_engine_again_after_idle_close(monkeypatch, sample_pptx):
    session, started = _powerpoint_session(monkeypatch)
    session.png(sample_pptx, 0)
    session._last_used = time.monotonic() - convert.IDLE_TIMEOUT_S - 1
    session._close_if_idle()

    session.png(sample_pptx, 0)

    assert len(started) == 2


def test_session_close_says_goodbye_to_the_engine(monkeypatch, sample_pptx):
    session, started = _powerpoint_session(monkeypatch)
    session.png(sample_pptx, 0)

    session.close()

    assert started[0].commands[-1] == {"command": "close"}


def test_session_script_does_not_close_powerpoint_started_by_a_person():
    """Скрипт сессии не завершает чужой PowerPoint: проверяем сам PS-текст, не запуская его."""
    assert "$alreadyRunning" in convert._PS_SESSION
    assert "-not $alreadyRunning" in convert._PS_SESSION
    assert "$app.Quit()" in convert._PS_SESSION


def test_libreoffice_session_converts_file_once_for_all_slides(monkeypatch, sample_pptx):
    widths: list[int] = []

    def fake_to_png(pptx_path, out_dir, width_px):
        widths.append(width_px)
        pages = []
        for index in range(3):
            page = out_dir / f"slide-{index + 1}.png"
            page.write_bytes(PNG + str(index).encode("ascii"))
            pages.append(page)
        return pages

    monkeypatch.setattr(convert, "_libreoffice_to_png", fake_to_png)
    session = convert._LibreOfficeSession()

    first = session.png(sample_pptx, 0, 900)
    third = session.png(sample_pptx, 2, 900)

    assert first != third
    assert widths == [900]  # soffice зовётся на файл целиком один раз
    with pytest.raises(ConverterUnavailable):
        session.png(sample_pptx, 9, 900)


def test_get_session_is_one_for_the_whole_process(monkeypatch):
    monkeypatch.setattr(convert, "_soffice_path", lambda: "soffice")
    monkeypatch.setattr(convert, "_powerpoint_registered", lambda: False)
    convert.close_session()
    try:
        session = convert.get_session()
        assert isinstance(session, convert._LibreOfficeSession)
        assert convert.get_session() is session
    finally:
        convert.close_session()


# ---------- настоящий движок: по умолчанию пропускается, чтобы не открывать окна ----------

def _require_live_engine():
    if os.environ.get(_LIVE_ENV) != "1":
        pytest.skip(
            f"живой прогон настоящего движка отключён по умолчанию (не открываем окна и не "
            f"трогаем живой PowerPoint без согласия): установите {_LIVE_ENV}=1"
        )
    if not available():
        pytest.skip("на этой машине нет ни soffice, ни PowerPoint")


def _powerpoint_process_running() -> bool:
    result = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command",
         "(Get-Process -Name POWERPNT -ErrorAction SilentlyContinue) -ne $null"],
        capture_output=True, timeout=30, check=False,
    )
    return result.stdout.decode("utf-8", "replace").strip().lower() == "true"


def test_to_pdf_with_real_engine_produces_valid_pdf(sample_pptx, tmp_path):
    _require_live_engine()
    out = to_pdf(sample_pptx, tmp_path / "out.pdf")
    assert out.is_file()
    assert out.read_bytes()[:5] == b"%PDF-"


def test_to_png_with_real_engine_produces_three_images_of_requested_width(sample_pptx, tmp_path):
    _require_live_engine()
    try:
        pages = to_png(sample_pptx, tmp_path / "png", width_px=1280)
    except ConverterUnavailable as exc:
        pytest.skip(f"движок есть, но растеризация pdf в png недоступна: {exc}")
    assert len(pages) == 3
    for page in pages:
        with Image.open(page) as img:
            assert img.width == 1280


def test_powerpoint_process_survives_if_already_running(sample_pptx, tmp_path, monkeypatch):
    _require_live_engine()
    if "powerpoint" not in available() or sys.platform != "win32":
        pytest.skip("PowerPoint недоступен")
    if not _powerpoint_process_running():
        pytest.skip("PowerPoint не был запущен до теста — нечего проверять на выживание")
    monkeypatch.setenv("DESIGNER_CONVERTER", "powerpoint")
    to_pdf(sample_pptx, tmp_path / "out.pdf")
    assert _powerpoint_process_running()
