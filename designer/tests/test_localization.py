"""Locale is a display operation: native exports and numerical evidence agree."""
import copy
import json
import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from pydantic import ValidationError
from pptx import Presentation

from designer.api.schemas import DeckCreateRequest
from designer.contracts import (ChartSpec, DataProvenance, DeckPlan, Element, Scene, Series,
                                SlideIntent, TableSpec, Tokens, Margins)
from designer.insights.labels import field_label
from designer.localization import (category_label, evidence_fingerprint, language_findings,
                                  language_instruction, localize_plan, localize_slide, migrate_deck_locale)
from designer.viz.pptx_native import add_chart, add_table
from designer.export.svg_chart import render_chart


def analytical(**overrides):
    props = dict(
        id="s1", kind="chart", title="actual_sales: максимум 12 — «North»",
        key_message="Среди 2 полных наблюдений по полю «actual_sales» максимальное значение 12: «North».",
        notes="Источник: original.csv. Расчёт: max(actual_sales) = 12.",
        provenance=DataProvenance(dataset_id="dataset-x", source_sha256="original-sha",
            content_hash="canonical-hash", file_name="original.csv", table_index=0,
            columns=["region", "actual_sales"], total_rows=2, row_lines=[2, 3], excluded_row_lines=[],
            displayed_row_lines=[2, 3], visualization="bar", aggregation="none",
            missing_policy="exclude_incomplete_selected_fields", formula="max(actual_sales) = 12", unit="USD"),
        chart=ChartSpec(type="bar", title="actual_sales · region", unit="USD",
            categories=["North", "ACME"], x_axis_title="actual_sales", y_axis_title="region",
            series=[Series(name="actual_sales", values=[12, 0])]))
    props.update(overrides)
    return SlideIntent(**props)


def plan(slide):
    return DeckPlan(title="Итоги", purpose="обзор", slides=[slide])


def test_default_and_supported_language_contract():
    assert DeckCreateRequest(design_system_id="ds", brief="brief").language == "ru"
    assert DeckCreateRequest(design_system_id="ds", brief="brief", language="en").language == "en"
    with pytest.raises(ValidationError):
        DeckCreateRequest(design_system_id="ds", brief="brief", language="de")
    with pytest.raises(ValueError):
        language_instruction("de")


@pytest.mark.parametrize("source,expected,unit", [
    ("actual_sales", "Продажи (факт)", ""), ("planned_profit", "Прибыль (план)", ""),
    ("target_leads", "Лиды (цель)", ""), ("revenue_USD", "Выручка", "USD"),
    ("ACME-API", "ACME-API", ""), ("unknown_field", "unknown_field", ""),
])
def test_vocabulary_is_field_based_not_dataset_or_template_based(source, expected, unit):
    assert field_label(source) == (expected, unit)


def test_labels_localize_consistently_without_mutating_source_or_numbers():
    source = analytical()
    old = source.model_dump()
    localized = localize_slide(source)
    assert localized.title == "Продажи (факт): максимум 12 — «Север»"
    assert localized.chart.categories == ["Север", "ACME"]
    assert localized.chart.series[0].name == "Продажи (факт)"
    assert localized.chart.x_axis_title == "Продажи (факт)"
    assert localized.provenance == source.provenance and localized.notes == source.notes
    assert source.model_dump() == old
    assert evidence_fingerprint(plan(source)) == evidence_fingerprint(plan(localized))
    assert localize_slide(localized) == localized


def test_table_zero_missing_ids_urls_names_and_numerical_strings_are_untouched():
    source = analytical(kind="table", chart=None, table=TableSpec(
        columns=["period", "actual_units", "planned_units", "note", "label", "reference_id", "brand"],
        rows=[["Week 007", "0", "", "actual missing", "Reference item 001", "REF-001", "Search"],
              ["Week 008", "-2.50 USD", "1e5", "original author prose", "https://example.com/Item", "001", "ACME"]]))
    result = localize_slide(source)
    assert result.table.rows[0] == ["Неделя 007", "0", "", "Факт отсутствует", "Запись справочника 001", "REF-001", "Search"]
    assert result.table.rows[1][1:] == source.table.rows[1][1:]
    assert evidence_fingerprint(plan(source)) == evidence_fingerprint(plan(result))
    assert category_label("Search", "company", "ru") == "Search"


def test_english_retains_informative_claims_and_all_numbers():
    source = analytical()
    result = localize_slide(source, "en")
    assert result.title == "actual sales: maximum 12 — «North»"
    assert "maximum value 12" in result.key_message
    assert re.findall(r"\d+", result.key_message) == re.findall(r"\d+", source.key_message)
    assert not re.search("[А-Яа-я]", result.title + result.key_message)
    assert evidence_fingerprint(plan(source)) == evidence_fingerprint(plan(result))


def test_english_preserves_percentage_claim_and_distribution_outlier_qualification():
    source = analytical(
        title="actual_sales: «North» — 62.00% показанной суммы",
        key_message="Максимальная доля: «North», 620 из 1000 (62.00%, округлено до 2 знаков). Доли относятся только к показанным строкам.")
    result = localize_slide(source, "en")
    assert "62.00% of the displayed total" in result.title
    assert "620 of 1000" in result.key_message
    assert "Shares refer only to the displayed rows" in result.key_message


def test_unknown_fields_are_preserved_with_visible_review_finding():
    scene = Scene(slide_id="s1", pattern_id="p", elements=[
        Element(id="e", type="text", box=(0, 0, 1, 1), text="net retention: customer cohort"),
        Element(id="source", type="text", role="source", box=(0, 0, 1, 1), text="Источник: original.csv"),
    ])
    findings = language_findings([scene], "ru")
    assert len(findings) == 1 and findings[0].check_id == "language_unresolved"
    assert "net retention" in findings[0].message
    assert "original.csv" not in findings[0].message


def test_template_residue_is_not_certified_as_complete():
    scene = Scene(slide_id="s1", pattern_id="p", elements=[
        Element(id="e", type="text", box=(0, 0, 1, 1), text="Click to add title")])
    assert "template_residue" in {f.check_id for f in language_findings([scene], "ru")}


@pytest.mark.parametrize("language", ["ru", "en"])
def test_native_chart_and_svg_share_labels_and_exact_values(tmp_path, language):
    chart = localize_slide(analytical(), language).chart
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tokens = Tokens(colors=[], fonts=[], type_scale=[], margins=Margins(left=0, top=0, right=0, bottom=0))
    add_chart(slide, chart, (0.1, 0.1, 0.8, 0.7), (prs.slide_width, prs.slide_height), tokens)
    path = tmp_path / "localized.pptx"
    prs.save(path)
    native = Presentation(path).slides[0].shapes[0].chart
    assert list(native.series[0].values) == [12, 0]
    assert native.series[0].name == chart.series[0].name
    svg = render_chart(chart, ["336699"])
    assert chart.categories[0] in svg
    assert native.category_axis.axis_title.text_frame.text == chart.y_axis_title
    with ZipFile(path) as archive:
        xml = archive.read("ppt/charts/chart1.xml").decode()
        assert chart.categories[0] in xml


def test_migration_refuses_unreviewed_editor_revisions(monkeypatch):
    class State:
        raw = {"revision": "manual-edit"}
    monkeypatch.setattr("designer.edit._load", lambda *_: State())
    with pytest.raises(ValueError, match="editor revisions"):
        migrate_deck_locale("test-deck", language="ru")
    with pytest.raises(ValueError, match="supports ru"):
        migrate_deck_locale("test-deck", language="en")
    State.raw = {}
    State.plan = plan(analytical())
    State.plan.language = "en"
    with pytest.raises(ValueError, match="создана на английском"):
        migrate_deck_locale("test-deck", language="ru")


def test_migration_snapshots_exports_and_keeps_provenance_without_edited_data_warning(tmp_path, monkeypatch):
    from designer import pipeline, store
    from designer.contracts import Deck
    from designer.layout.compose import compose
    from designer.layout.scene import build_scene
    from designer.layout.match import choose_pattern
    from designer.export.pptx_deck import export_pptx
    monkeypatch.setenv(store.DATA_DIR_ENV, str(tmp_path / "data"))
    monkeypatch.setattr("designer.pipeline.convert.available", lambda: [])
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[1])
    slide.shapes.title.text = "Template title"
    slide.placeholders[1].text = "Template body"
    buffer = BytesIO()
    source.save(buffer)
    ds = pipeline.import_template(buffer.getvalue(), "minimal.pptx")
    package = store.design_system_dir(ds.id)
    intent = analytical()
    original_plan = plan(intent)
    pattern = ds.patterns[0]
    spec = compose(intent, pattern, ds)
    scene = build_scene(spec, pattern, ds, package)
    deck_id = store.new_deck_id()
    store.init_deck(deck_id, ds.id, ["a"])
    store.save_deck_result(deck_id, Deck(id=deck_id, design_system_id=ds.id, variant="a",
                          plan=original_plan, specs=[spec], scenes=[scene]), [], variant="a")
    files = store.deck_variant_files_dir(deck_id, "a")
    export_pptx([spec], ds, package, files / "deck.pptx")
    before_pptx = (files / "deck.pptx").read_bytes()
    before_json = store.deck_variant_state_path(deck_id, "a").read_bytes()
    report = migrate_deck_locale(deck_id)
    backup = Path(report["backup"])
    assert (backup / "deck.json").read_bytes() == before_json
    assert (backup / "files/deck.pptx").read_bytes() == before_pptx
    after = store.load_deck_state(deck_id, "a")
    assert evidence_fingerprint(DeckPlan.model_validate(after["plan"])) == evidence_fingerprint(original_plan)
    assert after["plan"]["slides"][0]["provenance"] == intent.provenance.model_dump()
    assert not any(f["check_id"] == "data_source_edited" for f in after["findings"])
    assert "Продажи (факт)" in (files / "deck.markup.html").read_text(encoding="utf-8")
    native = Presentation(files / "deck.pptx")
    chart = next(s.chart for s in native.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values) == [12, 0]
    assert chart.series[0].name == "Продажи (факт)"
    assert after["revision"] == json.loads((files / "export-revision.json").read_text())["revision"]


def test_planner_and_writer_receive_explicit_locale_and_plan_does_not_trust_model_language():
    from designer.plan.planner import make_plan
    from designer.plan.writer import fill_slots
    class Client:
        systems = []
        def complete_json(self, *, system, user, schema, **kwargs):
            self.systems.append(system)
            if "slides" in schema["properties"]:
                return {"title": "Results", "purpose": "review", "language": "ru", "slides": [
                    {"id": "s1", "kind": "title", "title": "Results"},
                    {"id": "s2", "kind": "thanks", "title": "Thank you"}]}
            return {"title": "Results"}
    client = Client()
    result = make_plan("Review", "", "", 2, client, language="en")
    assert result.language == "en"
    fill_slots(result.slides[0], {"title": 40}, {}, 0, client, language="en")
    assert len(client.systems) == 2
    assert all("Required presentation language: English (en)" in s for s in client.systems)


def test_locale_ignoring_model_gets_one_retry_then_fails():
    from designer.plan.planner import make_plan
    class Client:
        calls = 0
        def complete_json(self, **kwargs):
            self.calls += 1
            return {"title": "Results", "purpose": "review", "slides": [
                {"id": "s1", "kind": "title", "title": "English prose " * 20},
                {"id": "s2", "kind": "thanks", "title": "Thank you"}]}
    client = Client()
    with pytest.raises(ValueError, match="language: ru"):
        make_plan("Итоги", "", "", 2, client, language="ru")
    assert client.calls == 2
