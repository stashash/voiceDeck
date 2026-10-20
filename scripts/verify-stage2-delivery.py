"""Evidence checks over generated files and persisted live API state."""
from pathlib import Path
import hashlib
import json
import sys
import urllib.request
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

ROOT=Path(__file__).resolve().parents[1]
BASE="http://127.0.0.1:8090"
OUT=ROOT/"artifacts/stage2-demo"
challenge=json.loads((OUT/"final-v3/run-manifest.json").read_text(encoding="utf-8"))
challenge_only="--challenge-only" in sys.argv
capture={} if challenge_only else json.loads((OUT/"screencast/capture-manifest.json").read_text(encoding="utf-8"))
if not challenge_only:
    assert capture["status"]=="passed" and not capture["browserErrors"]
assert challenge["pdf_count"]==9 and len(challenge["runs"])==9
assert len({r["design_system_id"] for r in challenge["templates"]})==3
expected={"time_series.xlsx":"line","category_shares.json":"pie","distribution.csv":"histogram",
          "category_comparison.csv":"grouped_column","missing_zero.csv":"table","reference_201.csv":"table"}
states={}
design_systems={}
results=[]
files=[]
source_hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (OUT/"content-v2").iterdir() if p.is_file()}
records=[*challenge["runs"]]
if not challenge_only:
    records.append({"deck_id":capture["deckId"],"variant":"a","output_dir":str(OUT/"screencast"),"unknown_template":True})
for record in records:
    deck_id=record["deck_id"]
    if deck_id not in states:
        with urllib.request.urlopen(f"{BASE}/decks/{deck_id}",timeout=60) as response:
            states[deck_id]=json.load(response)
    variant=record["variant"]
    state=states[deck_id]["variants"][variant]
    assert state["status"]=="done",(deck_id,variant)
    slides=state["plan"]["slides"]
    data_slides=[s for s in slides if s.get("provenance")]
    found={}
    for slide in data_slides:
        p=slide["provenance"]
        name=p["file_name"]
        found.setdefault(name,set()).add(p["visualization"])
        assert slide["source_note"] and p["source_sha256"] and p["formula"]
        assert p["source_sha256"]==source_hashes[name],(name,"source hash mismatch")
        if slide.get("chart") and slide["chart"]["type"] not in {"pie","kpi"}:
            assert slide["chart"].get("x_axis_title") and slide["chart"].get("y_axis_title"),(name,"axes")
    for name,kind in expected.items():
        assert found.get(name)=={kind},(deck_id,variant,name,found.get(name))
    reference=[s for s in data_slides if s["provenance"]["file_name"]=="reference_201.csv" and s.get("table")]
    lines=[line for s in reference for line in s["provenance"]["displayed_row_lines"]]
    assert len(lines)==201 and len(set(lines))==201,"Lost/duplicated reference rows"
    first_appendix=next(i for i,s in enumerate(slides) if s in reference)
    assert all(s.get("provenance",{}).get("page_count",0)>1 for s in slides[first_appendix:]),"Appendix interrupts main story"
    missing=next(s for s in data_slides if s["provenance"]["file_name"]=="missing_zero.csv")
    values=[str(v) for row in missing["table"]["rows"] for v in row]
    assert "8" in values and "0" in values and any(v in {"","—"} for v in values),values
    trend=next(s for s in data_slides if s["provenance"]["file_name"]=="time_series.xlsx")
    assert "30" in trend["title"] and "USD" in trend["title"]
    directory=Path(record["output_dir"])
    for name in ["deck.pdf","deck.pptx","deck.html"]:
        path=directory/name
        assert path.is_file() and path.stat().st_size>1000,path
        files.append({"path":str(path.relative_to(ROOT)),"bytes":path.stat().st_size,"sha256":hashlib.sha256(path.read_bytes()).hexdigest()})
    prs=Presentation(directory/"deck.pptx")
    assert len(prs.slides)==len(slides)==len(state["slide_images"])
    for native_slide, scene in zip(prs.slides, state["scenes"]):
        bar_specs=[e["chart"] for e in scene["elements"] if e.get("chart",{} ) and e["chart"]["type"]=="bar"]
        bar_charts=[sh.chart for sh in native_slide.shapes if sh.has_chart and sh.chart.chart_type==XL_CHART_TYPE.BAR_CLUSTERED]
        assert len(bar_specs)==len(bar_charts)
        for chart, spec in zip(bar_charts,bar_specs):
            assert chart.category_axis.axis_title.text_frame.text==spec["y_axis_title"], "Horizontal category/Y axis title mismatch"
            assert chart.value_axis.axis_title.text_frame.text==(spec["x_axis_title"] or spec["unit"]), "Horizontal value/X axis title mismatch"
            assert [str(c.label) for c in chart.plots[0].categories]==spec["categories"]
            assert [list(series.values) for series in chart.series]==[series["values"] for series in spec["series"]]
    native_charts=[sh.chart for sl in prs.slides for sh in sl.shapes if sh.has_chart]
    assert len(native_charts)>=5,"Expected editable native charts"
    dsid=state["design_system_id"]
    if dsid not in design_systems:
        with urllib.request.urlopen(f"{BASE}/design-systems/{dsid}") as response:design_systems[dsid]=json.load(response)
    ds=design_systems[dsid]
    accent=next(c["hex"] for c in ds["tokens"]["colors"] if c["role"]=="accent")
    family=next(f["family"] for f in ds["tokens"]["fonts"] if f["role"]=="body")
    legend=(ds["tokens"].get("chart_style") or {}).get("legend_position")
    for chart in native_charts:
        assert chart.font.name==family
        if chart.chart_type==XL_CHART_TYPE.LINE:
            assert str(chart.series[0].format.line.color.rgb)==accent
        if legend:
            assert chart.has_legend and chart.legend.position=={"bottom":XL_LEGEND_POSITION.BOTTOM,"top":XL_LEGEND_POSITION.TOP,
                    "left":XL_LEGEND_POSITION.LEFT,"right":XL_LEGEND_POSITION.RIGHT}[legend]
    if record.get("unknown_template"):
        assert next(f["family"] for f in ds["tokens"]["fonts"] if f["role"]=="body")=="Play"
        assert next(c["hex"] for c in ds["tokens"]["colors"] if c["role"]=="accent")=="713CDA"
        assert ds.get("brand_rules",{}).get("logo_asset")
        for i,sl in enumerate(prs.slides):
            assert any(sh.shape_type==13 for sh in sl.shapes),(i+1,"missing logo")
        assert "Play" in (directory/"deck.html").read_text(encoding="utf-8")
    results.append({"deckId":deck_id,"variant":variant,"slides":len(slides),"datasets":{k:sorted(v) for k,v in found.items()},
                    "referenceRows":len(lines),"nativeEditableCharts":len(native_charts),"unknownTemplate":bool(record.get("unknown_template"))})
evidence={"status":"passed","syntheticData":True,"challengePresentations":9,"unknownTemplatePresentations":0 if challenge_only else 1,
          "browserErrors":capture.get("browserErrors",[]),"checks":results,"files":files,
          "limits":["This verifies supported paths and supplied examples, not every possible Figma/PDF document.",
                    "Automatic contextual vision audit samples reference appendix boundaries; deterministic checks cover every slide."]}
(OUT/("challenge-verification.json" if challenge_only else "delivery-verification.json")).write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"status":"passed","presentations":len(results),"files":len(files)}))
