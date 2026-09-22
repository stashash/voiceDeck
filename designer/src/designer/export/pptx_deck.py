"""Сборка pptx из слайдов-образцов. Владелец: задача T-08.

Каждая инструкция вёрстки это клон слайда-образца: текст встаёт в слоты, блоки
перекладываются общим модулем размещения, диаграмма и таблица кладутся родными
объектами PowerPoint. Исходные слайды в конце удаляются, макеты и мастера остаются.
"""
from __future__ import annotations

import copy
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from designer.contracts import Box, DesignSystem, Pattern, RepeatGroup, SlideSpec
from designer.export.pptx_clone import clone_slide
from designer.export.pptx_text import set_text
from designer.layout.units import map_shape_box, place_units
from designer.parse.package import SOURCE_NAME
from designer.viz.pptx_native import add_chart, add_table

IDENTITY = (1.0, 1.0, 0.0, 0.0)

VIZ_COVER = 0.6
"""Доля площади фигуры внутри области: выше неё фигура образца уступает место диаграмме."""

KEEP_SIZE_KIND = "image"
"""Что не тянется вместе с блоком: картинки и значки только смещаются."""

_NOT_A_SHAPE = (qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst"))


# ---------- фигуры клона ----------

@dataclass
class _Placed:
    """Фигура собранного слайда: её XML, рамка в долях слайда и перевод координат."""

    element: object
    box: Box
    transform: tuple[float, float, float, float]
    kind: str
    dropped: bool = False


def _compose(outer: tuple, inner: tuple) -> tuple:
    return (
        outer[0] * inner[0],
        outer[1] * inner[1],
        outer[0] * inner[2] + outer[2],
        outer[1] * inner[3] + outer[3],
    )


def _group_transform(group) -> tuple:
    """Перевод координат детей группы в координаты слайда. Считается так же, как при разборе."""
    props = group._element.find(qn("p:grpSpPr"))
    xfrm = props.find(qn("a:xfrm")) if props is not None else None
    if xfrm is None:
        return IDENTITY
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return IDENTITY
    ch_cx = int(ch_ext.get("cx") or 0) or 1
    ch_cy = int(ch_ext.get("cy") or 0) or 1
    sx = int(ext.get("cx") or 0) / ch_cx
    sy = int(ext.get("cy") or 0) / ch_cy
    tx = int(off.get("x") or 0) - int(ch_off.get("x") or 0) * sx
    ty = int(off.get("y") or 0) - int(ch_off.get("y") or 0) * sy
    return (sx, sy, tx, ty)


def _kind_of(shape) -> str:
    if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
        return "image"
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    if shape.has_text_frame:
        return "text"
    return "shape"


def _box_of(shape, transform: tuple, slide_size: tuple[int, int]) -> Box:
    slide_w, slide_h = slide_size
    x = transform[0] * int(shape.left or 0) + transform[2]
    y = transform[1] * int(shape.top or 0) + transform[3]
    w = transform[0] * int(shape.width or 0)
    h = transform[1] * int(shape.height or 0)
    return (x / slide_w, y / slide_h, w / slide_w, h / slide_h)


def _walk(shapes, slide_size, transform=IDENTITY, out=None) -> dict[int, _Placed]:
    """Плоский указатель фигур слайда по их идентификаторам, с раскрытием групп."""
    if out is None:
        out = {}
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _walk(shape.shapes, slide_size, _compose(transform, _group_transform(shape)), out)
            continue
        out[shape.shape_id] = _Placed(
            element=shape._element,
            box=_box_of(shape, transform, slide_size),
            transform=transform,
            kind=_kind_of(shape),
        )
    return out


def _set_box(item: _Placed, box: Box, slide_size: tuple[int, int]) -> None:
    """Ставит фигуру в рамку слайда, пересчитав её в координаты родителя."""
    slide_w, slide_h = slide_size
    sx, sy, tx, ty = item.transform
    element = item.element
    element.x = int(round((box[0] * slide_w - tx) / sx))
    element.y = int(round((box[1] * slide_h - ty) / sy))
    element.cx = max(int(round(box[2] * slide_w / sx)), 1)
    element.cy = max(int(round(box[3] * slide_h / sy)), 1)
    item.box = box


def _drop(item: _Placed | None) -> None:
    if item is None or item.dropped:
        return
    parent = item.element.getparent()
    if parent is not None:
        parent.remove(item.element)
    item.dropped = True


def _live(placed: dict[int, _Placed], shape_id: int) -> _Placed | None:
    item = placed.get(shape_id)
    return None if item is None or item.dropped else item


# ---------- заполнение слайда ----------

def _fill_slots(placed: dict[int, _Placed], spec: SlideSpec, pattern: Pattern, slide_size: tuple[int, int]) -> None:
    slots = {slot.id: slot for slot in pattern.slots}
    for slot_id, text in spec.slot_text.items():
        slot = slots.get(slot_id)
        if slot is None:
            continue
        item = _live(placed, slot.shape_id)
        if item is not None:
            if abs(item.box[2] - slot.box[2]) > 0.005:
                _set_box(item, slot.box, slide_size)  # разбор сузил слот, чтобы текст не лёг на картинку макета
            set_text(item.element, text, spec.fitted_size_pt.get(slot_id), wrap=slot.role != "number")


def _group_of(pattern: Pattern, group_id: str | None) -> RepeatGroup | None:
    if group_id:
        return next((g for g in pattern.groups if g.id == group_id), None)
    return pattern.groups[0] if pattern.groups else None


def _unit_map(group: RepeatGroup, unit, placed: dict[int, _Placed]) -> dict[str, int]:
    """Слот блока -> фигура этого блока: по рамке относительно левого верхнего угла блока."""
    result: dict[str, int] = {}
    taken: set[int] = set()
    for slot in group.unit_slots:
        best, best_cost = None, None
        for shape_id in unit.shape_ids:
            if shape_id in taken:
                continue
            item = _live(placed, shape_id)
            if item is None or item.kind != "text":
                continue
            cost = (
                abs(item.box[0] - unit.box[0] - slot.box[0])
                + abs(item.box[1] - unit.box[1] - slot.box[1])
                + abs(item.box[2] - slot.box[2])
                + abs(item.box[3] - slot.box[3])
            )
            if best_cost is None or cost < best_cost:
                best, best_cost = shape_id, cost
        if best is not None:
            result[slot.id] = best
            taken.add(best)
    return result


def _next_id(slide):
    """Выдаёт свободные идентификаторы фигур этого слайда."""
    used = [
        int(node.get("id"))
        for node in slide._element.iter(qn("p:cNvPr"))
        if (node.get("id") or "").isdigit()
    ]
    counter = max(used, default=1)

    def issue() -> int:
        nonlocal counter
        counter += 1
        return counter

    return issue


def _move_unit(placed, shape_ids, old_box: Box, new_box: Box, slide_size) -> None:
    for shape_id in shape_ids:
        item = _live(placed, shape_id)
        if item is None:
            continue
        keep = item.kind == KEEP_SIZE_KIND
        _set_box(item, map_shape_box(old_box, new_box, item.box, keep_size=keep), slide_size)


def _copy_unit(placed, shape_ids, unit_map, old_box: Box, new_box: Box, slide_size, issue) -> dict[str, int]:
    """Новый блок копией фигур образца; возвращает своё соответствие слот -> фигура."""
    copied: dict[int, int] = {}
    for shape_id in shape_ids:
        item = _live(placed, shape_id)
        if item is None:
            continue
        element = copy.deepcopy(item.element)
        parent = item.element.getparent()
        if parent is None:
            continue
        parent.append(element)
        new_id = _renumber(element, issue)
        fresh = _Placed(element=element, box=item.box, transform=item.transform, kind=item.kind)
        keep = item.kind == KEEP_SIZE_KIND
        _set_box(fresh, map_shape_box(old_box, new_box, item.box, keep_size=keep), slide_size)
        placed[new_id] = fresh
        copied[shape_id] = new_id
    return {slot_id: copied[shape_id] for slot_id, shape_id in unit_map.items() if shape_id in copied}


def _renumber(element, issue) -> int:
    """Новые идентификаторы фигурам копии: на слайде они обязаны быть уникальными."""
    first = 0
    for node in element.iter(qn("p:cNvPr")):
        value = issue()
        node.set("id", str(value))
        first = first or value
    return first


def _fill_repeat_group(
    placed, group: RepeatGroup, unit_texts: list[dict[str, str]], count: int,
    fitted_size_pt: dict[str, float], ds: DesignSystem, issue,
) -> None:
    """Перекладывает и заполняет блоки одной группы на count блоков. Общая часть для главной и связанных групп."""
    boxes = place_units(group, count)
    units = group.units
    kept = min(count, len(units))
    maps = [_unit_map(group, units[i], placed) for i in range(kept)]

    for unit in units[kept:]:  # лишние блоки уходят со всеми своими фигурами
        for shape_id in unit.shape_ids:
            _drop(placed.get(shape_id))
    for i in range(kept):
        _move_unit(placed, units[i].shape_ids, units[i].box, boxes[i], ds.slide_size_emu)
    for i in range(kept, count):  # недостающие блоки это копия последнего
        maps.append(
            _copy_unit(
                placed, units[kept - 1].shape_ids, maps[kept - 1],
                boxes[kept - 1], boxes[i], ds.slide_size_emu, issue,
            )
        )
    for unit_text, unit_map in zip(unit_texts, maps):
        for slot_id, text in unit_text.items():
            item = _live(placed, unit_map.get(slot_id, -1))
            if item is not None:
                set_text(item.element, text, fitted_size_pt.get(slot_id))


def _fill_group(placed, spec: SlideSpec, pattern: Pattern, ds: DesignSystem, issue) -> None:
    group = _group_of(pattern, spec.group_id)
    if group is None or not group.units or not spec.unit_text:
        return
    count = len(spec.unit_text)
    _fill_repeat_group(placed, group, spec.unit_text, count, spec.fitted_size_pt, ds, issue)

    groups = {g.id: g for g in pattern.groups}
    for linked_id in group.linked_group_ids:  # связанная группа перекладывается тем же числом блоков
        linked_group = groups.get(linked_id)
        unit_texts = spec.linked_unit_text.get(linked_id)
        if linked_group is None or not linked_group.units or not unit_texts:
            continue
        _fill_repeat_group(placed, linked_group, unit_texts, count, spec.fitted_size_pt, ds, issue)


def _content_box(pattern: Pattern, ds: DesignSystem) -> Box:
    """Запасная рамка под диаграмму: поля шаблона, ниже заголовка."""
    margins = ds.tokens.margins
    top = margins.top
    for slot in pattern.slots:
        if slot.role in ("title", "subtitle"):
            top = max(top, slot.box[1] + slot.box[3] + 0.02)
    width = max(1.0 - margins.left - margins.right, 0.1)
    return (margins.left, top, width, max(1.0 - top - margins.bottom, 0.1))


def _viz_box(pattern: Pattern, area_id: str | None, ds: DesignSystem) -> Box:
    area = next((a for a in pattern.areas if a.id == area_id), None)
    return area.box if area is not None else _content_box(pattern, ds)


def _covered(box: Box, area: Box) -> float:
    """Какая доля фигуры лежит внутри области."""
    width = max(0.0, min(box[0] + box[2], area[0] + area[2]) - max(box[0], area[0]))
    height = max(0.0, min(box[1] + box[3], area[1] + area[3]) - max(box[1], area[1]))
    size = box[2] * box[3]
    return width * height / size if size > 0 else 0.0


def _fill_viz(slide, placed, spec: SlideSpec, pattern: Pattern, ds: DesignSystem) -> None:
    if spec.chart is None and spec.table is None:
        return
    box = spec.viz_box if spec.viz_box is not None else _viz_box(pattern, spec.viz_area_id, ds)
    for item in list(placed.values()):
        if not item.dropped and _covered(item.box, box) >= VIZ_COVER:
            _drop(item)
    if spec.chart is not None:
        add_chart(slide, spec.chart, box, ds.slide_size_emu, ds.tokens, theme=pattern.theme)
    else:
        add_table(slide, spec.table, box, ds.slide_size_emu, ds.tokens, theme=pattern.theme)


def _prune_groups(slide) -> None:
    """Группы, оставшиеся без фигур после удаления блоков, на слайде не нужны."""
    for group in reversed(list(slide._element.iter(qn("p:grpSp")))):
        if any(child.tag not in _NOT_A_SHAPE for child in group):
            continue
        parent = group.getparent()
        if parent is not None:
            parent.remove(group)


def _remove_shapes(placed: dict[int, _Placed], shape_ids: list[int]) -> None:
    """Фигуры, которые решила убрать вёрстка: пустые группы, образцы в рамке визуализации, заглушки."""
    for shape_id in shape_ids:
        _drop(placed.get(shape_id))


def _full_bleed_picture(layout, slide_size: tuple[int, int]) -> bool:
    """Макет несёт картинку на весь кадр: в ней могут быть нарисованы подложки и номера блоков образца."""
    w, h = slide_size
    return any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE and (shape.width or 0) >= w * 0.85 and (shape.height or 0) >= h * 0.85
        for shape in layout.shapes
    )


def _plainest_layout(master):
    """Макет того же мастера без картинок и с наименьшим числом собственных фигур."""
    def weight(layout):
        pictures = sum(1 for s in layout.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        own = sum(1 for s in layout.shapes if not s.is_placeholder)
        return (pictures, own, len(layout.shapes))
    return min(master.slide_layouts, key=weight)


def _relayout_if_bare(slide, spec: SlideSpec, pattern: Pattern, ds: DesignSystem) -> None:
    """Слайд с урезанным рядом блоков или с диаграммой уходит на чистый макет.

    Подложки карточек и водяные номера образца бывают нарисованы в картинке макета, а не фигурами:
    убрать четвёртую карточку из картинки нельзя, поэтому меняется сам макет.
    """
    layout = slide.slide_layout
    if not _full_bleed_picture(layout, ds.slide_size_emu):
        return
    unit_shapes = {sid for g in pattern.groups for u in g.units for sid in u.shape_ids}
    cut = bool(unit_shapes & set(spec.remove_shape_ids))
    if not cut and spec.viz_box is None:
        return
    plain = _plainest_layout(layout.slide_master)
    if plain is layout or _full_bleed_picture(plain, ds.slide_size_emu):
        return
    for rel in slide.part.rels.values():
        if rel.reltype.endswith("/slideLayout"):
            rel._target = plain.part  # python-pptx не даёт сменить макет иначе
            break


def _fill_slide(slide, spec: SlideSpec, pattern: Pattern, ds: DesignSystem) -> None:
    placed = _walk(slide.shapes, ds.slide_size_emu)
    issue = _next_id(slide)
    _relayout_if_bare(slide, spec, pattern, ds)
    _remove_shapes(placed, spec.remove_shape_ids)
    _fill_slots(placed, spec, pattern, ds.slide_size_emu)
    _fill_group(placed, spec, pattern, ds, issue)
    _fill_viz(slide, placed, spec, pattern, ds)
    _prune_groups(slide)
    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes


def _drop_samples(prs, count: int) -> None:
    """Убирает слайды-образцы: остаются только собранные, в порядке инструкций."""
    slide_ids = prs.slides._sldIdLst
    for node in list(slide_ids)[:count]:
        rel_id = node.rId
        slide_ids.remove(node)
        prs.part.drop_rel(rel_id)


def export_pptx(specs: list[SlideSpec], ds: DesignSystem, package_dir: Path, out_path: Path) -> Path:
    """Собирает pptx по инструкциям вёрстки и возвращает путь к сохранённому файлу."""
    prs = Presentation(str(Path(package_dir) / SOURCE_NAME))
    patterns = {pattern.id: pattern for pattern in ds.patterns}
    samples = len(prs.slides)
    for spec in specs:
        pattern = patterns.get(spec.pattern_id)
        if pattern is None:
            raise KeyError(f"в дизайн-системе нет паттерна {spec.pattern_id}")
        slide = clone_slide(prs, pattern.source_slide - 1)
        _fill_slide(slide, spec, pattern, ds)
    _drop_samples(prs, samples)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    Presentation(str(out_path))  # файл обязан открываться заново
    return out_path
