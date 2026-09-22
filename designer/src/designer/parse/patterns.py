"""Слайды-образцы как паттерны: слоты, повторяющиеся блоки, тип слайда. Владелец: задача T-02."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from designer.contracts import (
    Area,
    DecorShape,
    LayoutInfo,
    Pattern,
    RepeatGroup,
    RepeatUnit,
    SlideKind,
    Slot,
    TextStyle,
)
from designer.parse import geometry as geo

DEFAULT_SIZE_PT = 18.0
"""Кегль, если он не найден ни у фрагмента, ни у абзаца, ни у фигуры, ни у заполнителя."""

CHAR_WIDTH = 0.52
"""Средняя ширина знака в долях кегля."""

LINE_HEIGHT = 1.2
"""Интерлиньяж в долях кегля."""

EMU_PER_PT = 12700

ICON_MAX = (0.07, 0.12)
"""Картинка меньше этой рамки считается значком."""

DARK_LUMA = 0.5
"""Ниже этой яркости фон считается тёмным."""

IDENTITY = (1.0, 1.0, 0.0, 0.0)

NUMBER_RE = re.compile(r"^\d[\d\s.,:+\-/×xX]*\s*(%|[^\W\d_]{1,3})?$")

NUMBER_MASK_RE = re.compile(r"^([xхXХ]{1,4}\s*%|[xхXХ]{2,4})$")
"""Маска числа в образце: «ххх%», «x%», «XX»."""

NUMBER_MAX_CHARS = 6
"""Длина числового образца: «43 %», «1 200», «5 лет»."""

NUMBER_RATIO = 2.0
"""Во сколько раз кегль крупного числа больше кегля основного текста слайда."""

BIG_NUMBER_AREA = 0.05
"""Доля слайда под числом, с которой слайд держится на этом числе."""

BIG_NUMBER_MAX = 3
"""Сколько крупных чисел ещё читается как слайд с числом."""

GRAY_SAT = 0.1
"""Насыщенность, ниже которой заливка считается серой."""

GRAY_LUMA = (0.5, 0.9)
"""Яркость серой заглушки: светлее фона, темнее белого."""

PLACEHOLDER_GEOMS = {"ellipse", "roundRect", "round1Rect", "round2SameRect", "round2DiagRect"}
"""Круг, овал и скруглённый прямоугольник: в них ставят фото."""

HINT_CHARS = 30
"""Длина текста-подсказки внутри заглушки."""

PLACEHOLDER_MIN_AREA = 0.001
"""Доля слайда, меньше которой фигура это значок оформления, а не место под фото."""

FLAT_VARIANCE = 600
"""Разброс цвета картинки, ниже которого на ней нет деталей."""

PLACEHOLDER_SHARE = 0.15
"""Доля слайда под заглушкой, с которой слайд без своей картинки выглядит недоделанным."""

LAYOUT_PICTURE_MIN_AREA = 0.05
"""Доля кадра, с которой картинка макета считается иллюстрацией, а не значком."""

LAYOUT_BACKDROP_AREA = 0.5
"""Доля кадра, выше которой картинка макета это подсветка фона: образец ставит текст прямо на неё."""

LAYOUT_PICTURE_GAP = 0.02
"""Зазор между текстом и картинкой макета, доля ширины кадра."""

PLATE_PAD = 0.02
"""Отступ текста от нижнего края плашки, доля высоты кадра."""

UNIT_DECOR_INSIDE = 0.9
"""Доля фигуры оформления внутри блока, с которой она принадлежит блоку."""

TEAM_CAPTION_CHARS = 40
"""Длина подписи блока, которая читается как имя и должность, а не как абзац."""

LINE_THIN = 0.01
"""Толщина фигуры в долях слайда, ниже которой это линия, а не фигура."""

LINE_GEOMS = ("line", "straightConnector", "bentConnector", "curvedConnector")
"""Имена готовых форм линий и соединителей."""

QUOTE_MARKS = "«“„\"'"

TITLE_PH = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)

_PH_ROLE = {
    PP_PLACEHOLDER.TITLE: "title",
    PP_PLACEHOLDER.CENTER_TITLE: "title",
    PP_PLACEHOLDER.VERTICAL_TITLE: "title",
    PP_PLACEHOLDER.SUBTITLE: "subtitle",
    PP_PLACEHOLDER.BODY: "body",
    PP_PLACEHOLDER.VERTICAL_BODY: "body",
    PP_PLACEHOLDER.OBJECT: "body",
    PP_PLACEHOLDER.FOOTER: "footer",
    PP_PLACEHOLDER.DATE: "footer",
    PP_PLACEHOLDER.SLIDE_NUMBER: "footer",
}


# ---------- плоский список фигур ----------

@dataclass
class ShapeInfo:
    """Фигура слайда после раскрытия групп: рамка в долях слайда, текст, стиль."""

    shape_id: int
    kind: str
    box: geo.Box
    text: str = ""
    size_pt: float = DEFAULT_SIZE_PT
    tail_size_pt: float | None = None
    """Кегль текста после первого переноса строки, если он задан у самого фрагмента."""
    style: TextStyle = field(default_factory=TextStyle)
    placeholder: str | None = None
    image_part: str | None = None
    fill_color: str | None = None
    geom: str | None = None
    flat_image: bool = False

    @property
    def has_text(self) -> bool:
        return bool(self.text.strip())


def _compose(outer: tuple, inner: tuple) -> tuple:
    return (
        outer[0] * inner[0],
        outer[1] * inner[1],
        outer[0] * inner[2] + outer[2],
        outer[1] * inner[3] + outer[3],
    )


def _group_transform(group) -> tuple:
    """Перевод координат детей группы в координаты слайда."""
    props = group._element.find(qn("p:grpSpPr"))
    xfrm = props.find(qn("a:xfrm")) if props is not None else None
    if xfrm is None:
        return IDENTITY
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    if off is None or ext is None or ch_off is None or ch_ext is None:
        return IDENTITY
    ch_cx = int(ch_ext.get("cx") or 0) or 1
    ch_cy = int(ch_ext.get("cy") or 0) or 1
    sx = int(ext.get("cx") or 0) / ch_cx
    sy = int(ext.get("cy") or 0) / ch_cy
    tx = int(off.get("x") or 0) - int(ch_off.get("x") or 0) * sx
    ty = int(off.get("y") or 0) - int(ch_off.get("y") or 0) * sy
    return (sx, sy, tx, ty)


def _shape_kind(shape) -> str:
    if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
        return "image"
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        return "table"
    if shape.has_text_frame:
        return "text"
    return "shape"


def _shape_box(shape, transform: tuple, slide_size: tuple[int, int]) -> geo.Box:
    sw, sh = slide_size
    left = int(shape.left or 0)
    top = int(shape.top or 0)
    width = int(shape.width or 0)
    height = int(shape.height or 0)
    x = transform[0] * left + transform[2]
    y = transform[1] * top + transform[3]
    w = transform[0] * width
    h = transform[1] * height
    return (x / sw, y / sh, w / sw, h / sh)


def flatten_shapes(container, slide_size, theme, transform=IDENTITY, out=None) -> list[ShapeInfo]:
    """Плоский список фигур с раскрытием групп и пересчётом координат."""
    if out is None:
        out = []
    for shape in container:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            flatten_shapes(
                shape.shapes, slide_size, theme, _compose(transform, _group_transform(shape)), out
            )
            continue
        box = _shape_box(shape, transform, slide_size)
        kind = _shape_kind(shape)
        text = shape.text_frame.text.strip() if shape.has_text_frame else ""
        size = _font_size(shape)
        info = ShapeInfo(
            shape_id=shape.shape_id,
            kind=kind,
            box=box,
            text=text[:300],
            size_pt=size,
            tail_size_pt=_tail_size(shape),
            style=_text_style(shape, size, theme),
            placeholder=_placeholder_role(shape),
            image_part=_blip_part_name(shape),
            fill_color=_shape_fill(shape, theme),
            geom=_shape_geom(shape),
            flat_image=_flat_gray_image(shape) if kind == "image" else False,
        )
        out.append(info)
    return out


def _shape_geom(shape) -> str | None:
    """Имя готовой формы фигуры: rect, ellipse, roundRect."""
    props = shape._element.find(qn("p:spPr"))
    if props is None:
        return None
    node = props.find(qn("a:prstGeom"))
    return node.get("prst") if node is not None else None


def _shape_fill(shape, theme: ThemeInfo) -> str | None:
    """Цвет собственной заливки фигуры; картинка в заливке здесь не разбирается."""
    props = shape._element.find(qn("p:spPr"))
    if props is None or props.find(qn("a:noFill")) is not None:
        return None
    color, _ = _fill_to_color(props, None, theme)
    return color


def _placeholder_role(shape) -> str | None:
    if not shape.is_placeholder:
        return None
    try:
        return _PH_ROLE.get(shape.placeholder_format.type, "other")
    except (KeyError, ValueError):
        return "other"


def _blip_part_name(shape) -> str | None:
    """Имя части пакета с картинкой: по нему слой ассетов узнаёт свой файл."""
    if shape.shape_type not in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
        return None
    try:
        return Path(str(shape.part.related_part(shape._element.blip_rId).partname)).stem
    except Exception:
        return None


# ---------- кегль и стиль ----------

def _xml_size(el) -> float | None:
    if el is None:
        return None
    raw = el.get("sz")
    return int(raw) / 100 if raw else None


def _own_size(shape) -> float | None:
    if not shape.has_text_frame:
        return None
    body = shape.text_frame._txBody
    for para in body.findall(qn("a:p")):
        for run in para.findall(qn("a:r")):
            size = _xml_size(run.find(qn("a:rPr")))
            if size:
                return size
        props = para.find(qn("a:pPr"))
        if props is not None:
            size = _xml_size(props.find(qn("a:defRPr")))
            if size:
                return size
    styles = body.find(qn("a:lstStyle"))
    if styles is not None:
        level = styles.find(qn("a:lvl1pPr"))
        if level is not None:
            size = _xml_size(level.find(qn("a:defRPr")))
            if size:
                return size
    return None


def _inherited_size(shape) -> float | None:
    """Кегль от заполнителя макета, затем мастера, затем общих стилей мастера."""
    if not shape.is_placeholder:
        return None
    try:
        fmt = shape.placeholder_format
        idx, ph_type = fmt.idx, fmt.type
    except (KeyError, ValueError):
        return None
    layout = getattr(shape.part, "slide_layout", None)
    master = getattr(layout, "slide_master", None)
    for holder in (layout, master):
        if holder is None:
            continue
        for candidate in holder.placeholders:
            other = candidate.placeholder_format
            if other.idx == idx or other.type == ph_type:
                size = _own_size(candidate)
                if size:
                    return size
    if master is not None:
        styles = master._element.find(qn("p:txStyles"))
        if styles is not None:
            name = "p:titleStyle" if ph_type in TITLE_PH else "p:bodyStyle"
            node = styles.find(qn(name))
            if node is not None:
                level = node.find(qn("a:lvl1pPr"))
                if level is not None:
                    size = _xml_size(level.find(qn("a:defRPr")))
                    if size:
                        return size
    return None


def _autofit_scale(shape) -> float:
    if not shape.has_text_frame:
        return 1.0
    body_props = shape.text_frame._txBody.find(qn("a:bodyPr"))
    if body_props is None:
        return 1.0
    fit = body_props.find(qn("a:normAutofit"))
    if fit is None:
        return 1.0
    raw = fit.get("fontScale")
    return int(raw) / 100000 if raw else 1.0


def _font_size(shape) -> float:
    size = _own_size(shape)
    if size is None:
        size = _inherited_size(shape)
    if size is None:
        size = DEFAULT_SIZE_PT
    return round(size * _autofit_scale(shape), 1)


def _tail_size(shape) -> float | None:
    """Кегль первого фрагмента после первого переноса строки: подпись под числом в той же рамке."""
    if not shape.has_text_frame:
        return None
    seen_break = False
    for para in shape.text_frame._txBody.findall(qn("a:p")):
        for node in para:
            if node.tag == qn("a:br"):
                seen_break = True
            elif node.tag == qn("a:r") and seen_break and (node.findtext(qn("a:t")) or "").strip():
                props = node.find(qn("a:rPr"))
                size = props.get("sz") if props is not None else None
                return int(size) / 100 if size else None
        seen_break = True  # следующий абзац тоже начинается с новой строки
    return None


def _first_run_props(shape):
    if not shape.has_text_frame:
        return None, None
    body = shape.text_frame._txBody
    for para in body.findall(qn("a:p")):
        for run in para.findall(qn("a:r")):
            props = run.find(qn("a:rPr"))
            if props is not None:
                return props, para.find(qn("a:pPr"))
        props = para.find(qn("a:pPr"))
        if props is not None:
            return props.find(qn("a:defRPr")), props
    return None, None


_ALIGN = {"l": "left", "ctr": "center", "r": "right", "just": "left", "dist": "left"}


def _text_style(shape, size_pt: float, theme) -> TextStyle:
    if not shape.has_text_frame:
        return TextStyle()
    props, para_props = _first_run_props(shape)
    family = None
    color = None
    bold = italic = False
    if props is not None:
        latin = props.find(qn("a:latin"))
        if latin is not None:
            family = theme.font(latin.get("typeface"))
        color = _fill_color(props, theme)
        bold = props.get("b") == "1"
        italic = props.get("i") == "1"
    align = None
    if para_props is not None:
        align = _ALIGN.get(para_props.get("algn") or "")
    return TextStyle(
        family=family, size_pt=size_pt, bold=bold, italic=italic, color=color, align=align
    )


# ---------- цвета темы и фон ----------

@dataclass
class ThemeInfo:
    colors: dict[str, str] = field(default_factory=dict)
    clr_map: dict[str, str] = field(default_factory=dict)
    major: str | None = None
    minor: str | None = None

    def color(self, name: str | None) -> str | None:
        if not name:
            return None
        key = self.clr_map.get(name, name)
        return self.colors.get(key)

    def font(self, typeface: str | None) -> str | None:
        if not typeface:
            return None
        if typeface.startswith("+mj"):
            return self.major
        if typeface.startswith("+mn"):
            return self.minor
        return typeface


def _theme_of(master) -> ThemeInfo:
    info = ThemeInfo()
    clr_map = master._element.find(qn("p:clrMap"))
    if clr_map is not None:
        info.clr_map = dict(clr_map.attrib)
    part = None
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            break
    if part is None:
        return info
    try:
        root = etree.fromstring(part.blob)
    except Exception:
        return info
    scheme = root.find(".//" + qn("a:clrScheme"))
    if scheme is not None:
        for child in scheme:
            name = child.tag.split("}")[-1]
            srgb = child.find(qn("a:srgbClr"))
            system = child.find(qn("a:sysClr"))
            if srgb is not None:
                info.colors[name] = srgb.get("val", "").upper()
            elif system is not None:
                info.colors[name] = (system.get("lastClr") or "").upper()
    fonts = root.find(".//" + qn("a:fontScheme"))
    if fonts is not None:
        major = fonts.find(qn("a:majorFont"))
        minor = fonts.find(qn("a:minorFont"))
        if major is not None and major.find(qn("a:latin")) is not None:
            info.major = major.find(qn("a:latin")).get("typeface")
        if minor is not None and minor.find(qn("a:latin")) is not None:
            info.minor = minor.find(qn("a:latin")).get("typeface")
    return info


def _apply_mods(rgb: tuple[int, int, int], el) -> tuple[int, int, int]:
    r, g, b = rgb
    for mod in el:
        name = mod.tag.split("}")[-1]
        raw = mod.get("val")
        if not raw:
            continue
        value = int(raw) / 100000
        if name in ("shade", "lumMod"):
            r, g, b = r * value, g * value, b * value
        elif name == "tint":
            r = r + (255 - r) * (1 - value)
            g = g + (255 - g) * (1 - value)
            b = b + (255 - b) * (1 - value)
        elif name == "lumOff":
            r, g, b = r + 255 * value, g + 255 * value, b + 255 * value
    clamp = lambda v: max(0, min(255, int(round(v))))
    return clamp(r), clamp(g), clamp(b)


def _color_of(el, theme: ThemeInfo) -> str | None:
    """Шесть шестнадцатеричных знаков из узла цвета с учётом правок яркости."""
    if el is None:
        return None
    srgb = el.find(qn("a:srgbClr"))
    if srgb is not None:
        return _hex(_apply_mods(_rgb(srgb.get("val")), srgb))
    scheme = el.find(qn("a:schemeClr"))
    if scheme is not None:
        base = theme.color(scheme.get("val"))
        if base:
            return _hex(_apply_mods(_rgb(base), scheme))
    system = el.find(qn("a:sysClr"))
    if system is not None:
        return (system.get("lastClr") or "").upper() or None
    return None


def _rgb(value: str | None) -> tuple[int, int, int]:
    raw = (value or "000000").upper()
    return int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)


def _hex(rgb: tuple[int, int, int]) -> str:
    return "%02X%02X%02X" % rgb


def _fill_color(el, theme: ThemeInfo) -> str | None:
    solid = el.find(qn("a:solidFill"))
    if solid is not None:
        return _color_of(solid, theme)
    return None


def _luma(color: str) -> float:
    r, g, b = _rgb(color)
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255


def _image_average(part) -> str | None:
    try:
        from PIL import Image

        with Image.open(BytesIO(part.blob)) as img:
            small = img.convert("RGB").resize((1, 1))
            return _hex(small.getpixel((0, 0)))
    except Exception:
        return None


def _saturation(color: str) -> float:
    r, g, b = _rgb(color)
    top = max(r, g, b)
    return 0.0 if top == 0 else (top - min(r, g, b)) / top


def _is_gray(color: str | None) -> bool:
    """Серая заливка заглушки: цвета нет, яркость между фоном и белым."""
    if not color:
        return False
    return _saturation(color) < GRAY_SAT and GRAY_LUMA[0] <= _luma(color) <= GRAY_LUMA[1]


_FLAT_IMAGES: dict[str, bool] = {}


def _flat_gray_image(shape) -> bool:
    """Картинка одного серого тона без деталей: её ставят вместо будущего фото."""
    try:
        image = shape.image
        key = image.sha1
    except Exception:
        return False
    if key in _FLAT_IMAGES:
        return _FLAT_IMAGES[key]
    value = False
    try:
        from PIL import Image

        with Image.open(BytesIO(image.blob)) as img:
            data = img.convert("RGB").resize((16, 16)).tobytes()
        pixels = [tuple(data[i:i + 3]) for i in range(0, len(data), 3)]
        avg = tuple(sum(p[i] for p in pixels) // len(pixels) for i in range(3))
        spread = sum(
            sum((p[i] - avg[i]) ** 2 for i in range(3)) for p in pixels
        ) / len(pixels)
        value = spread <= FLAT_VARIANCE and _is_gray(_hex(avg))
    except Exception:
        value = False
    if len(_FLAT_IMAGES) > 512:
        _FLAT_IMAGES.clear()
    _FLAT_IMAGES[key] = value
    return value


def _is_photo_placeholder(info: ShapeInfo) -> bool:
    """Серая заглушка под фото: круг, овал или скруглённый прямоугольник без своего содержания."""
    if geo.area(info.box) < PLACEHOLDER_MIN_AREA:
        return False  # точка списка и тонкая линия под заглушку не годятся
    if info.kind == "image":
        return info.flat_image
    if info.geom not in PLACEHOLDER_GEOMS:
        return False
    if len(info.text.strip()) > HINT_CHARS or _is_number(info.text):
        return False  # число внутри фигуры это содержание слайда, а не подсказка
    return _is_gray(info.fill_color)


def _fill_to_color(node, part, theme: ThemeInfo) -> tuple[str | None, str | None]:
    """Из узла заливки вернуть (цвет, имя картинки)."""
    if node is None:
        return None, None
    solid = node.find(qn("a:solidFill"))
    if solid is not None:
        return _color_of(solid, theme), None
    grad = node.find(qn("a:gradFill"))
    if grad is not None:
        stops = [_color_of(gs, theme) for gs in grad.iter(qn("a:gs"))]
        stops = [s for s in stops if s]
        if stops:
            mixed = tuple(
                sum(_rgb(s)[i] for s in stops) // len(stops) for i in range(3)
            )
            return _hex(mixed), None
    blip = node.find(qn("a:blipFill"))
    if blip is not None and part is not None:
        embed = blip.find(qn("a:blip"))
        rel_id = embed.get(qn("r:embed")) if embed is not None else None
        if rel_id:
            try:
                image_part = part.related_part(rel_id)
            except (KeyError, AttributeError):
                return None, None
            return _image_average(image_part), hashlib.sha1(image_part.blob).hexdigest()[:12]
    ref = node.find(qn("p:bgRef"))
    if ref is not None:
        return _color_of(ref, theme), None
    return None, None


def _background_node(part):
    bg = part._element.find(qn("p:cSld"))
    return bg.find(qn("p:bg")) if bg is not None else None


def _background(slide, theme: ThemeInfo) -> tuple[str | None, str | None]:
    layout = slide.slide_layout
    master = layout.slide_master
    for holder in (slide, layout, master):
        node = _background_node(holder)
        if node is None:
            continue
        props = node.find(qn("p:bgPr"))
        color, asset = _fill_to_color(props if props is not None else node, holder.part, theme)
        if color:
            return color, asset
    return None, None


def _full_bleed_fill(shapes: list[ShapeInfo], slide, theme: ThemeInfo) -> tuple[str | None, str | None]:
    """Подложка на весь кадр: картинка или залитая фигура поверх фона."""
    covering = [
        s
        for s in shapes
        if s.box[2] >= geo.FULL_BLEED and s.box[3] >= geo.FULL_BLEED and abs(s.box[0]) < 0.05
    ]
    if not covering:
        return None, None
    by_id = {s.shape_id: s for s in covering}
    for shape in slide.shapes:
        info = by_id.get(shape.shape_id)
        if info is None:
            continue
        if info.kind == "image":
            try:
                # Тот же идентификатор, что у Asset.id в слое ассетов: первые 12 знаков sha1 содержимого.
                return _image_average(shape.image), shape.image.sha1[:12]
            except Exception:
                continue
        props = shape._element.find(qn("p:spPr"))
        color, asset = _fill_to_color(props, slide.part, theme)
        if color:
            return color, asset
    return None, None


# ---------- слоты, области, блоки ----------

def _capacity(box: geo.Box, size_pt: float, slide_pt: tuple[float, float]) -> tuple[int, int]:
    width_pt = box[2] * slide_pt[0]
    height_pt = box[3] * slide_pt[1]
    line = max(size_pt * LINE_HEIGHT, 1.0)
    lines = max(1, int(height_pt / line))
    per_line = max(1, int(width_pt / max(size_pt * CHAR_WIDTH, 0.1)))
    return per_line * lines, lines


def _is_number(text: str) -> bool:
    value = text.strip()
    return bool(value) and len(value) <= NUMBER_MAX_CHARS and bool(
        NUMBER_RE.match(value) or NUMBER_MASK_RE.match(value)
    )


def _is_captioned_number(info: ShapeInfo) -> bool:
    """Число и подпись в одной рамке: первая строка число, после переноса подпись вдвое мельче.

    Так в шаблоне бывает нарисован слайд с крупным числом: «ххх%», перенос, «данные показателя».
    Подпись того же кегля, что и число («2019», перенос, «Запуск»), это пункт хронологии.
    """
    parts = re.split(r"[\v\n]", info.text.strip(), maxsplit=1)
    if len(parts) < 2 or not parts[1].strip() or not _is_number(parts[0]) or not info.tail_size_pt:
        return False
    return info.size_pt >= info.tail_size_pt * NUMBER_RATIO


def _numeric(info: ShapeInfo) -> bool:
    return _is_number(info.text) or _is_captioned_number(info)


def _body_size(items: list[ShapeInfo]) -> float:
    """Кегль основного текста слайда: середина кеглей нечисловых подписей."""
    sizes = sorted(i.size_pt for i in items if i.has_text and not _numeric(i))
    if not sizes:
        return DEFAULT_SIZE_PT
    return sizes[len(sizes) // 2]


def _is_big_number(slot: Slot, body_pt: float) -> bool:
    """Крупное число: слот с числом и кеглем не меньше двух кеглей основного текста слайда."""
    return slot.role == "number" and (slot.style.size_pt or 0) >= body_pt * NUMBER_RATIO


def _is_text_slot(info: ShapeInfo) -> bool:
    if info.kind != "text":
        return False
    if _is_photo_placeholder(info):
        return False  # короткий текст внутри серой фигуры это подсказка, а не место под текст
    return info.has_text or info.placeholder in ("title", "subtitle", "body")


def _make_slot(info: ShapeInfo, role: str, slide_pt, origin: geo.Box | None = None) -> Slot:
    box = info.box
    if origin is not None:
        box = (box[0] - origin[0], box[1] - origin[1], box[2], box[3])
    chars, lines = _capacity(info.box, info.size_pt, slide_pt)
    return Slot(
        id=f"s{info.shape_id}",
        role=role,
        shape_id=info.shape_id,
        box=box,
        style=info.style,
        max_chars=chars,
        max_lines=lines,
        sample_text=info.text,
    )


def _slide_roles(items: list[ShapeInfo]) -> dict[int, str]:
    """Роли текстовых фигур вне блоков."""
    roles: dict[int, str] = {}
    if not items:
        return roles
    sizes = sorted((i.size_pt for i in items), reverse=True)
    top_size = sizes[0]
    title_id = None
    holders = [i for i in items if i.placeholder == "title"]
    if holders:
        title_id = min(holders, key=lambda i: i.box[1]).shape_id
    else:
        upper = [i for i in items if i.box[1] < 0.5 and len(i.text) <= 120]
        if upper:
            best = max(upper, key=lambda i: (i.size_pt, -i.box[1]))
            if best.size_pt >= top_size * 0.9:
                title_id = best.shape_id
    for info in items:
        if info.shape_id == title_id:
            roles[info.shape_id] = "title"
        elif info.placeholder == "subtitle":
            roles[info.shape_id] = "subtitle"
        elif info.placeholder == "footer" or (info.box[1] > 0.86 and info.size_pt <= 14):
            roles[info.shape_id] = "footer"
        elif _numeric(info):
            roles[info.shape_id] = "number"
        elif info.size_pt >= top_size * 0.7:
            roles[info.shape_id] = "heading"
        elif info.size_pt <= 12:
            roles[info.shape_id] = "caption"
        else:
            roles[info.shape_id] = "body"
    return roles


def _unit_roles(items: list[ShapeInfo]) -> dict[int, str]:
    """Роли внутри блока: heading по кеглю и порядку, остальное body."""
    roles: dict[int, str] = {}
    if not items:
        return roles
    ordered = sorted(items, key=lambda i: (round(i.box[1], 3), i.box[0]))
    top_size = max(i.size_pt for i in ordered)
    head_done = False
    for info in ordered:
        if _numeric(info):
            roles[info.shape_id] = "number"
        elif not head_done and info.size_pt >= top_size - 0.01:
            roles[info.shape_id] = "heading"
            head_done = True
        else:
            roles[info.shape_id] = "body"
    return roles


def _reading_order(info: ShapeInfo) -> tuple[float, float]:
    return (round(info.box[1], 2), round(info.box[0], 2))


def _area_kind(info: ShapeInfo) -> str:
    if info.kind in ("chart", "table"):
        return info.kind
    if _is_photo_placeholder(info):
        return "image"
    if info.box[2] <= ICON_MAX[0] and info.box[3] <= ICON_MAX[1]:
        return "icon"
    return "image"


def _is_area(info: ShapeInfo) -> bool:
    return info.kind in ("image", "chart", "table") or _is_photo_placeholder(info)


def _decor_kind(info: ShapeInfo) -> str:
    """Вид фигуры оформления: картинка, линия, текст или фигура."""
    if info.kind == "image":
        return "image"
    long_side, thin_side = max(info.box[2], info.box[3]), min(info.box[2], info.box[3])
    if (info.geom or "").startswith(LINE_GEOMS) or (thin_side <= LINE_THIN < long_side):
        return "line"
    if info.has_text:
        return "text"
    return "shape" if info.kind in ("shape", "text") else "other"


def _make_decor(items: list[ShapeInfo]) -> list[DecorShape]:
    """Фигуры оформления с рамками: по ним вёрстка решает, мешает ли оформление содержимому."""
    return [
        DecorShape(shape_id=i.shape_id, box=i.box, kind=_decor_kind(i))
        for i in items
    ]


def _bar_groups(items: list[ShapeInfo], axis: int) -> list[list[ShapeInfo]]:
    """Фигуры одной толщины по стороне axis: полосы одного ряда не разъезжаются по сотым."""
    out: list[list[ShapeInfo]] = []
    for info in sorted(items, key=lambda i: -i.box[axis]):
        for group in out:
            head = group[0].box[axis]
            if head > 0 and abs(head - info.box[axis]) <= geo.BAR_SIZE_TOL * head:
                group.append(info)
                break
        else:
            out.append([info])
    return [group for group in out if len(group) >= geo.BAR_MIN]


def _chart_sample(items: list[ShapeInfo]) -> tuple[Area, set[int]] | None:
    """Ряд цветных полос или столбиков разной длины на общей оси: образец диаграммы из фигур."""
    painted = [i for i in items if i.kind == "image" or i.fill_color]
    for axis in (3, 2):
        for group in _bar_groups(painted, axis):
            box = geo.chart_sample([i.box for i in group])
            if box is not None:
                first = min(i.shape_id for i in group)
                return Area(id=f"c{first}", kind="chart", box=box), {i.shape_id for i in group}
    return None


def _make_area(info: ShapeInfo, origin: geo.Box | None = None) -> Area:
    box = info.box
    if origin is not None:
        box = (box[0] - origin[0], box[1] - origin[1], box[2], box[3])
    return Area(
        id=f"a{info.shape_id}",
        kind=_area_kind(info),
        box=box,
        shape_id=info.shape_id,
        placeholder=_is_photo_placeholder(info),
    )


# ---------- тип слайда ----------

@dataclass
class SlideFacts:
    slots: list[Slot]
    areas: list[Area]
    groups: list[RepeatGroup]
    texts: list[ShapeInfo]
    decor: list[geo.Box] = field(default_factory=list)
    body_pt: float = DEFAULT_SIZE_PT

    @property
    def has_axis(self) -> bool:
        """Длинная тонкая фигура оформления: ось таймлайна."""
        return any(b[2] >= 0.5 and b[3] <= 0.03 for b in self.decor)

    @property
    def big_numbers(self) -> int:
        """Сколько на слайде крупных чисел вместе с числами внутри блоков."""
        free = sum(1 for s in self.slots if _is_big_number(s, self.body_pt))
        in_units = sum(
            len([s for s in g.unit_slots if _is_big_number(s, self.body_pt)]) * len(g.units)
            for g in self.groups
        )
        return free + in_units


def _holds_a_number(facts: SlideFacts) -> bool:
    """Слайд держится на числе: крупное число занимает заметное место и их не больше трёх."""
    if not 1 <= facts.big_numbers <= BIG_NUMBER_MAX:
        return False
    top_pt = max((s.style.size_pt or 0 for s in facts.slots), default=0)
    return any(
        _is_big_number(s, facts.body_pt) and geo.area(s.box) >= BIG_NUMBER_AREA
        and (s.style.size_pt or 0) >= top_pt
        for s in facts.slots
    )


def _short_captions(slots: list[Slot]) -> bool:
    """Подписи блока короткие: имя и должность, а не абзац текста."""
    return all(len(s.sample_text) <= TEAM_CAPTION_CHARS for s in slots)


def _classify(facts: SlideFacts) -> tuple[SlideKind, float]:
    areas = facts.areas
    groups = facts.groups
    if any(a.kind == "table" for a in areas):
        return SlideKind.table, 0.95
    if any(a.kind == "chart" for a in areas):
        return SlideKind.chart, 0.9
    if _holds_a_number(facts):
        return SlideKind.big_number, 0.8

    sizes = [t.size_pt for t in facts.texts] or [DEFAULT_SIZE_PT]
    top_size = max(sizes)

    if groups:
        group = max(groups, key=lambda g: len(g.units))
        # Номера и фото связанной группы принадлежат тому же смысловому блоку:
        # ряд кружков с цифрами и ряд подписей под ними это одни и те же шаги.
        linked = [g for g in groups if g.id in group.linked_group_ids]
        count = len(group.units)
        heads = [s for s in group.unit_slots if s.role in ("heading", "title")]
        captions = group.unit_slots + [s for g in linked for s in g.unit_slots]
        numbers = [s for g in [group, *linked] for s in g.unit_slots if s.role == "number"]
        photos = [a for g in [group, *linked] for a in g.unit_areas if a.kind == "image"]
        holders = [a for a in photos if a.placeholder]
        if count >= 3 and holders and _short_captions(captions):
            return SlideKind.team, 0.8
        if count >= 3 and numbers:
            return SlideKind.steps, 0.8
        if count >= 3 and photos and all(abs(a.box[2] - a.box[3]) < 0.08 for a in photos):
            return SlideKind.team, 0.7
        if count >= 3 and len(group.unit_slots) >= 2 and heads:
            return SlideKind.cards, 0.8
        if count == 2:
            return SlideKind.compare, 0.6
        if count >= 3 and group.direction == "row" and facts.has_axis:
            return SlideKind.timeline, 0.6
        if count >= 3 and group.direction == "column":
            return SlideKind.bullets, 0.6
        if count >= 3:
            return SlideKind.cards, 0.5
        return SlideKind.cards, 0.4

    quotes = [t for t in facts.texts if t.text[:1] in QUOTE_MARKS and len(t.text) > 30]
    if quotes:
        return SlideKind.quote, 0.75

    big_images = [a for a in areas if a.kind == "image" and geo.area(a.box) >= 0.12]
    body_slots = [s for s in facts.slots if s.role in ("body", "heading", "caption")]
    if big_images and body_slots:
        return SlideKind.image_text, 0.65

    columns = [s for s in facts.slots if s.role in ("body", "heading")]
    if len(columns) >= 3 and len({round(s.box[0], 2) for s in columns}) == 1:
        return SlideKind.bullets, 0.6

    title = next((s for s in facts.slots if s.role == "title"), None)
    if title is not None and top_size >= 28:
        rest = [s for s in facts.slots if s is not title]
        if not rest:
            return SlideKind.section, 0.5
        if len(rest) <= 3 and all(s.role in ("subtitle", "caption", "footer", "body") for s in rest):
            return SlideKind.title, 0.5
    return SlideKind.other, 0.2


# ---------- сборка паттерна ----------

def _build_group(cand: geo.RepeatCandidate, by_id: dict[int, ShapeInfo], slide_pt, index: int) -> RepeatGroup:
    boxes = cand.unit_boxes
    origin = boxes[0]
    units = [
        RepeatUnit(index=i, box=box, shape_ids=[f.key for f in unit])
        for i, (box, unit) in enumerate(zip(boxes, cand.units))
    ]
    first = sorted((by_id[f.key] for f in cand.units[0]), key=_reading_order)
    texts = [i for i in first if _is_text_slot(i)]
    roles = _unit_roles(texts)
    unit_slots = [_make_slot(i, roles[i.shape_id], slide_pt, origin) for i in texts]
    unit_areas = [_make_area(i, origin) for i in first if _is_area(i)]
    return RepeatGroup(
        id=f"g{index}",
        direction=cand.direction,
        cols=cand.cols,
        rows=cand.rows,
        step=cand.step,
        unit_size=cand.unit_size,
        units=units,
        unit_slots=unit_slots,
        unit_areas=unit_areas,
        max_units=len(units),
    )


def _link_groups(groups: list[RepeatGroup]) -> None:
    """Связать группы, которые повторяются синхронно: ряд номеров и ряд подписей под ними."""
    for i, first in enumerate(groups):
        for second in groups[i + 1:]:
            if (first.direction, first.cols, first.rows) != (second.direction, second.cols, second.rows):
                continue
            if not geo.in_lockstep([u.box for u in first.units], [u.box for u in second.units]):
                continue
            first.linked_group_ids.append(second.id)
            second.linked_group_ids.append(first.id)


def _primary_group_id(groups: list[RepeatGroup]) -> str | None:
    """Главная группа: больше текстовых слотов в блоке, при равенстве больше площадь блока."""
    if not groups:
        return None
    best = max(groups, key=lambda g: (len(g.unit_slots), g.unit_size[0] * g.unit_size[1]))
    return best.id


def _needs_images(groups: list[RepeatGroup], areas: list[Area], primary_id: str | None) -> bool:
    """Паттерн держится на фото: заглушка в каждом блоке главной группы или крупная заглушка."""
    primary = next((g for g in groups if g.id == primary_id), None)
    if primary is not None and any(a.placeholder for a in primary.unit_areas):
        return True
    return any(a.placeholder and geo.area(a.box) > PLACEHOLDER_SHARE for a in areas)


def _layout_pictures(slide, slide_size) -> list[geo.Box]:
    """Картинки макета слайда не на весь кадр."""
    return _pictures_of(slide.slide_layout, slide_size)


def _pictures_of(layout, slide_size) -> list[geo.Box]:
    """Картинки макета не на весь кадр: заполнитель шаблона бывает шире них, а текст на них не читается."""
    boxes = []
    for shape in layout.shapes:
        if shape.is_placeholder or _shape_kind(shape) != "image":
            continue
        box = _shape_box(shape, IDENTITY, slide_size)
        if box[2] < geo.FULL_BLEED or box[3] < geo.FULL_BLEED:
            boxes.append(box)
    return boxes


def _attach_to_units(groups: list[RepeatGroup], decor: list[ShapeInfo]) -> list[ShapeInfo]:
    """Оформление внутри блока едет вместе с блоком; возвращает то, что осталось оформлением слайда.

    Повтор собирается из фигур, которые есть в каждом блоке. Плашка заголовка, которой нет
    у одного из блоков (последняя карточка залита целиком), в повтор не попадает и без этой
    правки остаётся на старом месте, когда блоки перестраиваются, и ложится на чужой заголовок.
    """
    rest = []
    for info in decor:
        unit = next((u for g in groups for u in g.units
                     if geo.covered(info.box, u.box) >= UNIT_DECOR_INSIDE
                     and geo.area(info.box) < geo.area(u.box)), None)
        if unit is None:
            rest.append(info)
        else:
            unit.shape_ids.append(info.shape_id)
    return rest


def _clip_to_plate(slot: Slot, decor: list[ShapeInfo], slide_pt) -> Slot:
    """Слот, который начинается на плашке и выходит за её нижний край, укорачивается до плашки.

    Образец держит на плашке строку-другую, а рамка текста у него уходит ниже плашки:
    по такой рамке подгонка разрешает лишнюю строку, и она ложится под край плашки.
    """
    x, y, w, h = slot.box
    bottom = y + h
    for info in decor:
        if info.kind not in ("shape", "text") or info.has_text:
            continue
        px, py, pw, ph = info.box
        if pw >= geo.FULL_BLEED and ph >= geo.FULL_BLEED:
            continue  # подложка на весь кадр это фон
        inside = px - 1e-3 <= x and x + w <= px + pw + 1e-3 and py <= y < py + ph
        if inside and y + h > py + ph + 1e-3:
            bottom = min(bottom, py + ph - PLATE_PAD)
    size = slot.style.size_pt or DEFAULT_SIZE_PT
    if bottom >= y + h - 1e-6 or (bottom - y) * slide_pt[1] < 2 * size * LINE_HEIGHT:
        return slot  # плашка на одну строку: рамка текста шире неё ради отступа, а не ради строк
    box = (x, y, w, bottom - y)
    chars, lines = _capacity(box, size, slide_pt)
    return slot.model_copy(update={"box": box, "max_chars": chars, "max_lines": lines})


def _clip_to_layout(slot: Slot, slide, slide_size) -> Slot:
    """Слот сужается до левого края картинки макета, на которую он заходит справа.

    Заполнитель титула в шаблоне бывает на всю ширину, а справа в макете стоит иллюстрация:
    короткий текст образца до неё не доходит, длинный лёг бы на неё.
    """
    x, y, w, h = slot.box
    right = x + w
    for px, py, pw, ph in _layout_pictures(slide, slide_size):
        if py < y + h and py + ph > y and x < px < right and pw * ph > LAYOUT_PICTURE_MIN_AREA:
            right = min(right, px - LAYOUT_PICTURE_GAP)
    if right >= x + w - 1e-6 or right - x < w * 0.3:
        return slot
    return slot.model_copy(update={"box": (x, y, right - x, h)})


def _pattern_of_slide(slide, number: int, theme: ThemeInfo, slide_size, slide_pt) -> Pattern:
    shapes = flatten_shapes(slide.shapes, slide_size, theme)
    by_id = {s.shape_id: s for s in shapes}
    frames = [
        geo.Frame(key=s.shape_id, box=s.box, sig=geo.signature(s.kind, s.box, s.has_text))
        for s in shapes
    ]
    field_box = geo.content_box([s.box for s in shapes])
    cands = geo.find_repeats(frames)
    body_pt = _body_size(shapes)

    groups: list[RepeatGroup] = []
    taken: set[int] = set()
    for cand in cands:
        group = _build_group(cand, by_id, slide_pt, len(groups) + 1)
        if not group.unit_slots and not group.unit_areas:
            continue  # повтор без содержимого это оформление, а не блок
        group.max_units = geo.fit_units(
            cand.unit_boxes[0], cand.unit_size, cand.step, cand.cols, cand.rows, field_box
        )
        groups.append(group)
        taken.update(cand.keys)
    _link_groups(groups)
    primary_id = _primary_group_id(groups)

    free = sorted((s for s in shapes if s.shape_id not in taken), key=_reading_order)
    texts = [s for s in free if _is_text_slot(s)]
    roles = _slide_roles(texts)
    slots = [_clip_to_layout(_make_slot(s, roles[s.shape_id], slide_pt), slide, slide_size) for s in texts]

    plain = [s for s in free if not _is_text_slot(s) and not _is_photo_placeholder(s)]
    sample = _chart_sample(plain)
    bars = sample[1] if sample is not None else set()
    areas = [_make_area(s) for s in free if _is_area(s) and s.shape_id not in bars]
    if sample is not None:
        # Полосы образца сами по себе не место под картинку: они и есть диаграмма.
        areas.append(sample[0])

    used = taken | {s.shape_id for s in texts} | {a.shape_id for a in areas if a.shape_id}
    decor = _attach_to_units(groups, [s for s in shapes if s.shape_id not in used])
    slots = [_clip_to_plate(slot, decor, slide_pt) for slot in slots]

    color, asset = _background(slide, theme)
    bleed_color, bleed_asset = _full_bleed_fill(shapes, slide, theme)
    if bleed_color:
        color, asset = bleed_color, bleed_asset or asset
    dark = color is not None and _luma(color) < DARK_LUMA

    facts = SlideFacts(slots, areas, groups, texts, [s.box for s in decor], body_pt)
    kind, confidence = _classify(facts)
    return Pattern(
        id=f"p{number:03d}",
        source_slide=number,
        layout_name=slide.slide_layout.name,
        kind=kind,
        kind_confidence=confidence,
        theme="dark" if dark else "light",
        background_asset=asset,
        background_color=color,
        needs_images=_needs_images(groups, areas, primary_id),
        primary_group_id=primary_id,
        slots=slots,
        areas=areas,
        groups=groups,
        decor_shape_ids=[s.shape_id for s in decor],
        decor=_make_decor(decor),
    )


def extract_patterns(pptx_path: Path) -> list[Pattern]:
    """Разобрать каждый слайд презентации в паттерн вёрстки."""
    pres = Presentation(str(pptx_path))
    slide_size = (pres.slide_width or 1, pres.slide_height or 1)
    slide_pt = (slide_size[0] / EMU_PER_PT, slide_size[1] / EMU_PER_PT)
    themes: dict[int, ThemeInfo] = {}
    patterns = []
    for number, slide in enumerate(pres.slides, start=1):
        master = slide.slide_layout.slide_master
        key = id(master)
        if key not in themes:
            themes[key] = _theme_of(master)
        patterns.append(_pattern_of_slide(slide, number, themes[key], slide_size, slide_pt))
    return patterns


def extract_layouts(pptx_path: Path) -> list[LayoutInfo]:
    """Макеты презентации с заполнителями в виде слотов."""
    pres = Presentation(str(pptx_path))
    slide_size = (pres.slide_width or 1, pres.slide_height or 1)
    slide_pt = (slide_size[0] / EMU_PER_PT, slide_size[1] / EMU_PER_PT)
    layouts = []
    for master_index, master in enumerate(pres.slide_masters):
        theme = _theme_of(master)
        for layout in master.slide_layouts:
            holders = []
            for shape in layout.placeholders:
                info = ShapeInfo(
                    shape_id=shape.shape_id,
                    kind=_shape_kind(shape),
                    box=_shape_box(shape, IDENTITY, slide_size),
                    text=(shape.text_frame.text.strip() if shape.has_text_frame else "")[:300],
                    size_pt=_font_size(shape),
                    placeholder=_placeholder_role(shape),
                )
                info.style = _text_style(shape, info.size_pt, theme)
                holders.append(_make_slot(info, info.placeholder or "other", slide_pt))
            picture = any(
                _shape_kind(shape) == "image"
                and (shape.width or 0) >= slide_size[0] * geo.FULL_BLEED
                and (shape.height or 0) >= slide_size[1] * geo.FULL_BLEED
                for shape in layout.shapes
            )
            pictures = [box for box in _pictures_of(layout, slide_size)
                        if LAYOUT_PICTURE_MIN_AREA < box[2] * box[3] <= LAYOUT_BACKDROP_AREA]
            layouts.append(
                LayoutInfo(name=layout.name, master_index=master_index, placeholders=holders,
                           full_bleed_picture=picture, pictures=pictures)
            )
    return layouts
