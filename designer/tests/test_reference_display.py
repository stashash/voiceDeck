import copy
import hashlib
import json
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from designer import pipeline, store
from designer.api.app import app
from designer.contracts import ChartSpec, Series, Deck, DeckPlan, SlideIntent, SlideKind, Item
from designer.export.sources import collect_sources, export_source_package
from designer.insights.suggest import Suggestion, insight_slide
from designer.localization import localize_plan
from designer.reference_display import table_pages, recompose_reference_plan, migrate_deck_reference_display

FIXTURES = Path(__file__).parent / "fixtures" / "semantic_stage2"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv(store.DATA_DIR_ENV, str(tmp_path / "data"))
    return TestClient(app)


def materialized(client, name="reference_201.csv", mode="compact"):
    raw = (FIXTURES / name).read_bytes()
    saved = client.post("/insights/plan", files={"files": (name, raw)}).json()
    ds = saved["datasets"][0]
    candidate = next(c for c in ds["plan"]["candidates"] if c["kind"] == "table")
    response = client.post("/insights/plan/materialize-bundle", data={"bundle_id": saved["id"],
        "selections": json.dumps([{"dataset_id": ds["dataset_id"], "candidate_kind": "table",
                                  "required_fields": candidate["required_fields"], "table_display": mode}])})
    assert response.status_code == 200, response.text
    suggestions = store.load_insights(response.json()["id"])["suggestions"]
    slides = [insight_slide(Suggestion(**s)) for s in suggestions]
    for i, slide in enumerate(slides):
        slide.id = f"r{i}"
    return DeckPlan(title="Test", purpose="test", slides=slides)


def test_compact_is_honest_and_full_is_lossless_without_duplicate_summary(client):
    compact = materialized(client)
    assert len(compact.slides) == 1
    slide = compact.slides[0]
    assert len(slide.table.rows) == 7
    assert "не репрезентативная" in slide.key_message
    assert "7 из 201" in slide.title
    assert slide.provenance.table_display == "compact"
    assert slide.provenance.row_lines == list(range(2, 203))
    localized = localize_plan(compact).slides[0]
    assert len(localized.source_note.splitlines()) == 2
    assert "HTML" in localized.source_note.splitlines()[1]
    assert "архив" in localized.source_note.splitlines()[1]
    assert "показано 7 из 201" in slide.provenance.formula
    full = materialized(client, mode="full")
    assert len(full.slides) == 29
    snapshot = store.load_dataset(slide.provenance.dataset_id)
    assert [r for s in full.slides for r in s.table.rows] == snapshot["table"]["rows"]
    assert [n for s in full.slides for n in s.provenance.displayed_row_lines] == snapshot["table"]["row_lines"]
    assert all(s.kind == SlideKind.table for s in full.slides)


def test_source_rows_are_physical_scope_not_contiguous_assumption():
    lines = [3, 5, 8, 9, 12, 14, 20, 23]
    rows = [[str(i)] for i in range(8)]
    pages = table_pages({"visual": "table"}, {"table": {"rows": rows, "headers": ["id"], "row_lines": lines}})
    assert len(pages) == 2
    assert pages[0]["provenance"]["displayed_row_lines"] == lines[:7]
    assert "3, 5, 8, 9, 12, 14, 20" in pages[0]["coverage"]
    assert pages[1]["provenance"]["displayed_row_lines"] == [23]


def test_missing_zero_is_small_analytical_table_not_lookup(client):
    plan = localize_plan(materialized(client, "missing_zero.csv"))
    assert len(plan.slides) == 1
    slide = plan.slides[0]
    assert slide.title == "Ноль — значение, пропуск — отсутствие данных"
    assert "Справочник" not in slide.title
    assert len(slide.table.rows) == 4
    assert any("" in row for row in slide.table.rows)
    assert any("0" in row for row in slide.table.rows)
    assert slide.provenance.page_count is None


def test_en_compact_has_english_code_owned_text_and_unchanged_evidence(client):
    original = materialized(client)
    plan = localize_plan(original, "en")
    assert plan.slides[0].title == "Source preview · 7 of 201 rows"
    assert "not a representative sample" in plan.slides[0].key_message
    assert plan.slides[0].provenance == original.slides[0].provenance
    assert "Shown: 7 of 201" in plan.slides[0].source_note


def test_offline_exports_contain_exact_source_all_rows_and_search(client, tmp_path):
    plan = materialized(client)
    for name in ("deck.html", "deck.markup.html"):
        (tmp_path / name).write_text("<html><body><main>slides</main></body></html>", encoding="utf-8")
    (tmp_path / "deck.pptx").write_bytes(b"test presentation")
    manifest = export_source_package(plan, tmp_path)
    entry = manifest["datasets"][0]
    raw = (FIXTURES / "reference_201.csv").read_bytes()
    with zipfile.ZipFile(tmp_path / "deck.sources.zip") as archive:
        assert archive.read(entry["original_path"]) == raw
        table = json.loads(archive.read(entry["table_path"]))
        assert len(table["rows"]) == 201
        assert table["row_lines"] == list(range(2, 203))
    with zipfile.ZipFile(tmp_path / "deck.bundle.zip") as archive:
        assert archive.read(entry["original_path"]) == raw
        assert "deck.pptx" in archive.namelist()
        markup = archive.read("deck.html").decode()
        assert 'id="source-filter"' in markup and "REF-201" in markup
        assert markup.count("<tr>") == 202  # full table + header
        assert "localhost" not in markup
        assert "stopPropagation" in markup  # typing does not navigate slides
    response = client.get(f"/datasets/{entry['dataset_id']}/source")
    assert response.content == raw
    assert response.headers["X-Source-SHA256"] == hashlib.sha256(raw).hexdigest()
    assert "attachment" in response.headers["Content-Disposition"]


def test_exports_fail_closed_for_missing_or_corrupted_original(client, tmp_path):
    plan = materialized(client)
    p = plan.slides[0].provenance
    p.source_sha256 = "f" * 64
    with pytest.raises(ValueError, match="checksum"):
        export_source_package(plan, tmp_path)
    assert not (tmp_path / "deck.sources.zip").exists()
    p.dataset_id = "missing"
    with pytest.raises(ValueError, match="missing"):
        collect_sources(plan)


def test_migration_preserves_chart_and_all_other_main_data_and_roundtrips_modes(client):
    plan = localize_plan(materialized(client, mode="full"))
    small = localize_plan(materialized(client, "missing_zero.csv")).slides[0]
    small.id = "small"
    main = SlideIntent(id="main", kind=SlideKind.chart, title="Author title", notes="Author notes",
                       chart=ChartSpec(type="bar", categories=["A", "B"], series=[Series(name="Exact", values=[0, 12.5])]))
    plan.slides = [main, small, *plan.slides]
    before = plan.model_dump()
    compact, origin, removed = recompose_reference_plan(plan)
    assert plan.model_dump() == before
    assert compact.slides[0] == main
    assert len(compact.slides) == 3
    assert compact.slides[1].table == small.table
    full, _, _ = recompose_reference_plan(compact, "full")
    assert len(full.slides) == 31
    assert [row for s in full.slides[2:] for row in s.table.rows] == [row for s in plan.slides[2:] for row in s.table.rows]
    forged = compact.model_copy(deep=True)
    forged.slides[-1].table.rows[0][0] = "edited"
    with pytest.raises(ValueError, match="differ"):
        recompose_reference_plan(forged)


def test_legacy_summary_replaced_in_place_no_empty_slide_and_original_ids_preserved(client):
    plan = localize_plan(materialized(client, mode="full"))
    summary = plan.slides[0].model_copy(deep=True)
    summary.id, summary.kind, summary.table = "summary", SlideKind.big_number, None
    summary.provenance.displayed_row_lines = []
    summary.provenance.page_count = None
    summary.provenance.table_display = None
    last = SlideIntent(id="last", kind=SlideKind.cta, title="Ask a real question")
    plan.slides = [summary, last, *plan.slides]
    compact, origin, removed = recompose_reference_plan(plan)
    assert [s.id for s in compact.slides] == ["summary", "last"]
    assert compact.slides[0].table is not None
    assert compact.slides[1] == last
    assert len(removed) == 30


def test_migration_archive_and_history_are_exact_and_export_requires_no_model(client, tmp_path, monkeypatch):
    from designer.layout.compose import compose
    from designer.layout.scene import build_scene
    monkeypatch.setattr("designer.edit.convert.available", lambda: False)
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[1])
    slide.shapes.title.text = "Template"
    slide.placeholders[1].text = "Body"
    buffer = BytesIO()
    source.save(buffer)
    ds = pipeline.import_template(buffer.getvalue(), "minimal.pptx")
    plan = localize_plan(materialized(client, mode="full"))
    package = store.design_system_dir(ds.id)
    specs = [compose(s, ds.patterns[0], ds) for s in plan.slides]
    scenes = [build_scene(s, ds.patterns[0], ds, package) for s in specs]
    deck_id = store.new_deck_id()
    store.init_deck(deck_id, ds.id, ["a"])
    store.save_deck_result(deck_id, Deck(id=deck_id, design_system_id=ds.id, variant="a",
                          plan=plan, specs=specs, scenes=scenes), [], variant="a")
    state_path = store.deck_variant_state_path(deck_id, "a")
    before = state_path.read_bytes()
    files = store.deck_variant_files_dir(deck_id, "a")
    (files / "deck.pptx").write_bytes(b"old pptx")
    (files / "deck.pdf").write_bytes(b"old pdf with old pages")
    stale_image = store.deck_variant_slides_dir(deck_id, "a") / "slide-029.png"
    stale_image.write_bytes(b"old image")
    report = migrate_deck_reference_display(deck_id)
    assert report["slides_before"] == 29 and report["slides_after"] == 1
    assert (Path(report["backup"]) / "deck.json").read_bytes() == before
    assert (Path(report["backup"]) / "files/deck.pptx").read_bytes() == b"old pptx"
    assert list(store.deck_variant_history_dir(deck_id, "a").glob("*.json"))
    assert len(Presentation(files / "deck.pptx").slides) == 1
    assert not (files / "deck.pdf").exists()  # converter unavailable: never publish stale PDF
    assert not stale_image.exists()
    assert (files / "deck.sources.zip").exists()
    assert "REF-201" in (files / "deck.html").read_text(encoding="utf-8")
    assert report["main_evidence_unchanged"] and report["model_calls"] == 0
    repeat = migrate_deck_reference_display(deck_id)
    assert repeat["changed"] is False


def test_bad_mode_and_full_budget_are_explicit_errors():
    with pytest.raises(ValueError, match="table_display"):
        table_pages({}, {}, "typo")
    rows = [[str(i)] for i in range(220)]
    snapshot = {"table": {"headers": ["id"], "rows": rows, "row_lines": list(range(220))}}
    assert len(table_pages({"visual": "table"}, snapshot)) == 1
    with pytest.raises(ValueError, match="предел"):
        table_pages({"visual": "table"}, snapshot, "full")


def test_instruction_echo_cleanup_keeps_real_qualification_and_numeric_slides(client):
    from designer.reference_display import clean_instruction_echoes
    cover = SlideIntent(id="cover", kind=SlideKind.title, title="Учебный пилот", key_message="Результаты")
    echo = SlideIntent(id="echo", kind=SlideKind.bullets, title="Слайды уже готовы",
                       items=[Item(heading="Темы", body="Слайды по каждой теме уже подготовлены."),
                              Item(heading="Важно", body="Все значения в пилоте являются синтетическими.")])
    numeric = echo.model_copy(deep=True)
    numeric.id, numeric.key_message = "numeric", "Стоимость подготовки — 150 рублей"
    plan = DeckPlan(title="T", purpose="test", slides=[cover, echo, numeric])
    result, removed, modified = clean_instruction_echoes(plan)
    assert removed == ["echo"] and modified == ["cover"]
    assert "синтетическими" in result.slides[0].key_message
    assert result.slides[1] == numeric
    assert "синтетическими" not in plan.slides[0].key_message
