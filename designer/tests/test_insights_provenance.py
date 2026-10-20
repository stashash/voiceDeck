from __future__ import annotations

import base64
import hashlib
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches, Pt

from designer import pipeline, store
from designer.api.app import app, get_llm_client
from designer.contracts import Deck, DeckPlan
from designer.export.html import render_deck
from designer.export.pptx_deck import export_pptx
from designer.insights.suggest import Suggestion, insight_slide
from designer.layout.compose import compose
from designer.layout.scene import build_scene
from designer.parse.package import build_package


FIXTURES = Path(__file__).parent / "fixtures" / "semantic_stage2"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DESIGNER_DATA_DIR", str(tmp_path / "data"))
    return TestClient(app)


def bundle(client, names=("time_series.csv", "reference_201.csv")):
    response = client.post("/insights/plan", files=[
        ("files", (name, (FIXTURES / name).read_bytes(), "text/csv")) for name in names
    ], data={"goal": "show trend"})
    assert response.status_code == 200, response.text
    return response.json()


def choices(bundle):
    return [{"dataset_id": item["dataset_id"], "candidate_kind": item["plan"]["selected_candidate"]["kind"],
             "required_fields": item["plan"]["selected_candidate"]["required_fields"]}
            for item in bundle["datasets"]]


def materialize(client, saved):
    response = client.post("/insights/plan/materialize-bundle", data={
        "bundle_id": saved["id"], "selections": json.dumps(choices(saved)),
    })
    assert response.status_code == 200, response.text
    return response.json()


def test_bundle_persists_original_bytes_normalized_rows_profile_and_plan(client):
    saved = bundle(client)
    assert client.get(f"/insights/plan/{saved['id']}").json() == saved
    for item in saved["datasets"]:
        snapshot = store.load_dataset(item["dataset_id"])
        raw = (FIXTURES / item["file_name"]).read_bytes()
        assert base64.b64decode(snapshot["source_base64"]) == raw
        assert hashlib.sha256(raw).hexdigest() == snapshot["source_sha256"] == item["source_sha256"]
        assert len(snapshot["table"]["rows"]) == item["rows"]
        assert len(snapshot["table"]["row_lines"]) == item["rows"]
    repeated = bundle(client)
    assert [item["dataset_id"] for item in saved["datasets"]] == [item["dataset_id"] for item in repeated["datasets"]]


def test_persisted_bundle_materializes_multiple_sources_and_reaches_deck_generation(client, monkeypatch):
    saved = bundle(client)
    response = materialize(client, saved)
    assert [item["id"] for item in response["suggestions"]] == [f"s{i}" for i in range(1, 32)]
    stored = store.load_insights(response["id"])
    assert stored["bundle_id"] == saved["id"]
    assert len(stored["analyses"]) == 2
    assert stored["suggestions"][1]["visual"] == "big_number"
    reference = stored["suggestions"][2]
    assert reference["provenance"]["total_rows"] == 201
    assert len(reference["provenance"]["row_lines"]) == 201
    assert len(reference["provenance"]["displayed_row_lines"]) == 7
    pages = stored["suggestions"][2:]
    assert len(pages) == 29
    assert [page["provenance"]["page_index"] for page in pages] == list(range(1, 30))
    assert all(page["provenance"]["page_count"] == 29 for page in pages)
    assert sorted(line for page in pages for line in page["provenance"]["displayed_row_lines"]) == list(range(2, 203))
    assert pages[-1]["provenance"]["displayed_row_lines"] == [198, 199, 200, 201, 202]
    captured = {}
    monkeypatch.setattr(pipeline, "generate_deck", lambda *a, **kw: captured.update(kw))
    app.dependency_overrides[get_llm_client] = lambda: object()
    try:
        deck = client.post("/decks", json={"design_system_id": "test", "brief": "Общий обзор",
                           "variants": ["a"], "insights_id": response["id"], "insight_ids": ["s1", "s2"]})
    finally:
        app.dependency_overrides.pop(get_llm_client, None)
    assert deck.status_code == 200, deck.text
    intents = [insight_slide(suggestion) for suggestion in captured["insights"]]
    assert {intent.provenance.dataset_id for intent in intents} == {item["dataset_id"] for item in saved["datasets"]}
    plan = DeckPlan(title="Обзор", purpose="test", slides=intents)
    store.save_deck_result(deck.json()["deck_id"], Deck(id=deck.json()["deck_id"], design_system_id="test",
                           variant="a", plan=plan), [])
    restored = DeckPlan.model_validate(store.load_deck_state(deck.json()["deck_id"])["plan"])
    assert restored.slides[1].provenance == intents[1].provenance


@pytest.mark.parametrize("case", ["foreign", "duplicate", "forged_fields", "corrupt_source", "empty"])
def test_bundle_rejects_unverified_selections(client, case):
    saved = bundle(client)
    selected = choices(saved)
    if case == "foreign": selected[0]["dataset_id"] = "absent"
    elif case == "duplicate": selected.append(selected[0])
    elif case == "forged_fields": selected[0]["required_fields"] = ["invented"]
    elif case == "empty": selected = []
    else:
        path = store.data_dir() / "datasets" / f"{selected[0]['dataset_id']}.json"
        snapshot = json.loads(path.read_text(encoding="utf-8"))
        snapshot["source_base64"] = base64.b64encode(b"changed").decode()
        path.write_text(json.dumps(snapshot), encoding="utf-8")
    response = client.post("/insights/plan/materialize-bundle", data={
        "bundle_id": saved["id"], "selections": json.dumps(selected),
    })
    assert response.status_code == 422, response.text


def test_missing_rows_and_zero_have_exact_lineage(client):
    response = client.post("/insights/plan", files={"files": ("observed.csv", b"category,value\na,0\nb,10\nc,\nd,20", "text/csv")})
    saved = response.json()
    item = saved["datasets"][0]
    candidate = next(c for c in item["plan"]["candidates"] if c["kind"] in ("bar", "column"))
    response = client.post("/insights/plan/materialize-bundle", data={"bundle_id": saved["id"],
        "selections": json.dumps([{"dataset_id": item["dataset_id"], "candidate_kind": candidate["kind"],
                                    "required_fields": candidate["required_fields"]}])})
    assert response.status_code == 200, response.text
    evidence = response.json()["suggestions"][0]["provenance"]
    assert evidence["row_lines"] == [2, 3, 5]
    assert evidence["excluded_row_lines"] == [4]
    assert "3 из 4" in response.json()["suggestions"][0]["coverage"]


def test_source_is_visible_in_pptx_html_and_saved_scene(client, tmp_path):
    saved = bundle(client, ("time_series.csv",))
    response = materialize(client, saved)
    suggestion = Suggestion(**store.load_insights(response["id"])["suggestions"][0])
    intent = insight_slide(suggestion)
    intent.id = "s1"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(11), Inches(1))
    title.text = "Исходный заголовок"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    source = tmp_path / "template.pptx"
    prs.save(source)
    package = tmp_path / "package"
    ds = build_package(source, package)
    pattern = ds.patterns[0]
    spec = compose(intent, pattern, ds)
    scene = build_scene(spec, pattern, ds, package)
    source_element = next(el for el in scene.elements if el.id == "data-source")
    assert intent.title == "Выручка: +36000 USD между 2026-01 и 2026-06 (+30.00%)"
    assert intent.chart.unit == "USD"
    assert "time_series.csv" in source_element.text
    assert "Охват: 6 из 6" in source_element.text
    assert spec.viz_box[1] + spec.viz_box[3] <= .90
    pptx = export_pptx([spec], ds, package, tmp_path / "result.pptx")
    exported = Presentation(pptx).slides[0]
    chart = next(shape.chart for shape in exported.shapes if shape.has_chart)
    assert chart.value_axis.tick_labels.number_format == "#,##0"
    assert chart.value_axis.tick_labels.number_format_is_linked is False
    assert chart.value_axis.axis_title.text_frame.text == "Выручка, USD"
    text = "\n".join(shape.text for shape in exported.shapes if shape.has_text_frame)
    assert source_element.text in text
    assert suggestion.provenance.source_sha256 in exported.notes_slide.notes_text_frame.text
    deck = Deck(id="proof", design_system_id=ds.id, variant="a", plan=DeckPlan(title="Данные", purpose="test", slides=[intent]),
                specs=[spec], scenes=[scene])
    html = render_deck(deck, ds, package)
    assert "time_series.csv" in html and "Охват: 6 из 6" in html
