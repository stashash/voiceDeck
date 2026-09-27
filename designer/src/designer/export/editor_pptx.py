"""Native, editable PPTX export of the editor's canonical document."""
from __future__ import annotations

import base64
import io
from urllib.parse import urlparse

import httpx
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from designer import store
from designer.editor_documents import Document, safe_image_url


def image_bytes(url: str) -> bytes:
    if not safe_image_url(url):
        raise ValueError('Unsupported image URL')
    if url.startswith('data:'):
        data = base64.b64decode(url.split(',', 1)[1], validate=True)
    elif url.startswith('/editor/assets/'):
        data = (store.data_dir() / 'editor-assets' / url.rsplit('/', 1)[1]).read_bytes()
    elif url.startswith('/design-systems/'):
        _, _, ds, relative = url.split('/', 3)
        path = store.design_system_asset_path(ds, relative)
        if not path:
            raise ValueError('Invalid image path')
        data = path.read_bytes()
    else:
        # No arbitrary URL fetches, redirects, or private network access from documents.
        if urlparse(url).hostname not in ('upload.wikimedia.org', 'thumb.wikimedia.org'):
            raise ValueError('Unsupported image host')
        with httpx.Client(timeout=5, follow_redirects=False) as client:
            with client.stream('GET', url) as response:
                response.raise_for_status()
                chunks, size = [], 0
                for chunk in response.iter_bytes():
                    size += len(chunk)
                    if size > 8*1024*1024:
                        raise ValueError('Image exceeds 8 MB')
                    chunks.append(chunk)
                data = b''.join(chunks)
    if len(data) > 8*1024*1024:
        raise ValueError('Image exceeds 8 MB')
    with Image.open(io.BytesIO(data)) as image:
        if image.width * image.height > 24_000_000:
            raise ValueError('Image exceeds 24 megapixels')
        if image.format not in ('PNG', 'JPEG', 'WEBP'):
            raise ValueError('Only PNG, JPEG and WebP are supported')
        if image.format == 'WEBP':
            output = io.BytesIO()
            image.save(output, format='PNG')
            return output.getvalue()
        image.verify()
    return data


def rgb(value: str):
    return RGBColor.from_string(value.lstrip('#'))


def export_document(document: Document) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    px = lambda value: Pt(value * .75)
    images: dict[str, bytes] = {}
    def picture(shapes, url, x, y, width, height, cover=False):
        if url not in images:
            images[url] = image_bytes(url)
        blob = images[url]
        with Image.open(io.BytesIO(blob)) as image:
            iw, ih = image.size
        scale = max(width/iw, height/ih) if cover else min(width/iw, height/ih)
        if cover:
            item = shapes.add_picture(io.BytesIO(blob), px(x), px(y), px(width), px(height))
            item.crop_left = item.crop_right = max(0, (iw*scale-width)/(2*iw*scale))
            item.crop_top = item.crop_bottom = max(0, (ih*scale-height)/(2*ih*scale))
        else:
            item = shapes.add_picture(io.BytesIO(blob), px(x+(width-iw*scale)/2), px(y+(height-ih*scale)/2), px(iw*scale), px(ih*scale))
        return item
    for page in document.pages:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(page.background)
        if page.backgroundUrl:
            picture(slide.shapes, page.backgroundUrl, 0, 0, 1280, 720, True)
        groups = {}
        for c in page.components:
            shapes = slide.shapes
            if c.group:
                if c.group not in groups:
                    groups[c.group] = slide.shapes.add_group_shape()
                shapes = groups[c.group].shapes
            box = (px(c.x), px(c.y), px(c.width), px(c.height))
            if c.kind == 'image' and c.url:
                shape = picture(shapes, c.url, c.x, c.y, c.width, c.height, c.fit == 'cover')
            elif c.kind == 'table' and c.rows:
                shape = shapes.add_table(len(c.rows), len(c.rows[0]), *box)
                for ri, row in enumerate(c.rows):
                    for ci, value in enumerate(row):
                        cell = shape.table.cell(ri, ci)
                        cell.text = value
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(page.background if c.fill == 'transparent' else c.fill)
                        cell.vertical_anchor = MSO_ANCHOR.TOP
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.size = px(c.size)
                            paragraph.font.name = c.font or 'Arial'
                            paragraph.font.color.rgb = rgb(c.color)
                            paragraph.font.bold = ri == 0 or bool(c.bold)
            elif c.kind == 'chart' and c.labels and c.values:
                data = CategoryChartData()
                data.categories = c.labels
                data.add_series(c.text or 'Данные', c.values)
                shape = shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *box, data)
                shape.chart.has_legend = False
                shape.chart.chart_style = 10
            else:
                shape = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if c.radius else MSO_SHAPE.RECTANGLE, *box)
                shape.fill.background() if c.fill == 'transparent' else shape.fill.solid()
                if c.fill != 'transparent':
                    shape.fill.fore_color.rgb = rgb(c.fill)
                    if c.opacity is not None:
                        alpha = OxmlElement('a:alpha')
                        alpha.set('val', str(round(c.opacity*100000)))
                        shape.fill._xPr.solidFill.srgbClr.append(alpha)
                if c.radius:
                    shape.adjustments[0] = min(.5, c.radius/min(c.width, c.height))
                shape.line.fill.background()
                if c.borderWidth and c.borderColor:
                    shape.line.color.rgb = rgb(c.borderColor)
                    shape.line.width = px(c.borderWidth)
                frame = shape.text_frame
                frame.word_wrap = True
                frame.vertical_anchor = MSO_ANCHOR.TOP
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = px(16 if c.kind == 'card' else 0)
                frame.text = c.text
                for p in frame.paragraphs:
                    p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(c.align, PP_ALIGN.LEFT)
                    p.font.name = c.font or 'Arial'
                    p.font.size = px(c.size)
                    p.font.bold = c.bold if c.bold is not None else c.kind == 'title'
                    p.font.italic = bool(c.italic)
                    p.font.color.rgb = rgb(c.color)
            shape.name = c.id
            shape.rotation = c.rotation or 0
            for effect in shape._element.xpath('.//a:effectRef'):
                effect.set('idx', '0')
        slide.notes_slide.notes_text_frame.text = '\n'.join(filter(None, [page.notes, *[c.attribution for c in page.components]]))
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()
