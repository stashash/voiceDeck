"""Тесты add_chart: тип, ряды, категории, цвета, рамка, лимит рядов. Владелец: T-05, тема и кегль — T-20."""
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt

from designer.contracts import ChartSpec, ColorToken, FontToken, Margins, Series, Tokens, TypeStep
from designer.viz import palette
from designer.viz.pptx_native import add_chart

SLIDE_SIZE = (9144000, 5143500)


def _tokens() -> Tokens:
    return Tokens(
        colors=[
            ColorToken(hex="0077FF", role="accent", share=0.3, source="theme"),
            ColorToken(hex="FF7A00", role="accent_alt", share=0.2, source="theme"),
            ColorToken(hex="EEEEEE", role="surface", share=0.4, source="usage"),
            ColorToken(hex="999999", role="text_muted", share=0.1, source="usage"),
        ],
        fonts=[FontToken(family="PT Sans", role="body", share=0.8)],
        type_scale=[
            TypeStep(size_pt=28, role="heading", share=0.2),
            TypeStep(size_pt=10, role="caption", share=0.1),
        ],
        margins=Margins(left=0.03, top=0.06, right=0.03, bottom=0.06),
    )


def _blank_slide():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_SIZE
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide


def _roundtrip(prs: Presentation, tmp_path: Path) -> Presentation:
    path = tmp_path / "out.pptx"
    prs.save(str(path))
    return Presentation(str(path))


def _spec(chart_type: str, series_count: int = 2, categories_count: int = 3) -> ChartSpec:
    series = [
        Series(name=f"ряд{i}", values=[float(v + i) for v in range(categories_count)])
        for i in range(series_count)
    ]
    return ChartSpec(
        type=chart_type, categories=[f"кат{i}" for i in range(categories_count)], series=series, unit="шт",
    )


@pytest.mark.parametrize("chart_type,xl_type", [
    ("column", XL_CHART_TYPE.COLUMN_CLUSTERED),
    ("bar", XL_CHART_TYPE.BAR_CLUSTERED),
    ("line", XL_CHART_TYPE.LINE),
    ("pie", XL_CHART_TYPE.PIE),
    ("donut", XL_CHART_TYPE.DOUGHNUT),
])
def test_add_chart_creates_right_type_series_and_categories(tmp_path, chart_type, xl_type):
    tokens = _tokens()
    prs, slide = _blank_slide()
    series_count = 1 if chart_type in ("pie", "donut") else 2
    spec = _spec(chart_type, series_count=series_count)

    add_chart(slide, spec, (0.1, 0.1, 0.5, 0.4), SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    chart = reopened.slides[0].shapes[0].chart
    assert chart.chart_type == xl_type
    assert len(chart.series) == series_count
    assert len(chart.plots[0].categories) == 3


def test_add_chart_first_series_fill_is_accent(tmp_path):
    tokens = _tokens()
    prs, slide = _blank_slide()
    spec = _spec("column", series_count=2)

    add_chart(slide, spec, (0.1, 0.1, 0.5, 0.4), SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    chart = reopened.slides[0].shapes[0].chart
    assert str(chart.series[0].format.fill.fore_color.rgb) == "0077FF"


def test_add_chart_box_matches_within_one_percent(tmp_path):
    tokens = _tokens()
    prs, slide = _blank_slide()
    spec = _spec("column", series_count=1)
    box = (0.1, 0.2, 0.5, 0.3)

    add_chart(slide, spec, box, SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    shape = reopened.slides[0].shapes[0]
    slide_w, slide_h = SLIDE_SIZE
    assert shape.left == pytest.approx(box[0] * slide_w, rel=0.01)
    assert shape.top == pytest.approx(box[1] * slide_h, rel=0.01)
    assert shape.width == pytest.approx(box[2] * slide_w, rel=0.01)
    assert shape.height == pytest.approx(box[3] * slide_h, rel=0.01)


def test_add_chart_more_than_five_series_raises():
    tokens = _tokens()
    _, slide = _blank_slide()
    spec = _spec("column", series_count=6)

    with pytest.raises(ValueError):
        add_chart(slide, spec, (0.1, 0.1, 0.5, 0.4), SLIDE_SIZE, tokens)


# ---------- тема: T-20 ----------

def _dark_tokens() -> Tokens:
    return Tokens(
        colors=[
            ColorToken(hex="0B1220", role="background", share=0.5, source="theme"),
            ColorToken(hex="66CCFF", role="accent", share=0.3, source="theme"),
            ColorToken(hex="F2F2F2", role="text", share=0.2, source="usage"),
            ColorToken(hex="4A5568", role="text_muted", share=0.1, source="usage"),
        ],
        fonts=[FontToken(family="PT Sans", role="body", share=0.8)],
        type_scale=[
            TypeStep(size_pt=28, role="heading", share=0.2),
            TypeStep(size_pt=10, role="caption", share=0.1),
        ],
        margins=Margins(left=0.03, top=0.06, right=0.03, bottom=0.06),
    )


def test_add_chart_dark_theme_axis_labels_meet_contrast(tmp_path):
    tokens = _dark_tokens()
    prs, slide = _blank_slide()
    spec = _spec("column", series_count=1)

    add_chart(slide, spec, (0.1, 0.1, 0.5, 0.4), SLIDE_SIZE, tokens, theme="dark")

    reopened = _roundtrip(prs, tmp_path)
    chart = reopened.slides[0].shapes[0].chart
    category_hex = str(chart.category_axis.tick_labels.font.color.rgb)
    value_hex = str(chart.value_axis.tick_labels.font.color.rgb)
    assert palette.contrast_ratio(category_hex, "0B1220") >= 4.5
    assert palette.contrast_ratio(value_hex, "0B1220") >= 4.5


def test_add_chart_dark_theme_colors_the_legend(tmp_path):
    tokens = _dark_tokens()
    prs, slide = _blank_slide()
    spec = _spec("column", series_count=2)

    add_chart(slide, spec, (0.1, 0.1, 0.5, 0.4), SLIDE_SIZE, tokens, theme="dark")

    reopened = _roundtrip(prs, tmp_path)
    chart = reopened.slides[0].shapes[0].chart
    legend_hex = str(chart.legend.font.color.rgb)
    assert palette.contrast_ratio(legend_hex, "0B1220") >= 4.5


def test_add_chart_light_theme_leaves_font_color_unset(tmp_path):
    tokens = _dark_tokens()
    prs, slide = _blank_slide()
    spec = _spec("column", series_count=1)

    add_chart(slide, spec, (0.1, 0.1, 0.5, 0.4), SLIDE_SIZE, tokens)

    reopened = _roundtrip(prs, tmp_path)
    chart = reopened.slides[0].shapes[0].chart
    assert chart.font.color.type is None


# ---------- кегль подписей по высоте рамки: T-20 ----------

def _tokens_with_body_step() -> Tokens:
    tokens = _tokens()
    tokens.type_scale.append(TypeStep(size_pt=18, role="body", share=0.3))
    return tokens


def _chart_label_size(tmp_path, tokens, box, name: str):
    prs, slide = _blank_slide()
    add_chart(slide, _spec("column", series_count=1), box, SLIDE_SIZE, tokens)
    path = tmp_path / name
    prs.save(str(path))
    return Presentation(str(path)).slides[0].shapes[0].chart


def test_add_chart_label_size_grows_with_the_frame(tmp_path):
    """Просторной рамке достаётся ступень крупнее: мелкая подпись там читается как брак."""
    tokens = _tokens_with_body_step()
    roomy = _chart_label_size(tmp_path, tokens, (0.1, 0.0, 0.5, 1.0), "roomy.pptx")
    tight = _chart_label_size(tmp_path, tokens, (0.1, 0.1, 0.5, 0.3), "tight.pptx")
    assert roomy.font.size > tight.font.size
    assert roomy.font.size == Pt(18)


def test_add_chart_labels_never_go_below_the_smallest_step(tmp_path):
    """Тесная рамка не опускает подписи ниже самой мелкой ступени шкалы."""
    tokens = _tokens_with_body_step()
    small = min(step.size_pt for step in tokens.type_scale)
    chart = _chart_label_size(tmp_path, tokens, (0.1, 0.1, 0.5, 0.1), "small.pptx")
    assert chart.font.size == Pt(small)
    assert chart.value_axis.tick_labels.font.size == Pt(small)
    assert chart.category_axis.tick_labels.font.size == Pt(small)
