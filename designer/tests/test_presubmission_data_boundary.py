"""Adversarial evidence for the data and rendering boundary, independent of demo fixtures."""
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation

from designer.audit.checks_integrity import check_chart_data, _is_placeholder
from designer.contracts import ChartSpec, DesignSystem, Element, Margins, Scene, Series, Tokens
from designer.export.svg_chart import render_chart
from designer.insights.plan import AnalyticalGoal, plan_table
from designer.insights.table import read_table
from designer.viz.pptx_native import add_chart


def tokens():
    return Tokens(colors=[], fonts=[], type_scale=[], margins=Margins(left=0, right=0, top=0, bottom=0))


def plan(text, goal="compare"):
    return plan_table(read_table(text.encode(), "adversarial.csv"), goal)


def test_percent_level_change_is_in_points_with_relative_change_separately_named():
    candidate = plan("period,rate_%\n2026-02,12\n2026-01,10\n", "trend").selected_candidate
    assert candidate.row_lines == (3, 2)
    assert "+2 п. п." in candidate.headline
    assert "+20.00%" in candidate.headline
    assert "от 10 %" in candidate.statement and "до 12 %" in candidate.statement
    assert "= +2 п. п." in candidate.formula


@pytest.mark.parametrize("header,rows", [
    ("amount", "A,100 $\nB,100 ₽"),
    ("amount", "A,100 $\nB,200 $"),
    ("amount_USD", "A,100 $\nB,100 €"),
    ("rate", "A,10%\nB,20%"),
])
def test_units_cannot_be_stripped_or_compared_as_same_quantity(header, rows):
    result = plan(f"category,{header}\n{rows}\n", "shares")
    assert result.selected_candidate.kind == "table"
    assert any("единицы в ячейках" in warning for warning in result.warnings)
    assert all(not c.renderable for c in result.candidates if c.kind != "table")


def test_explicit_matching_header_and_cell_units_remain_chartable():
    result = plan("category,amount_USD\nA,100 $\nB,200 $\n")
    assert result.selected_candidate.kind == "bar"


def test_percentage_levels_are_not_renormalized_to_a_different_denominator():
    result = plan("category,rate_%\nA,20\nB,30\n", "shares")
    assert not any(c.kind == "pie" for c in result.candidates)
    assert any("исходный знаменатель" in warning for warning in result.warnings)
    complete = plan("category,share_%\nA,20\nB,80\n", "shares")
    assert complete.selected_candidate.kind == "pie"
    assert complete.selected_candidate.denominator == 100


def test_explicit_multimeasure_selection_is_not_silently_narrowed():
    result = plan("period,a_USD,b_USD\n2026-01,1,2\n2026-02,3,4\n",
                  AnalyticalGoal("trend", measures=("a_USD", "b_USD")))
    assert result.selected_candidate.kind == "table"
    assert all(not c.renderable for c in result.candidates if c.kind != "table")
    assert any("пропадают выбранные показатели" in warning for warning in result.warnings)


def test_native_horizontal_bar_titles_follow_physical_axes_and_data_survives(tmp_path):
    spec = ChartSpec(type="bar", categories=["North", "South"], unit="USD",
                     x_axis_title="Revenue, USD", y_axis_title="Region",
                     series=[Series(name="Revenue", values=[-20, 10])])
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_chart(slide, spec, (.1, .1, .8, .7), (prs.slide_width, prs.slide_height), tokens())
    path = tmp_path / "signed-bar.pptx"
    prs.save(path)
    chart = Presentation(path).slides[0].shapes[0].chart
    assert chart.category_axis.axis_title.text_frame.text == "Region"
    assert chart.value_axis.axis_title.text_frame.text == "Revenue, USD"
    assert tuple(chart.series[0].values) == (-20, 10)


@pytest.mark.parametrize("kind", ["bar", "line", "column"])
@pytest.mark.parametrize("values", [[-10, 0, 20], [-30, -20, -10], [0, 0, 0], [1e308, -1e308, 0]])
def test_signed_svg_geometry_is_bounded_and_preserves_axis_titles(kind, values):
    spec = ChartSpec(type=kind, categories=["A", "B", "C"], x_axis_title="X <axis>", y_axis_title="Y & unit",
                     series=[Series(name="s", values=values)])
    root = ET.fromstring(render_chart(spec, []))
    assert {el.text for el in root.iter("text")} >= {"X <axis>", "Y & unit"}
    if kind == "line":
        points = list(root.iter("circle"))
        assert len(points) == 3
        assert all(20 <= float(el.attrib["cy"]) <= 310 for el in points)
        # Larger values are always higher, even when all values are negative.
        for i, a in enumerate(values):
            for j, b in enumerate(values):
                if a < b:
                    assert float(points[i].attrib["cy"]) > float(points[j].attrib["cy"])
    else:
        bars = [el for el in root.iter("rect") if el.attrib.get("class") == "chart-bar"]
        assert len(bars) == 3
        for el in bars:
            x, y, w, h = (float(el.attrib[k]) for k in ("x", "y", "width", "height"))
            assert w >= 0 and h >= 0
            assert 49.9 <= x <= x + w <= 580.1
            assert 19.9 <= y <= y + h <= 310.1


@pytest.mark.parametrize("kind", ["pie", "donut"])
def test_full_observed_share_is_visible_when_other_categories_are_zero(kind):
    spec = ChartSpec(type=kind, categories=["A", "B"], series=[Series(name="s", values=[10, 0])])
    root = ET.fromstring(render_chart(spec, []))
    circles = [el for el in root.iter("circle") if el.attrib.get("class") == "chart-slice"]
    assert len(circles) == 1 and float(circles[0].attrib["r"]) > 0


@pytest.mark.parametrize("kind,values", [("line", [1]), ("bar", [1]), ("column", [1]),
                                         ("line", [1, float("nan")]), ("pie", [-1, 2]),
                                         ("pie", [0, 0]), ("pie", [1e308, 1e308])])
def test_bad_legacy_chart_data_is_flagged_and_both_exports_refuse_it(kind, values):
    spec = ChartSpec(type=kind, categories=["A", "B"], series=[Series(name="s", values=values)], unit="units")
    scene = Scene(slide_id="s1", pattern_id="p1", elements=[Element(id="c1", type="chart", box=(0, 0, 1, 1), chart=spec)])
    ds = DesignSystem(id="d", source_file="x.pptx", slide_size_emu=(100, 100), tokens=tokens())
    findings = check_chart_data([scene], ds)
    assert len(findings) == 1 and findings[0].severity == "error"
    with pytest.raises(ValueError):
        render_chart(spec, [])
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    with pytest.raises(ValueError):
        add_chart(slide, spec, (.1, .1, .8, .7), (prs.slide_width, prs.slide_height), tokens())


def test_legitimate_russian_text_is_not_a_placeholder_but_literal_placeholder_is():
    assert not _is_placeholder("Текст должен сохранять исходные данные")
    assert not _is_placeholder("Проверяем заголовок и источник")
    assert _is_placeholder("Текст") and _is_placeholder("Заголовок") and _is_placeholder("TODO: заменить")
