from pathlib import Path
import json
import pytest
from designer.brandbook_import import parse_brandbook_source, BrandbookImportError
from designer.brand_assets import save_brand_assets
from designer.brandbook_json import parse_brand_tokens, resolve_tokens
from designer.contracts import BrandRules, ChartStyle, DesignSystem, Element, TextStyle
from designer.layout.brand import apply_brand

EXAMPLES = Path(__file__).resolve().parents[2] / "frontend/public/brandbooks"


def test_zip_applies_explicit_logo_grid_typography_and_chart_rules(tmp_path):
    result = parse_brandbook_source((EXAMPLES / "local-brandbook.zip").read_bytes(), "brand.zip")
    assert len(result.tokens.guides_x) == 12
    assert result.tokens.chart_style.legend_position == "bottom"
    assets, rules, metadata = save_brand_assets(tmp_path, result)
    assert rules.logo_asset and rules.clear_space == .015
    assert any(asset.id == rules.logo_asset and asset.kind == "logo" for asset in assets)
    assert all((tmp_path / asset.path).is_file() for asset in assets)
    ds = DesignSystem(id="example", source_file="x.pptx", slide_size_emu=(12192000, 6858000), tokens=result.tokens, assets=assets, brand_rules=rules)
    title = Element(id="title", type="text", role="title", box=(.06, .05, .90, .10), text="Headline", style=TextStyle(family="Arial"))
    scene = apply_brand([title], ds)
    logo = next(el for el in scene if el.role == "brand_logo")
    assert title.box[0] + title.box[2] <= logo.box[0] - rules.clear_space + 1e-9
    assert title.style.family == "Play"
    assert len(metadata) == 2


def test_image_materials_without_token_json_recover_palette_and_logo(tmp_path):
    result = parse_brandbook_source((EXAMPLES / "local-assets-only.zip").read_bytes(), "brand.zip")
    assert any(color.hex == "713CDA" for color in result.tokens.colors)
    _, rules, _ = save_brand_assets(tmp_path, result)
    assert rules.logo_asset


def test_pdf_recovers_evidenced_roles_fonts_grid_and_logo(tmp_path):
    result = parse_brandbook_source((EXAMPLES / "local-brandbook.pdf").read_bytes(), "guide.pdf")
    assert any(color.role == "accent" and color.hex == "713CDA" for color in result.tokens.colors)
    assert any(font.role == "heading" and font.family == "Play" for font in result.tokens.fonts)
    assert result.tokens.margins.left == .06
    assert len(result.tokens.guides_x) == 12
    assert result.assets and result.rules["logo"]


def test_bad_logo_box_and_grid_rules_fail_before_persistence():
    for rules in ({"logo_box": [.95, .02, .2, .1]}, {"grid_columns": 24, "grid_gutter": .1}, {"logo": "../private.png"}):
        with pytest.raises(BrandbookImportError):
            parse_brandbook_source(json.dumps({"tokens": {"colors": {"accent": "713CDA"}}, "rules": rules}).encode(), "bad.json")


def test_chart_style_and_guides_follow_selected_priority():
    template = parse_brand_tokens({"colors": {"accent": "0000FF"}, "chart_style": {"legend_position": "right"}, "guides_x": [.1, .5]})
    brand = parse_brand_tokens({"colors": {"accent": "713CDA"}, "chart_style": {"legend_position": "bottom"}, "guides_x": [.06, .4]})
    assert resolve_tokens(template, brand).chart_style.legend_position == "bottom"
    assert resolve_tokens(template, brand).guides_x == [.06, .4]
    assert resolve_tokens(template, brand, False).chart_style.legend_position == "right"


def test_logo_and_legend_are_present_in_actual_pptx_export(tmp_path, templates):
    from pptx import Presentation
    from pptx.enum.chart import XL_LEGEND_POSITION
    from designer.parse.package import build_package
    from designer.layout.match import choose_pattern
    from designer.layout.compose import compose
    from designer.export.pptx_deck import export_pptx
    from designer.contracts import ChartSpec, Series, SlideIntent, SlideKind
    package = tmp_path / "package"
    ds = build_package(templates[0], package)
    imported = parse_brandbook_source((EXAMPLES / "local-brandbook.zip").read_bytes(), "brand.zip")
    assets, rules, _ = save_brand_assets(package, imported)
    ds.tokens = resolve_tokens(ds.tokens, imported.tokens)
    ds.assets.extend(assets)
    ds.brand_rules = rules
    intent = SlideIntent(id="s1", kind=SlideKind.chart, title="Выручка выросла на 30%",
                         chart=ChartSpec(type="grouped_column", categories=["A", "B"],
                                         series=[Series(name="План", values=[10, 20]), Series(name="Факт", values=[12, 26])]))
    pattern = choose_pattern(intent, ds, [])
    path = export_pptx([compose(intent, pattern, ds)], ds, package, tmp_path / "result.pptx")
    slide = Presentation(path).slides[0]
    assert any(shape.shape_type == 13 and shape.image.sha1.startswith(rules.logo_asset) for shape in slide.shapes)
    chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
    assert chart.legend.position == XL_LEGEND_POSITION.BOTTOM
    assert chart.font.name == "Play"
